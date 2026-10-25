"""Génère la fiche Excel de reporting commercial de GIC AGROPELC.

Usage (depuis ce dossier) :
    python generer_fiche_commercial.py                 # écrit les deux fichiers dans ../
    python generer_fiche_commercial.py --exemple FICHIER   # fichier d'essai avec données fictives (vérification)

Deux livrables, issus du même code :
  - MODELE_VIERGE : à remettre à chaque commercial ;
  - DEMARRAGE_Sept2026 : le même classeur, déjà garni avec les informations de la fiche
    « ETAT DEPLOIEMENT SEPTEMBRE » (26 PDV, planning des semaines 1, 2 et 4). Le petit tableau
    « RECAP » des SIM de cette fiche n'est pas repris : le commercial l'a indiqué comme une erreur.
"""

from __future__ import annotations

import argparse
import datetime as dt
import math
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

sys.path.insert(0, str(Path(__file__).resolve().parent))

import _journaux as J  # noqa: E402
import _rapports as R  # noqa: E402
from _donnees import (  # noqa: E402
    A_RENSEIGNER,
    CAUSES,
    DEBUT_SEMAINE_FICHE,
    ETAPES,
    FICHE_PDV,
    FONCTIONS,
    FRAIS_TYPES,
    LOCALISATION_PROPRE,
    METRIQUES,
    MOTIFS_PERTE,
    MOUVEMENTS_STOCK,
    NOM_IMPOSE,
    NON,
    OBJETS_VISITE,
    OUI,
    OUI_NON,
    PAIEMENTS,
    PAIEMENTS_PREVUS,
    PREFIXE_NOM,
    PRODUITS,
    QUANTITES,
    REMISES,
    RESULTATS,
    SEMAINES_MOIS,
    SOURCES,
    STATUTS,
    STATUTS_COMMANDE,
    ST_PROSPECT,
    TYPE_DEDUIT_DU_LIBELLE,
    TYPE_FICHE_VERS_TYPE,
    TYPES_CLIENT,
    TYPES_CONTACT,
    TYPES_OPERATION,
    TYPES_PERIODE,
    UNITES,
    ZONES,
    cle_localisation,
    quartier_depuis_localisation,
)
from _fiche_client import fiche_client  # noqa: E402
from _outils import (  # noqa: E402
    BORDURE,
    COULEURS_ONGLETS,
    FMT_DATE,
    GRIS_CALC,
    JAUNE,
    VERT_FONCE,
    ARDOISE,
    BLEU,
    _appliquer_validation,
    construire_journal,
    definir_nom,
    entete_tableau,
    fill,
    mise_en_page,
    police,
    section,
)
from _tableau_bord import tableau_de_bord  # noqa: E402

VERSION = "2.0"
DATE_VERSION = "08/10/2026"
CAPACITE_LISTE = 40
PREMIERE_LISTE = 6


# --- Paramètres ---------------------------------------------------------------------------------------------------------------------


def feuille_parametres(wb) -> None:
    ws = wb.create_sheet("Paramètres")
    ws.sheet_properties.tabColor = COULEURS_ONGLETS["guide"]
    for col, largeur in zip("ABCDEF", (3, 46, 34, 40, 24, 18)):
        ws.column_dimensions[col].width = largeur
    ws["B1"] = "PARAMÈTRES — qui remplit cette fiche ?"
    ws["B1"].font = police(16, True, couleur=VERT_FONCE)
    ws["B2"] = "À remplir une seule fois, avant la première saisie. Les cellules jaunes se saisissent ; les cellules grises se calculent."
    ws["B2"].font = police(10, False, True, ARDOISE)
    ws["B3"] = "◄ Sommaire"
    ws["B3"].hyperlink = "#'LISEZ-MOI'!A1"
    ws["B3"].font = Font(name="Calibri", size=10, bold=True, color=BLEU, underline="single")
    ws.row_dimensions[1].height = 26

    section(ws, 4, "1. MON IDENTITÉ", 2, 6)
    champs = [
        (5, "Nom et prénom du commercial", "Commercial_Nom", None, "Ex. Awa Mbarga"),
        (6, "Code commercial", "Commercial_Code", None, "Ex. COM01 — à saisir une seule fois, avant tout le reste"),
        (7, "Fonction", "Commercial_Fonction", {"type": "liste", "nom": "L_Fonction"}, "Choisir dans la liste"),
        (8, "Zone / secteur principal", "Commercial_Zone", {"type": "liste", "nom": "L_Zone", "souple": True}, "Choisir dans la liste"),
        (9, "Téléphone", "Commercial_Tel", None, ""),
        (10, "Responsable direct", "Commercial_Resp", None, ""),
        (11, "Date de début d'utilisation de la fiche", "Commercial_Debut", {"type": "date_passee"}, "Choisir dans la liste de dates"),
    ]
    for ligne, libelle, nom, dv, aide in champs:
        c = ws.cell(ligne, 2, libelle)
        c.font = police(10, True, couleur=VERT_FONCE)
        c.border = BORDURE
        saisie = ws.cell(ligne, 3)
        saisie.fill = fill(JAUNE)
        saisie.border = BORDURE
        saisie.font = police(10, True)
        if nom == "Commercial_Debut":
            saisie.number_format = FMT_DATE
        if nom in ("Commercial_Tel", "Commercial_Code"):
            saisie.number_format = "@"
        if aide:
            ws.cell(ligne, 4, aide).font = police(9, False, True, ARDOISE)
        if dv:
            _appliquer_validation(ws, "C", ligne, ligne, dv)
        definir_nom(wb, nom, f"'Paramètres'!$C${ligne}")
    ws.cell(12, 2, "Identifiant utilisé dans toutes les feuilles").font = police(10, True, couleur=VERT_FONCE)
    ws.cell(12, 2).border = BORDURE
    ident = ws.cell(12, 3, f'=IF(C6="","{A_RENSEIGNER}",C6)')
    ident.fill = fill(GRIS_CALC)
    ident.border = BORDURE
    ident.font = police(10, True, couleur=ARDOISE)
    ws.cell(12, 4, "Calculé : tant que le code est vide, « À RENSEIGNER » apparaît dans les feuilles.").font = police(9, False, True, ARDOISE)
    definir_nom(wb, "Commercial_Id", "'Paramètres'!$C$12")

    section(ws, 14, "2. SEUILS D'ALERTE (colonne « Alerte » de la feuille Clients)", 2, 6)
    for ligne, libelle, valeur, nom in (
        (15, "Jours sans achat avant « client inactif »", 30, "P_JoursInactif"),
        (16, "Jours avant « prospect jamais visité »", 7, "P_JoursNonVisite"),
    ):
        c = ws.cell(ligne, 2, libelle)
        c.font = police(10, True, couleur=VERT_FONCE)
        c.border = BORDURE
        saisie = ws.cell(ligne, 3, valeur)
        saisie.fill = fill(JAUNE)
        saisie.border = BORDURE
        saisie.font = police(10, True)
        saisie.alignment = Alignment(horizontal="left")
        _appliquer_validation(ws, "C", ligne, ligne, {"type": "entier", "min": 1, "max": 365})
        definir_nom(wb, nom, f"'Paramètres'!$C${ligne}")

    section(ws, 18, "3. MON ÉQUIPE (pour le responsable commercial)", 2, 6)
    ws.cell(19, 2, "Listez ici les commerciaux suivis : ils alimentent les listes « Commercial » et la feuille Équipe. Un commercial seul peut laisser ce tableau vide.").font = police(9, False, True, ARDOISE)
    entete_tableau(ws, 20, ["Code", "Nom et prénom", "Fonction", "Zone", "Téléphone"], col_debut=2, hauteur=22)
    premiere, nb = 21, 15
    for r in range(premiere, premiere + nb):
        for col in range(2, 7):
            c = ws.cell(r, col)
            c.border = BORDURE
            c.font = police(10)
            c.fill = fill(JAUNE) if col == 2 else fill("FFFFFF")
            if col in (2, 6):
                c.number_format = "@"
        ws.row_dimensions[r].height = 18
    _appliquer_validation(ws, "D", premiere, premiere + nb - 1, {"type": "liste", "nom": "L_Fonction"})
    _appliquer_validation(ws, "E", premiere, premiere + nb - 1, {"type": "liste", "nom": "L_Zone", "souple": True})
    for cle, col in (("Equipe_Code", "B"), ("Equipe_Nom", "C"), ("Equipe_Fonction", "D"), ("Equipe_Zone", "E"), ("Equipe_Tel", "F")):
        definir_nom(wb, cle, f"'Paramètres'!${col}${premiere}:${col}${premiere + nb - 1}")
    definir_nom(wb, "L_Equipe", f"OFFSET('Paramètres'!$B${premiere},0,0,MAX(1,COUNTA('Paramètres'!$B${premiere}:$B${premiere + nb - 1})),1)")
    ws.freeze_panes = "A4"
    mise_en_page(ws, "portrait", f"A1:F{premiere + nb}")
    ws.sheet_view.showGridLines = False


# --- Listes ----------------------------------------------------------------------------------------------------------------------------


def feuille_listes(wb) -> dict:
    ws = wb.create_sheet("Listes")
    ws.sheet_properties.tabColor = COULEURS_ONGLETS["reference"]
    ws["A1"] = "LISTES DE RÉFÉRENCE — les choix des listes déroulantes"
    ws["A1"].font = police(16, True, couleur=VERT_FONCE)
    ws["A2"] = (
        "Pour ajouter un choix, écrivez-le à la suite d'une liste, sans ligne vide. Ne renommez pas un choix déjà utilisé. "
        "La direction renseigne ici les prix de référence des produits (colonne « Prix de référence »). Les listes de dates, d'heures et de mois se mettent à jour seules."
    )
    ws["A2"].font = police(10, False, True, ARDOISE)
    ws["A3"] = "◄ Sommaire"
    ws["A3"].hyperlink = "#'LISEZ-MOI'!A1"
    ws["A3"].font = Font(name="Calibri", size=10, bold=True, color=BLEU, underline="single")

    blocs = [
        ("Fonctions", "L_Fonction", FONCTIONS),
        ("Types de client", "L_TypeClient", TYPES_CLIENT),
        ("Sources de prospect", "L_Source", SOURCES),
        ("Statuts du compte (automatiques)", "L_Statut", STATUTS),
        ("Étapes du pipeline", "L_Etape", ETAPES),
        ("Types de contact", "L_TypeContact", TYPES_CONTACT),
        ("Objets de visite", "L_ObjetVisite", OBJETS_VISITE),
        ("Résultats de visite", "L_Resultat", RESULTATS),
        ("Motifs de perte", "L_MotifPerte", MOTIFS_PERTE),
        ("Unités", "L_Unite", UNITES),
        ("Modes de paiement", "L_Paiement", PAIEMENTS),
        ("Types d'opération", "L_TypeOp", TYPES_OPERATION),
        ("Statuts de commande (automatiques)", "L_StatutCmd", STATUTS_COMMANDE),
        ("Paiements prévus", "L_PaiementPrevu", PAIEMENTS_PREVUS),
        ("Mouvements de stock", "L_Mouvement", MOUVEMENTS_STOCK),
        ("Causes de casse ou de perte", "L_Cause", CAUSES),
        ("Quartiers / zones", "L_Zone", ZONES),
        ("Indicateurs d'objectif", "L_Metrique", METRIQUES),
        ("Types de période", "L_TypePeriode", TYPES_PERIODE),
        ("Oui / Non (cases à cocher)", "L_OuiNon", OUI_NON),
        ("Semaines du mois", "L_SemMois", SEMAINES_MOIS),
    ]
    col = 1
    derniere = PREMIERE_LISTE + CAPACITE_LISTE - 1

    def bloc(titre: str, nom: str, entetes: list, lignes: list, largeurs: list) -> int:
        nonlocal col
        ws.cell(4, col, titre).font = police(11, True, couleur=VERT_FONCE)
        for i, (en, lg) in enumerate(zip(entetes, largeurs)):
            c = ws.cell(5, col + i, en)
            c.font = police(10, True, couleur="FFFFFF")
            c.fill = fill(VERT_FONCE if i == 0 else ARDOISE)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDURE
            ws.column_dimensions[get_column_letter(col + i)].width = lg
        for k in range(CAPACITE_LISTE):
            for i in range(len(entetes)):
                c = ws.cell(PREMIERE_LISTE + k, col + i)
                c.border = BORDURE
                c.font = police(10, False, False, "000000" if i == 0 else ARDOISE)
                if i == 0:
                    c.fill = fill("FFFFFF")
        for k, valeurs in enumerate(lignes):
            for i, v in enumerate(valeurs):
                if v not in (None, ""):
                    ws.cell(PREMIERE_LISTE + k, col + i, v)
        lettre = get_column_letter(col)
        definir_nom(wb, nom, f"OFFSET('Listes'!${lettre}${PREMIERE_LISTE},0,0,MAX(1,COUNTA('Listes'!${lettre}${PREMIERE_LISTE}:${lettre}${derniere})),1)")
        debut = col
        col += len(entetes) + 1
        ws.column_dimensions[get_column_letter(col - 1)].width = 3
        return debut

    positions = {}
    for titre, nom, items in blocs:
        positions[nom] = bloc(titre, nom, ["Libellé", "Code (indicatif)"], [(lib, code) for lib, code in items], [30, 22])
    # Produits : libellé, unité, prix de référence, code
    debut = bloc(
        "Produits",
        "L_Produit",
        ["Produit", "Unité", "Prix de référence (XAF) — à renseigner par la direction", "Code (indicatif)"],
        [(lib, unite, prix, code) for lib, unite, prix, code in PRODUITS],
        [34, 12, 22, 22],
    )
    positions["L_Produit"] = debut
    a, d = get_column_letter(debut), get_column_letter(debut + 3)
    definir_nom(wb, "T_Produits", f"'Listes'!${a}${PREMIERE_LISTE}:${d}${derniere}")
    _appliquer_validation(ws, get_column_letter(debut + 1), PREMIERE_LISTE, derniere, {"type": "liste", "nom": "L_Unite"})
    for k in range(CAPACITE_LISTE):
        ws.cell(PREMIERE_LISTE + k, debut + 2).number_format = "#,##0"

    # Sélecteurs : dates, lundis, mois, heures, quantités, frais, remises (calculés ou fixes, sans trou)
    def bloc_selecteur(titre: str, nom: str, valeurs: list, fmt: str, largeur: float = 22) -> None:
        nonlocal col
        ws.cell(4, col, titre).font = police(11, True, couleur=VERT_FONCE)
        c = ws.cell(5, col, "Choix")
        c.font = police(10, True, couleur="FFFFFF")
        c.fill = fill(VERT_FONCE)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDURE
        ws.column_dimensions[get_column_letter(col)].width = largeur
        for k, valeur in enumerate(valeurs):
            cell = ws.cell(PREMIERE_LISTE + k, col, valeur)
            cell.number_format = fmt
            cell.border = BORDURE
            cell.font = police(10, False, False, ARDOISE)
            cell.fill = fill(GRIS_CALC)
        lettre = get_column_letter(col)
        definir_nom(wb, nom, f"'Listes'!${lettre}${PREMIERE_LISTE}:${lettre}${PREMIERE_LISTE + len(valeurs) - 1}")
        col += 2
        ws.column_dimensions[get_column_letter(col - 1)].width = 3

    lundi = "TODAY()-WEEKDAY(TODAY(),2)+1"
    bloc_selecteur("Dates passées (aujourd'hui en tête)", "L_DatesPassees", [f"=TODAY()-{k}" for k in range(60)], "ddd dd/mm/yyyy")
    bloc_selecteur("Dates à venir (aujourd'hui en tête)", "L_DatesAvenir", [f"=TODAY()+{k}" for k in range(90)], "ddd dd/mm/yyyy")
    bloc_selecteur("Lundis (semaine en cours, puis passées)", "L_LundisPasses", [f"={lundi}-7*{k}" for k in range(30)], "dd/mm/yyyy")
    bloc_selecteur("Lundis (semaine en cours, puis à venir)", "L_LundisAvenir", [f"={lundi}+7*{k}" for k in range(26)], "dd/mm/yyyy")
    bloc_selecteur("Mois (12 passés, en cours, 6 à venir)", "L_Mois", [f"=DATE(YEAR(TODAY()),MONTH(TODAY())-12+{k},1)" for k in range(19)], "mmmm yyyy")
    bloc_selecteur("Heures (tous les quarts d'heure)", "L_Heures", [(300 + 15 * k) / 1440 for k in range(69)], "hh:mm", 14)
    bloc_selecteur("Quantités proposées", "L_Quantites", QUANTITES, "General", 14)
    bloc_selecteur("Frais proposés (XAF)", "L_Frais", FRAIS_TYPES, "#,##0", 14)
    bloc_selecteur("Remises proposées (XAF)", "L_Remises", REMISES, "#,##0", 14)
    ws.row_dimensions[5].height = 48
    ws.freeze_panes = "A6"
    ws.sheet_view.zoomScale = 90
    return positions


# --- Données de départ ---------------------------------------------------------------------------------------------------------------


def _poser(ws, spec, ligne: int, **valeurs) -> None:
    for cle, valeur in valeurs.items():
        ws[f"{spec.lettres[cle]}{ligne}"] = valeur


def _etiquette(nom: str, numero: int) -> str:
    """Libellé de liste d'un client : « Resto La Paix [001] » (le numéro relie toutes les feuilles)."""
    return f"{nom} [{numero:03d}]"


def garnir_demarrage(wb, specs: dict) -> None:
    """Reprend la fiche « ETAT DEPLOIEMENT SEPTEMBRE » : une ligne complète par point de vente."""
    ws_p = wb["Paramètres"]
    ws_p["C6"] = "COM01"  # code de départ ; à remplacer par le code réel du commercial
    cl, pl = specs["Clients"], specs["Planning"]
    ws_c, ws_pl = wb["Clients"], wb["Planning"]

    lignes = []
    for sem in (1, 2, 4):
        for potentiel, loc, tel, type_fiche in FICHE_PDV[sem]:
            cle = cle_localisation(loc)
            type_client = TYPE_FICHE_VERS_TYPE.get(type_fiche or "") or TYPE_DEDUIT_DU_LIBELLE.get(cle, "Autre")
            if cle in NOM_IMPOSE:
                base = NOM_IMPOSE[cle]
            else:
                prefixe = PREFIXE_NOM.get(type_client, "PDV")
                base = f"{prefixe} {LOCALISATION_PROPRE[cle]}"
            lignes.append((sem, potentiel, cle, tel, type_client, base))

    total = Counter(ligne[-1] for ligne in lignes)
    vus: Counter = Counter()
    for i, (sem, potentiel, cle, tel, type_client, base) in enumerate(lignes):
        r = 6 + i
        vus[base] += 1
        nom = base if total[base] == 1 else f"{base} ({vus[base]})"
        valeurs = dict(
            nom=nom,
            type=type_client,
            tel=tel,
            repere=LOCALISATION_PROPRE[cle],
            source="Prospection terrain",
            etape="Nouveau",
            creation=dt.date(*DEBUT_SEMAINE_FICHE[sem]),
            potentiel=potentiel,
            obs=f"Repris de la fiche de septembre 2026 (semaine {sem}).",
        )
        quartier = quartier_depuis_localisation(cle)
        if quartier:
            valeurs["quartier"] = quartier
        _poser(ws_c, cl, r, **valeurs)
        _poser(
            ws_pl,
            pl,
            r,
            mois=dt.date(2026, 9, 1),
            semmois=sem,
            client=_etiquette(nom, i + 1),
            objet="Prospection / découverte",
            remarque="Repris de la fiche de septembre 2026.",
        )


def garnir_exemple(wb, specs: dict) -> None:
    """Jeu d'essai FICTIF, pour vérifier les formules (jamais livré)."""
    wb["Paramètres"]["C5"] = "Test Commercial"
    wb["Paramètres"]["C6"] = "T1"
    wb["Paramètres"]["C7"] = "Commercial terrain"
    for i, (code, nom) in enumerate((("T1", "Test Commercial"), ("T2", "Deuxième Test"))):
        wb["Paramètres"].cell(21 + i, 2, code)
        wb["Paramètres"].cell(21 + i, 3, nom)

    D = dt.date
    cl = specs["Clients"]
    ws = wb["Clients"]
    clients = [
        dict(nom="Resto La Paix", type="Restaurant", tel="677000001", quartier="Akwa", creation=D(2026, 10, 5), produit="Œufs frais - alvéole de 30", potentiel=10),
        dict(nom="Hôtel Dovil", type="Hôtel", tel="677000002", quartier="Bonanjo", etape="Intéressé", creation=D(2026, 10, 6)),
        dict(nom="Mme Ndongo", type="Particulier / ménage", tel="677000003", quartier="Deido", etape="Nouveau", creation=D(2026, 10, 7)),
        dict(nom="Grossiste Mboppi", type="Grossiste", tel="677000004", quartier="New Bell", creation=D(2026, 9, 20)),
    ]
    for i, c in enumerate(clients):
        _poser(ws, cl, 6 + i, **c)
    # Second commercial (fusion de deux fiches) : son client n° 1 porte le même numéro que celui de T1.
    _poser(ws, cl, 12, nom="Boulangerie T2", type="Boulangerie / pâtisserie", creation=D(2026, 10, 6), commercial="T2")
    ws["A12"] = 1

    L1, L2, L3, L4 = _etiquette("Resto La Paix", 1), _etiquette("Hôtel Dovil", 2), _etiquette("Mme Ndongo", 3), _etiquette("Grossiste Mboppi", 4)
    LT2 = _etiquette("Boulangerie T2", 1)
    vc, ws = specs["Visites_Contacts"], wb["Visites_Contacts"]
    visites = [
        dict(date=D(2026, 10, 5), heure=dt.time(8, 30), type="Visite sur place", client=L1, objet="Prise de commande", resultat="Commande prise", actiondate=D(2026, 10, 12), actiontexte="Rappeler"),
        dict(date=D(2026, 10, 5), heure=dt.time(10, 0), type="Visite sur place", client=L2, objet="Prospection / découverte", resultat="À relancer", actiondate=D(2026, 10, 8), actiontexte="Relancer le gérant"),
        dict(date=D(2026, 10, 6), heure=dt.time(9, 0), type="Appel téléphonique", client=L3, objet="Prospection / découverte", resultat="Rendez-vous fixé"),
        dict(date=D(2026, 10, 6), heure=dt.time(11, 0), type="Visite sur place", client=L1, objet="Livraison", resultat="Livraison effectuée"),
        dict(date=D(2026, 10, 7), heure=dt.time(9, 30), type="Visite sur place", client=L4, objet="Encaissement / recouvrement", resultat="Paiement encaissé"),
        dict(date=D(2026, 10, 7), heure=dt.time(14, 0), type="Visite sur place", client=L3, objet="Présentation produits et prix", resultat="Absent / fermé"),
        dict(date=D(2026, 10, 8), heure=dt.time(10, 15), type="Visite sur place", client=L2, objet="Prise de commande", resultat="Commande prise"),
    ]
    for i, v in enumerate(visites):
        _poser(ws, vc, 6 + i, **v)
    _poser(ws, vc, 20, date=D(2026, 10, 6), type="Visite sur place", client=LT2, objet="Prise de commande", resultat="Commande prise", commercial="T2")

    cm, ws = specs["Commandes"], wb["Commandes"]
    oeufs, poulet, decoupes = "Œufs frais - alvéole de 30", "Poulet de chair entier", "Découpes de volaille"
    commandes = [
        dict(date=D(2026, 10, 5), client=L1, type="Commande", produit=oeufs, qte=20, pu=2500, livreelle=D(2026, 10, 6)),
        dict(date=D(2026, 10, 5), client=L1, type="Commande", produit=poulet, qte=5, pu=4000, livreelle=D(2026, 10, 6)),
        dict(date=D(2026, 10, 8), client=L2, type="Précommande", produit=poulet, qte=100, pu=3800, livprevue=D(2026, 10, 20)),
        dict(date=D(2026, 10, 7), client=L4, type="Vente directe", produit=oeufs, qte=10, pu=2600),
        dict(date=D(2026, 10, 7), client=L3, type="Commande", produit=decoupes, qte=2, pu=3000, annulee=OUI),
    ]
    for i, c in enumerate(commandes):
        _poser(ws, cm, 6 + i, **c)
    _poser(ws, cm, 20, date=D(2026, 10, 6), client=LT2, type="Vente directe", produit=oeufs, qte=4, pu=2500, commercial="T2")

    en, ws = specs["Encaissements"], wb["Encaissements"]
    encaissements = [
        dict(date=D(2026, 10, 6), client=L1, montant=30000, mode="Espèces"),
        dict(date=D(2026, 10, 7), client=L4, montant=26000, mode="Orange Money", reference="OM123456"),
        dict(date=D(2026, 10, 7), client=L1, montant=20000, mode="Espèces", dateverse=D(2026, 10, 7)),
    ]
    for i, e in enumerate(encaissements):
        _poser(ws, en, 6 + i, **e)

    st, ws = specs["Stock"], wb["Stock"]
    stock = [
        dict(date=D(2026, 10, 5), mouvement="Reçu du magasin", produit=oeufs, qte=50),
        dict(date=D(2026, 10, 5), mouvement="Reçu du magasin", produit=poulet, qte=20),
        dict(date=D(2026, 10, 6), mouvement="Casse / avarie", produit=oeufs, qte=2, cause="Œufs fêlés ou cassés au transport", photo=OUI, valeursaisie=5000),
        dict(date=D(2026, 10, 7), mouvement="Retourné au magasin", produit=oeufs, qte=3),
    ]
    for i, s in enumerate(stock):
        _poser(ws, st, 6 + i, **s)

    jr, ws = specs["Journal"], wb["Journal"]
    for i, (jour, debut, fin, zone, frais) in enumerate(
        [(D(2026, 10, 5), dt.time(8, 0), dt.time(17, 0), "Akwa", 3000), (D(2026, 10, 6), dt.time(8, 15), dt.time(16, 30), "Deido", 2500), (D(2026, 10, 7), dt.time(8, 0), dt.time(17, 30), "New Bell", 2000)]
    ):
        _poser(ws, jr, 6 + i, date=jour, debut=debut, fin=fin, zone=zone, frais=frais, forts="Bonne journée", difficultes="Circulation dense", soutien="Stock d'œufs", demain="Visiter l'hôtel")

    ob, ws = specs["Objectifs"], wb["Objectifs"]
    objectifs = [
        dict(metrique="Visites sur place", periode="Semaine", debut=D(2026, 10, 5), valeur=20),
        dict(metrique="Chiffre d'affaires (XAF)", periode="Semaine", debut=D(2026, 10, 5), valeur=100000),
        dict(metrique="Nouveaux clients", periode="Semaine", debut=D(2026, 10, 5), valeur=1),
        dict(metrique="Chiffre d'affaires (XAF)", periode="Mois", debut=D(2026, 10, 1), valeur=400000),
    ]
    for i, o in enumerate(objectifs):
        _poser(ws, ob, 6 + i, **o)

    pl, ws = specs["Planning"], wb["Planning"]
    _poser(ws, pl, 6, mois=D(2026, 10, 1), semmois=2, jour=D(2026, 10, 5), client=L1, objet="Prise de commande")
    _poser(ws, pl, 7, mois=D(2026, 10, 1), semmois=2, jour=D(2026, 10, 6), client=L2, objet="Prospection / découverte")
    _poser(ws, pl, 8, mois=D(2026, 10, 1), semmois=2, client=L3, objet="Présentation produits et prix")

    bi, ws = specs["Bilans_Hebdo"], wb["Bilans_Hebdo"]
    _poser(ws, bi, 6, lundi=D(2026, 10, 5), reussites="2 nouveaux clients", difficultes="Rupture de calibre", besoins="Plus de stock", priorites="Relancer l'hôtel", marche="Concurrent à 2300 XAF")


# --- Guide (LISEZ-MOI) --------------------------------------------------------------------------------------------------------------------


def _hauteur(texte: str, largeur_car: float, ligne: float = 14.5, mini: float = 18) -> float:
    lignes = sum(max(1, math.ceil(len(partie) / largeur_car)) for partie in texte.split("\n"))
    return max(mini, lignes * ligne + 6)


def feuille_lisez_moi(ws, demarrage: bool) -> None:
    ws.sheet_properties.tabColor = COULEURS_ONGLETS["guide"]
    largeurs = (3, 26, 28, 22, 58, 14, 3)
    for col, largeur in zip("ABCDEFG", largeurs):
        ws.column_dimensions[col].width = largeur
    ws.sheet_view.showGridLines = False
    ws["B1"] = "FICHE DE REPORTING COMMERCIAL — GIC AGROPELC"
    ws["B1"].font = police(20, True, couleur=VERT_FONCE)
    ws.row_dimensions[1].height = 34
    ws["B2"] = f"Version {VERSION} · {DATE_VERSION} · Votre cahier de terrain, vos rapports et votre tableau de bord dans un seul fichier."
    ws["B2"].font = police(10, False, True, ARDOISE)
    ligne = 4
    LARGEUR_TOTALE = 138  # B:F, en caractères

    def titre(texte: str) -> None:
        nonlocal ligne
        section(ws, ligne, texte, 2, 6)
        ligne += 1

    def paragraphe(texte: str, gras: bool = False, couleur: str = "000000", taille: int = 10) -> None:
        nonlocal ligne
        ws.merge_cells(start_row=ligne, start_column=2, end_row=ligne, end_column=6)
        c = ws.cell(ligne, 2, texte)
        c.font = police(taille, gras, False, couleur)
        c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
        ws.row_dimensions[ligne].height = _hauteur(texte, LARGEUR_TOTALE - 6)
        ligne += 1

    # Colonnes de la fiche : B=26, C=28, D=22, E=58, F=14. Les tableaux de 2 ou 3 colonnes fusionnent des colonnes voisines.
    SPANS = {
        2: [(2, 2), (3, 6)],
        3: [(2, 3), (4, 5), (6, 6)],
        4: [(2, 2), (3, 3), (4, 4), (5, 5)],
        5: [(2, 2), (3, 3), (4, 4), (5, 5), (6, 6)],
    }

    def _largeur(debut_col: int, fin_col: int) -> float:
        return sum(largeurs[debut_col - 1:fin_col])

    def tableau(entetes: list, lignes_: list) -> None:
        nonlocal ligne
        spans = SPANS[len(entetes)]
        for (a, b), texte in zip(spans, entetes):
            for col in range(a, b + 1):
                c = ws.cell(ligne, col)
                c.fill = fill(ARDOISE)
                c.border = BORDURE
            ws.cell(ligne, a, texte)
            ws.cell(ligne, a).font = police(10, True, couleur="FFFFFF")
            ws.cell(ligne, a).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if b > a:
                ws.merge_cells(start_row=ligne, start_column=a, end_row=ligne, end_column=b)
        ws.row_dimensions[ligne].height = 24
        ligne += 1
        for valeurs in lignes_:
            haut = 18
            for i, ((a, b), texte) in enumerate(zip(spans, valeurs)):
                for col in range(a, b + 1):
                    ws.cell(ligne, col).border = BORDURE
                cell = ws.cell(ligne, a, texte)
                cell.font = police(10, i == 0)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                if b > a:
                    ws.merge_cells(start_row=ligne, start_column=a, end_row=ligne, end_column=b)
                haut = max(haut, _hauteur(str(texte), _largeur(a, b) - 2))
            ws.row_dimensions[ligne].height = haut
            ligne += 1

    titre("1. L'ESSENTIEL EN 30 SECONDES")
    paragraphe("Cette fiche remplace votre cahier : vous y notez ce que vous faites sur le terrain, et elle calcule toute seule vos rapports (jour, semaine, mois) et votre tableau de bord.")
    paragraphe("Vous ne tapez presque rien. Presque chaque case s'ouvre en liste : cliquez la case, puis la petite flèche à droite, et choisissez. Pour une DATE, la liste propose les jours (aujourd'hui en premier). Pour une HEURE, la liste propose chaque quart d'heure. Pour un CLIENT ou un PRODUIT, tapez les premières lettres (Excel 365 filtre la liste) ou ouvrez la liste.")
    paragraphe("Les « cases à cocher » sont des listes à un clic : ☑ Oui / ☐ Non.")
    paragraphe("Couleurs : case BLANCHE = à remplir · case GRISE = se calcule seule, n'y touchez pas · case JAUNE = votre choix dans un rapport (date, commercial) · case ROUGE = une information obligatoire manque.", gras=True)
    ligne += 1

    titre("2. CE QUE JE REMPLIS, ET QUAND")
    tableau(
        ["Quand", "Je remplis", "Dans la feuille", "Comment (en clair)", "Durée"],
        [
            ("UNE SEULE FOIS, le premier jour", "Mon identité", "Paramètres", "Nom, CODE COMMERCIAL (ex. COM01), fonction, zone. Ne changez plus le code ensuite.", "3 min"),
            ("UNE SEULE FOIS", "Mon portefeuille", "Clients", "Une ligne par client ou prospect : nom, type, téléphone, quartier, source, date de création. Ces clients apparaissent ensuite dans toutes les listes.", "10 à 30 min"),
            ("UNE SEULE FOIS (direction)", "Les prix de référence", "Listes", "Colonne « Prix de référence » des produits. Les commandes et les casses les utilisent toutes seules.", "5 min"),
            ("CHAQUE LUNDI (ou en fin de mois)", "Mes objectifs", "Objectifs", "Choisir l'indicateur, « Semaine » ou « Mois », le début dans la liste, la valeur à atteindre.", "3 min"),
            ("CHAQUE MATIN", "Ma journée", "Journal", "Date du jour (premier de la liste), heure de prise de service, zone du jour. Regardez « Rapport_Jour » : actions et visites prévues.", "1 min"),
            ("APRÈS CHAQUE VISITE OU APPEL", "Le contact", "Visites_Contacts", "Date, heure, type de contact, client, objet, résultat ; si une action est à faire : sa date et ce qu'il faut faire.", "1 min"),
            ("QUAND JE TROUVE UN NOUVEAU PROSPECT", "Le prospect", "Clients", "Nom, type, téléphone, quartier, source, étape, date de création. À faire AVANT d'enregistrer la visite.", "1 min"),
            ("À CHAQUE COMMANDE OU VENTE", "Les produits", "Commandes", "Date, client, type, produit, quantité (une ligne par produit). Livraison réelle quand c'est livré. Le prix et le montant se calculent.", "1 min"),
            ("À CHAQUE PAIEMENT REÇU", "L'argent", "Encaissements", "Date, client, montant, mode de paiement (+ référence si Orange Money, MoMo, virement ou chèque).", "30 s"),
            ("QUAND JE VERSE L'ARGENT À L'ENTREPRISE", "La date du versement", "Encaissements", "Sur les lignes d'espèces concernées : choisir la date du versement. « Versé ? » passe à ☑ Oui tout seul.", "30 s"),
            ("QUAND JE REÇOIS, RENDS OU CASSE DE LA MARCHANDISE", "Le mouvement", "Stock", "Date, mouvement, produit, quantité ; pour une casse : la cause et « Photo prise ? ».", "1 min"),
            ("CHAQUE SOIR", "Mon bilan du jour", "Journal", "Heure de fin, frais, réussites, difficultés, besoin de soutien, plan du lendemain.", "3 min"),
            ("CHAQUE SOIR", "Mon rapport du jour", "Rapport_Jour", "Choisir la date (case jaune) puis imprimer ou exporter en PDF et envoyer au responsable.", "1 min"),
            ("CHAQUE VENDREDI OU SAMEDI", "Mon bilan de la semaine", "Bilans_Hebdo", "Choisir le lundi dans la liste puis écrire les 5 rubriques. Vérifiez « Rapport_Semaine », préparez le « Planning » de la semaine suivante.", "10 min"),
            ("CHAQUE DÉBUT DE MOIS", "Mon plan et mon bilan du mois", "Plan_Mensuel, Tableau_de_bord", "Choisir le mois ; lire les résultats ; regarder l'évolution sur 12 semaines dans le tableau de bord.", "5 min"),
            ("RESPONSABLE COMMERCIAL, CHAQUE SEMAINE", "Le suivi de l'équipe", "Équipe, Tableau_de_bord, Bilans_Hebdo", "Choisir le lundi ; comparer les commerciaux ; écrire le commentaire du responsable dans Bilans_Hebdo.", "15 min"),
        ],
    )
    ligne += 1

    titre("3. REMPLIR UNE LIGNE, PAS À PAS")
    paragraphe("EXEMPLE 1 — Je viens de visiter « Resto La Paix » (feuille Visites_Contacts, première ligne vide)", gras=True, couleur=VERT_FONCE)
    paragraphe("1) Date : cliquez la case, puis la flèche : choisissez aujourd'hui (premier de la liste).   2) Heure : choisissez dans la liste.   3) Type de contact : « Visite sur place ».   4) Client : tapez « Resto » puis choisissez « Resto La Paix [001] ».   5) Objet et Résultat : choisissez dans les listes.   6) S'il y a quelque chose à faire : date de la prochaine action (liste) et ce qu'il faut faire (ex. « Rappeler jeudi »).   C'est tout : le nom, le quartier et la semaine s'affichent seuls, et la fiche du client se met à jour.")
    paragraphe("EXEMPLE 2 — Le client commande 20 alvéoles et 5 poulets (feuille Commandes, deux lignes)", gras=True, couleur=VERT_FONCE)
    paragraphe("Ligne 1 : date, client, type « Commande », produit « Œufs frais - alvéole de 30 », quantité 20. Ligne 2 : mêmes date, client et type, produit « Poulet de chair entier », quantité 5. Le numéro de commande, le prix (celui de la feuille Listes), le montant et le statut se remplissent seuls. Quand c'est livré : choisissez la « Livraison réelle ». À la première livraison, le prospect devient « Client » tout seul.")
    paragraphe("EXEMPLE 3 — Le client paie 30 000 XAF en espèces (feuille Encaissements)", gras=True, couleur=VERT_FONCE)
    paragraphe("Date, client, montant, mode « Espèces ». La colonne « Reste dû par le client » vous montre ce qu'il doit encore. Quand vous remettez l'argent à l'entreprise, choisissez la « Date du versement » sur cette ligne.")
    ligne += 1

    titre("4. CE QUI SE REMPLIT TOUT SEUL (les liens entre les feuilles)")
    tableau(
        ["Je saisis…", "…et cela apparaît automatiquement dans"],
        [
            ("un client (Clients)", "les listes de Visites_Contacts, Commandes, Encaissements et Planning ; la feuille Fiche_Client ; son code et son statut"),
            ("une visite ou un appel", "Clients (dernier contact, prochaine action, nombre de visites) ; Journal (compteurs du jour) ; Rapport_Jour ; Rapport_Semaine ; Tableau_de_bord"),
            ("une commande livrée", "Clients (devient « Client », CA, dernier achat) ; Journal (CA du jour) ; Stock (vendu et restant) ; tous les rapports"),
            ("un paiement", "Clients (encaissé et reste à encaisser) ; Journal ; « espèces détenues » des rapports ; taux de recouvrement du tableau de bord"),
            ("une casse ou une perte", "Stock (restant) ; Journal (casses) ; rapports ; tableau de bord"),
            ("un objectif", "Rapport_Semaine, Plan_Mensuel, Équipe et Tableau_de_bord (taux d'atteinte coloré)"),
            ("un prix de référence (Listes)", "Commandes (prix appliqué, montant) et Stock (valeur des casses)"),
            ("mon code commercial (Paramètres)", "la colonne « Commercial » de toutes les feuilles et les références des lignes"),
        ],
    )
    ligne += 1

    titre("5. LES 4 RÈGLES D'OR")
    paragraphe("1.  Une ligne = un fait (une visite, un produit commandé, un paiement…). Ne regroupez pas, ne fusionnez pas de cellules dans les feuilles de saisie.")
    paragraphe("2.  Choisissez dans les listes plutôt que de taper : c'est ce qui rend les totaux justes. N'écrivez jamais dans une case grise.")
    paragraphe("3.  Remplissez sans laisser de ligne vide au milieu. Ne supprimez ni n'insérez de colonne ; ne renommez pas les feuilles.")
    paragraphe("4.  Une case ROUGE = il manque une information obligatoire, ou une date mal écrite : corrigez-la, les rapports en dépendent.")
    ligne += 1

    titre("6. LES FEUILLES")
    tableau(
        ["Feuille", "Quand", "Qui", "À quoi elle sert"],
        [(a, c, d, b) for a, b, c, d in [
            ("Paramètres", "Identité, seuils d'alerte, équipe", "Une fois", "Tous"),
            ("Journal", "Une ligne par jour : heures, zone, frais, bilan (+ chiffres du jour calculés)", "Chaque jour", "Commercial"),
            ("Visites_Contacts", "Chaque visite, appel ou message, avec résultat et prochaine action", "Chaque contact", "Commercial"),
            ("Commandes", "Précommandes, commandes, ventes directes (une ligne par produit)", "Chaque commande", "Commercial"),
            ("Encaissements", "Paiements reçus et versement à l'entreprise", "Chaque paiement", "Commercial"),
            ("Clients", "Portefeuille, pipeline, alertes, CA par client", "Chaque nouveau contact", "Commercial"),
            ("Planning", "Visites à faire, semaine par semaine", "Chaque semaine", "Commercial"),
            ("Stock", "Marchandise reçue, rendue, cassée ou perdue", "Chaque mouvement", "Commercial"),
            ("Objectifs", "Cibles de la semaine ou du mois", "Chaque semaine / mois", "Commercial, responsable"),
            ("Bilans_Hebdo", "Le texte du rapport de la semaine (+ chiffres de la semaine)", "Chaque semaine", "Commercial, responsable"),
            ("Fiche_Client", "Rechercher un client et voir toute son histoire", "Au besoin", "Tous"),
            ("Rapport_Jour", "Rapport journalier imprimable", "Chaque soir", "Tous"),
            ("Rapport_Semaine", "Rapport hebdomadaire avec graphiques", "Chaque semaine", "Tous"),
            ("Plan_Mensuel", "Plan et résultats du mois, semaine par semaine", "Chaque mois", "Tous"),
            ("Équipe", "Comparaison des commerciaux", "Chaque semaine", "Responsable"),
            ("Listes", "Choix des listes déroulantes, prix de référence (modifiable)", "Au besoin", "Responsable"),
            ("Correspondances", "Où ira chaque colonne dans l'application (information)", "—", "Équipe projet"),
            ("Tableau_de_bord", "Indicateurs clés et leur évolution sur 12 semaines (dernière page)", "Chaque semaine", "Tous"),
        ]],
    )
    # liens vers les feuilles
    premiere_feuille = ligne - 18
    for i in range(18):
        cell = ws.cell(premiere_feuille + i, 2)
        cell.hyperlink = f"#'{cell.value}'!A1"
        cell.font = Font(name="Calibri", size=10, bold=True, color=BLEU, underline="single")
    ligne += 1

    titre("7. VOCABULAIRE")
    tableau(
        ["Mot", "Ce que cela veut dire"],
        [
            ("Prospect / Client / Perdu", "Un prospect n'a pas encore acheté ; il devient Client automatiquement à sa première livraison ; marquez « Compte perdu ? » s'il renonce (avec le motif)."),
            ("Précommande", "Commande prise avant que la marchandise soit disponible (ex. premiers poulets). Elle n'entre dans le chiffre d'affaires qu'à la livraison."),
            ("Vente directe", "Marchandise livrée et payée sur place, sans commande préalable : elle est livrée à la date de la vente."),
            ("Chiffre d'affaires réalisé", "Valeur des lignes livrées, comptée à la date de livraison."),
            ("Commande prévisionnelle par jour", "Quantité qu'un point de vente pourrait commander chaque jour, estimée à la prospection (Clients, « Potentiel estimé par jour »)."),
            ("Panier moyen", "Chiffre d'affaires divisé par le nombre de ventes (un client servi un jour = une vente)."),
            ("Taux de conversion", "Part des visites sur place qui se terminent par une commande prise."),
            ("Espèces détenues", "Espèces encaissées et pas encore versées à l'entreprise."),
        ],
    )
    ligne += 1

    if demarrage:
        titre("8. CE QUI A ÉTÉ REPRIS DE LA FICHE DE SEPTEMBRE 2026")
        paragraphe("Les chiffres de votre fiche « ÉTAT DE DÉPLOIEMENT SEPTEMBRE » sont déjà dans ce classeur : vous n'avez pas à les ressaisir, seulement à les compléter au fil des visites. Le petit tableau « RECAP » des SIM n'a pas été repris : il se trouvait là par erreur (modèle d'une autre activité).", couleur=ARDOISE)
        tableau(
            ["Dans votre fiche", "Dans ce classeur", "Où"],
            [
                ("26 points de vente des semaines 1, 2 et 4, avec leur numéro GSM", "Un prospect par ligne : nom, type, téléphone, localisation, source, statut", "Clients"),
                ("Commande prévisionnelle par jour de chaque point de vente", "Colonne « Potentiel estimé par jour »", "Clients"),
                ("Type de point de vente (RESTO, PARTICULIER, REVENDEUR)", "Type de client (Restaurant, Particulier / ménage, Revendeur)", "Clients"),
                ("Découpage en semaines 1, 2 et 4", "Une ligne de planning par point de vente avec son numéro de semaine (septembre 2026)", "Planning"),
                ("TOTAL COMMANDE PRÉVISIONNELLE JOUR", "Calculé : semaine 1 = 51, semaine 2 = 39, semaine 4 = 45, soit 135 par jour", "Plan_Mensuel"),
            ],
        )
        paragraphe("Valeurs de départ retenues (modifiables directement) : noms construits d'après le type et la localisation ; type déduit du libellé pour les semaines 1 et 2 quand il est explicite (hôtel, boulangerie, marché), sinon « Autre » ; étape « Nouveau », source « Prospection terrain », date de création = premier jour de la semaine de ciblage (1er, 7 et 21 septembre 2026) ; code commercial « COM01 » prérempli dans Paramètres, à remplacer par le vôtre.")
        ligne += 1

    paragraphe("Besoin d'aide ou d'une évolution de la fiche ? Contactez la direction commerciale. Cette fiche sera chargée dans l'application GIC AGROPELC à son ouverture ; gardez-la à jour, c'est ce qui permettra de reprendre votre historique.", couleur=ARDOISE)
    ws.freeze_panes = "A4"
    mise_en_page(ws, "portrait", f"A1:G{ligne}")


# --- Correspondances ----------------------------------------------------------------------------------------------------------------------


def feuille_correspondances(wb, specs: dict) -> None:
    ws = wb.create_sheet("Correspondances")
    ws.sheet_properties.tabColor = COULEURS_ONGLETS["reference"]
    for col, largeur in zip("ABCDEF", (22, 38, 12, 50, 14, 52)):
        ws.column_dimensions[col].width = largeur
    ws["A1"] = "CORRESPONDANCES — où ira chaque colonne dans l'application"
    ws["A1"].font = police(16, True, couleur=VERT_FONCE)
    ws["A2"] = (
        "Information pour l'équipe projet (INDICATIF, à valider au moment de l'import). Toutes les feuilles de saisie ont leur entête en ligne 5 et leurs données "
        "à partir de la ligne 6 ; les colonnes « Technique » (masquées) ne sont pas à importer."
    )
    ws["A2"].font = police(10, False, True, ARDOISE)
    ws["A3"] = "◄ Sommaire"
    ws["A3"].hyperlink = "#'LISEZ-MOI'!A1"
    ws["A3"].font = Font(name="Calibri", size=10, bold=True, color=BLEU, underline="single")
    entete_tableau(ws, 5, ["Feuille", "Colonne", "Nature", "Champ cible indicatif dans l'application", "Obligatoire", "Aide à la saisie"], hauteur=26)
    r = 6
    ordre = ["Journal", "Visites_Contacts", "Commandes", "Encaissements", "Clients", "Planning", "Stock", "Objectifs", "Bilans_Hebdo"]
    for nom in ordre:
        spec = specs[nom]
        for c in spec.cols:
            nature = {"in": "Saisie", "calc": "Calculé", "static": "Numéro"}[c.kind]
            if c.cache:
                nature = "Technique"
            valeurs = [nom, c.header, nature, c.tgt or "—", "Oui" if (c.req or c.req_si) else "", c.hint]
            for i, v in enumerate(valeurs):
                cell = ws.cell(r, 1 + i, v)
                cell.font = police(9, i == 0)
                cell.border = BORDURE
                cell.alignment = Alignment(vertical="center", wrap_text=(i == 3))
            r += 1
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:F{r - 1}"
    mise_en_page(ws, "landscape")


# --- Assemblage ---------------------------------------------------------------------------------------------------------------------------


def construire(demarrage: bool, exemple: bool = False) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    ws_lisez = wb.create_sheet("LISEZ-MOI")
    feuille_parametres(wb)

    specs = {}
    for fabrique in (
        J.journal_journal,
        J.journal_visites,
        J.journal_commandes,
        J.journal_encaissements,
        J.journal_clients,
        J.journal_planning,
        J.journal_stock,
        J.journal_objectifs,
        J.journal_bilans,
    ):
        spec = fabrique()
        construire_journal(wb, spec)
        specs[spec.nom] = spec

    # Liste dynamique des clients (libellé « Resto La Paix [012] »)
    cl = specs["Clients"]
    definir_nom(
        wb,
        "L_Clients",
        f"OFFSET('Clients'!${cl.lettres['label']}$6,0,0,MAX(1,COUNTIF('Clients'!${cl.lettres['nom']}$6:${cl.lettres['nom']}${cl.derniere},\"?*\")),1)",
    )

    fiche_client(wb)
    R.rapport_jour(wb)
    R.rapport_semaine(wb)
    R.plan_mensuel(wb)
    R.equipe(wb)
    feuille_listes(wb)
    feuille_correspondances(wb, specs)
    tableau_de_bord(wb)
    feuille_lisez_moi(ws_lisez, demarrage)

    if demarrage:
        garnir_demarrage(wb, specs)
    if exemple:
        garnir_exemple(wb, specs)

    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.properties.title = "Fiche de reporting commercial — GIC AGROPELC"
    wb.properties.creator = "GIC AGROPELC — Direction commerciale"
    wb.properties.description = f"Fiche de reporting commercial v{VERSION} ({DATE_VERSION})"
    for feuille in wb.worksheets:
        feuille.sheet_view.tabSelected = False
    wb.active = 0
    wb["LISEZ-MOI"].sheet_view.tabSelected = True
    return wb


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--exemple", help="écrit un fichier d'essai avec des données fictives (vérification) et s'arrête")
    args = parser.parse_args()
    if args.exemple:
        construire(demarrage=False, exemple=True).save(args.exemple)
        print(f"Fichier d'essai écrit : {args.exemple}")
        return
    dossier = Path(__file__).resolve().parent.parent
    cible_vierge = dossier / f"GIC_Reporting_Commercial_v{VERSION}_MODELE_VIERGE.xlsx"
    cible_demarrage = dossier / f"GIC_Reporting_Commercial_v{VERSION}_DEMARRAGE_Sept2026.xlsx"
    construire(demarrage=False).save(cible_vierge)
    construire(demarrage=True).save(cible_demarrage)
    print(f"Écrit : {cible_vierge.name}\nÉcrit : {cible_demarrage.name}")


if __name__ == "__main__":
    main()
