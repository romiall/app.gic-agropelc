"""Outils de construction des classeurs Excel de fiches terrain (openpyxl).

Principes (communs à toutes les fiches du dossier `fiches-terrain/`) :

- les feuilles de **saisie** sont des journaux : une ligne = un fait, entête en ligne 5, données à partir
  de la ligne 6, aucune cellule fusionnée dans le tableau, un type de valeur par colonne ;
- les feuilles de **rapport** ne contiennent que des formules (les indicateurs se calculent à partir des
  faits saisis, jamais à la main) ; seules les cellules de sélection (date, commercial) sont libres ;
- les listes déroulantes viennent de la feuille « Listes » (noms définis dynamiques) ;
- les colonnes calculées sont grisées et ne se saisissent pas.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Callable, Optional

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

# --- Palette ---------------------------------------------------------------------------------------------
VERT_FONCE = "1F4D3A"
VERT = "2E7D5B"
VERT_CLAIR = "E8F3EC"
GRIS_CALC = "F1F3F4"
GRIS_BORD = "C9CED2"
ARDOISE = "5F6B73"
JAUNE = "FFF4CC"
ROUGE_CLAIR = "FFC7CE"
ROUGE_TEXTE = "9C0006"
ORANGE = "E07B39"
BLEU = "2F6EBA"
SARCELLE = "1F7A8C"
VIOLET = "7A4FA3"
GRIS_ONGLET = "7F8C8D"

POLICE = "Calibri"

COULEURS_ONGLETS = {
    "guide": VERT_FONCE,
    "quotidien": BLEU,
    "portefeuille": SARCELLE,
    "ponctuel": ORANGE,
    "rapport": VIOLET,
    "reference": GRIS_ONGLET,
}

FIN = Side(style="thin", color=GRIS_BORD)
BORDURE = Border(left=FIN, right=FIN, top=FIN, bottom=FIN)

FMT_DATE = "dd/mm/yyyy"
FMT_DATE_ZERO_CACHE = "dd/mm/yyyy;;"
FMT_HEURE = "hh:mm"
FMT_XAF = "#,##0"
FMT_XAF_ZERO_CACHE = "#,##0;-#,##0;"
FMT_POURCENT = "0%"
FMT_TEXTE = "@"
FMT_ENTIER = "0"
FMT_GPS = "0.000000"


def fill(couleur: str) -> PatternFill:
    return PatternFill(start_color=couleur, end_color=couleur, fill_type="solid")


def police(taille: int = 10, gras: bool = False, italique: bool = False, couleur: str = "000000") -> Font:
    return Font(name=POLICE, size=taille, bold=gras, italic=italique, color=couleur)


def definir_nom(wb, nom: str, reference: str) -> None:
    """Nom défini au niveau du classeur (`reference` : `'Feuille'!$A$1:$A$9` ou une formule)."""
    wb.defined_names[nom] = DefinedName(nom, attr_text=reference)


def ref_plage(feuille: str, col: str, premiere: int, derniere: int) -> str:
    return f"'{feuille}'!${col}${premiere}:${col}${derniere}"


# --- Colonnes d'un journal -------------------------------------------------------------------------------------


@dataclass
class Col:
    """Une colonne d'un journal de saisie.

    kind : « in » (saisie), « calc » (formule), « static » (valeur fixe : numéro de ligne).
    f    : pour une colonne « calc », fonction (ligne, lettres) -> formule.
    dv   : validation de données (voir `_appliquer_validation`).
    tgt  : champ cible indicatif dans l'application (feuille « Correspondances »).
    req  : champ obligatoire dès que la ligne est commencée.
    req_si : obligatoire sous condition (fonction (lettres) -> formule sur la ligne 6).
    cache : colonne technique masquée.
    """

    key: str
    header: str
    width: float = 14
    kind: str = "in"
    fmt: Optional[str] = None
    hint: str = ""
    req: bool = False
    req_si: Optional[Callable[[dict], str]] = None
    dv: Optional[dict] = None
    f: Optional[Callable[[int, dict], str]] = None
    tgt: str = ""
    cache: bool = False
    align: Optional[str] = None
    date: bool = False  # colonne de dates : une valeur qui n'est pas une date est signalée en rouge
    nommer: bool = True  # enregistrer un nom défini `{prefixe}_{key}` pour la plage de la colonne


@dataclass
class Journal:
    nom: str
    titre: str
    sous_titre: str
    prefixe: str
    cols: list
    lignes: int
    onglet: str
    gel: str = "D6"
    cle: str = ""  # colonne dont la saisie « ouvre » une ligne (sinon la première colonne de saisie)
    entite_cible: str = ""
    # renseignés à la construction
    lettres: dict = field(default_factory=dict)
    premiere: int = 6
    derniere: int = 0


ENTETE = 5
PREMIERE = 6


def _groupes_contigus(lettres_in: list) -> list:
    """['C','D','E','G'] -> [('C','E'), ('G','G')] pour des plages COUNTA compactes."""
    if not lettres_in:
        return []
    from openpyxl.utils import column_index_from_string as idx

    ordre = sorted(lettres_in, key=idx)
    groupes = []
    debut = prec = ordre[0]
    for lettre in ordre[1:]:
        if idx(lettre) == idx(prec) + 1:
            prec = lettre
        else:
            groupes.append((debut, prec))
            debut = prec = lettre
    groupes.append((debut, prec))
    return groupes


# Listes proposées sous forme de sélecteurs (calendrier, heures, quantités…) : une valeur hors liste reste
# permise après un simple « OK » (style « information »), pour ne jamais bloquer le terrain.
LISTES_SELECTEUR = {
    "date_passee": ("L_DatesPassees", "Choisissez la date dans la liste (aujourd'hui en tête). Vous pouvez aussi en taper une autre, au format jj/mm/aaaa."),
    "date_avenir": ("L_DatesAvenir", "Choisissez la date dans la liste (aujourd'hui en tête). Vous pouvez aussi en taper une autre, au format jj/mm/aaaa."),
    "lundi_passe": ("L_LundisPasses", "Choisissez le lundi de la semaine dans la liste."),
    "lundi_avenir": ("L_LundisAvenir", "Choisissez le lundi de la semaine dans la liste."),
    "mois": ("L_Mois", "Choisissez le mois dans la liste."),
    "heure": ("L_Heures", "Choisissez l'heure dans la liste (tous les quarts d'heure). Vous pouvez aussi taper une autre heure, au format hh:mm."),
    "quantite": ("L_Quantites", "Choisissez une quantité dans la liste ou tapez-en une autre."),
    "frais": ("L_Frais", "Choisissez un montant dans la liste ou tapez-en un autre."),
    "remise": ("L_Remises", "Choisissez un montant dans la liste ou tapez-en un autre."),
}


def _liste(formule: str, style: str, titre: str, message: str) -> DataValidation:
    return DataValidation(
        type="list",
        formula1=formule,
        allow_blank=True,
        showErrorMessage=True,
        errorStyle=style,
        errorTitle=titre,
        error=message,
    )


def _appliquer_validation(ws, col_letter: str, premiere: int, derniere: int, dv: dict) -> None:
    """dv : {'type': 'liste', 'nom': 'L_Produit', 'souple': bool} | {'type': 'liste_formule', 'formule': '=…'}
    | un type de sélecteur de LISTES_SELECTEUR | {'type': 'nombre', 'min': 0} | {'type': 'positif'}
    | {'type': 'entier', 'min': 1, 'max': 5} | {'type': 'perso', 'formule': '...', 'message': '...'}"""
    plage = f"{col_letter}{premiere}:{col_letter}{derniere}"
    t = dv["type"]
    if t == "liste":
        souple = dv.get("souple", False)
        v = _liste(
            f"={dv['nom']}",
            "warning" if souple else "stop",
            "Valeur hors liste",
            "Cette valeur n'est pas dans la liste. Choisissez-la avec la flèche à droite de la cellule"
            + (" (vous pouvez aussi garder une autre valeur)." if souple else "."),
        )
    elif t == "liste_formule":
        v = _liste(dv["formule"], "information", "Valeur hors liste", "Choisissez dans la liste avec la flèche à droite de la cellule.")
    elif t in LISTES_SELECTEUR:
        nom, message = LISTES_SELECTEUR[t]
        v = _liste(f"={nom}", "information", "Valeur hors liste", message + " Cliquez sur OK pour la garder telle quelle.")
    elif t == "date":
        v = DataValidation(
            type="date",
            operator="between",
            formula1="=DATE(2024,1,1)",
            formula2="=DATE(2040,12,31)",
            allow_blank=True,
            showErrorMessage=True,
            errorStyle="stop",
            errorTitle="Date invalide",
            error="Saisissez une date valide au format jj/mm/aaaa (ex. 08/10/2026).",
        )
    elif t == "nombre":
        v = DataValidation(
            type="decimal",
            operator="greaterThanOrEqual",
            formula1=str(dv.get("min", 0)),
            allow_blank=True,
            showErrorMessage=True,
            errorStyle="stop",
            errorTitle="Nombre invalide",
            error=f"Saisissez un nombre supérieur ou égal à {dv.get('min', 0)}.",
        )
    elif t == "positif":
        v = DataValidation(
            type="decimal",
            operator="greaterThan",
            formula1="0",
            allow_blank=True,
            showErrorMessage=True,
            errorStyle="stop",
            errorTitle="Nombre invalide",
            error="Saisissez un nombre strictement positif.",
        )
    elif t == "entier":
        v = DataValidation(
            type="whole",
            operator="between",
            formula1=str(dv.get("min", 0)),
            formula2=str(dv.get("max", 9999999)),
            allow_blank=True,
            showErrorMessage=True,
            errorStyle="stop",
            errorTitle="Nombre entier attendu",
            error=f"Saisissez un nombre entier entre {dv.get('min', 0)} et {dv.get('max', 9999999)}.",
        )
    elif t == "perso":
        v = DataValidation(
            type="custom",
            formula1=dv["formule"],
            allow_blank=True,
            showErrorMessage=True,
            errorStyle="warning",
            errorTitle="Vérifiez la valeur",
            error=dv.get("message", "La valeur saisie semble incorrecte."),
        )
    else:  # pragma: no cover
        raise ValueError(t)
    ws.add_data_validation(v)
    v.add(plage)


def _en_tete_feuille(ws, titre: str, sous_titre: str, largeur_cols: int, retour: bool = True) -> None:
    ws["A1"] = titre
    ws["A1"].font = police(16, True, couleur=VERT_FONCE)
    ws["A2"] = sous_titre
    ws["A2"].font = police(10, False, True, ARDOISE)
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 18


def _lien_sommaire(ws, cellule: str = "A3") -> None:
    c = ws[cellule]
    c.value = "◄ Sommaire"
    c.hyperlink = "#'LISEZ-MOI'!A1"
    c.font = Font(name=POLICE, size=10, bold=True, color=BLEU, underline="single")


def _legende(ws, ligne: int, depart_col: int) -> None:
    items = [("Saisie", None, "000000"), ("Calculé", GRIS_CALC, ARDOISE), ("À compléter", ROUGE_CLAIR, ROUGE_TEXTE)]
    for i, (texte, couleur, ft) in enumerate(items):
        c = ws.cell(ligne, depart_col + i, texte)
        c.font = police(9, False, False, ft)
        c.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
        c.border = BORDURE
        if couleur:
            c.fill = fill(couleur)
    ws.row_dimensions[ligne].height = 22


def construire_journal(wb, spec: Journal):
    """Crée la feuille de saisie décrite par `spec` ; renvoie la feuille."""
    ws = wb.create_sheet(spec.nom)
    ws.sheet_properties.tabColor = COULEURS_ONGLETS[spec.onglet]
    cols = spec.cols
    spec.lettres = {c.key: get_column_letter(i + 1) for i, c in enumerate(cols)}
    L = spec.lettres
    spec.premiere = PREMIERE
    spec.derniere = PREMIERE + spec.lignes - 1
    derniere_col = get_column_letter(len(cols))

    _en_tete_feuille(ws, spec.titre, spec.sous_titre, len(cols))
    _lien_sommaire(ws, "A3")
    ws.merge_cells("A3:B3")
    _legende(ws, 3, 3)

    cles_in = [c for c in cols if c.kind == "in"]
    lettres_in = [L[c.key] for c in cles_in]
    groupes = _groupes_contigus(lettres_in)

    # Entêtes et aide
    for i, c in enumerate(cols, start=1):
        h = ws.cell(ENTETE, i, c.header)
        h.font = police(10, True, couleur="FFFFFF")
        h.fill = fill(VERT_FONCE if c.kind == "in" else ARDOISE)
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        h.border = BORDURE
        a = ws.cell(ENTETE - 1, i, c.hint)
        a.font = police(8, False, True, ARDOISE)
        a.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = c.width
        if c.cache:
            ws.column_dimensions[get_column_letter(i)].hidden = True
    ws.row_dimensions[ENTETE].height = 44
    ws.row_dimensions[ENTETE - 1].height = 40

    # Corps
    for r in range(spec.premiere, spec.derniere + 1):
        for i, c in enumerate(cols, start=1):
            cell = ws.cell(r, i)
            cell.border = BORDURE
            cell.font = police(10, False, False, "000000" if c.kind == "in" else ARDOISE)
            if c.kind != "in":
                cell.fill = fill(GRIS_CALC)
            if c.fmt:
                cell.number_format = c.fmt
            if c.align:
                cell.alignment = Alignment(horizontal=c.align, vertical="center")
            else:
                cell.alignment = Alignment(vertical="center")
            if c.kind == "static":
                cell.value = r - spec.premiere + 1
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c.kind == "calc" and c.f is not None:
                cell.value = c.f(r, L)
        ws.row_dimensions[r].height = 18

    # Validations
    for c in cols:
        if c.dv:
            dv = c.dv(L) if callable(c.dv) else c.dv
            _appliquer_validation(ws, L[c.key], spec.premiere, spec.derniere, dv)

    # Champs obligatoires manquants (ligne commencée)
    args_counta = ",".join(f"${a}{spec.premiere}:${b}{spec.premiere}" for a, b in groupes)
    rouge = dict(fill=fill(ROUGE_CLAIR), font=Font(color=ROUGE_TEXTE))
    for c in cols:
        lettre = L[c.key]
        plage = f"{lettre}{spec.premiere}:{lettre}{spec.derniere}"
        if c.req:
            formule = f'AND(COUNTA({args_counta})>0,{lettre}{spec.premiere}="")'
            ws.conditional_formatting.add(plage, FormulaRule(formula=[formule], **rouge))
        if c.req_si is not None:
            ws.conditional_formatting.add(plage, FormulaRule(formula=[c.req_si(L)], **rouge))
        if c.date:
            ws.conditional_formatting.add(
                plage, FormulaRule(formula=[f'AND({lettre}{spec.premiere}<>"",NOT(ISNUMBER({lettre}{spec.premiere})))'], **rouge)
            )

    # Noms définis par colonne
    for c in cols:
        if c.nommer:
            definir_nom(wb, f"{spec.prefixe}_{c.key}", ref_plage(spec.nom, L[c.key], spec.premiere, spec.derniere))

    # Présentation
    ws.freeze_panes = spec.gel
    ws.auto_filter.ref = f"A{ENTETE}:{derniere_col}{spec.derniere}"
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{ENTETE}:{ENTETE}"
    ws.print_options.gridLines = False
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    return ws


# --- Aides pour les feuilles de rapport ---------------------------------------------------------------------------


def section(ws, ligne: int, texte: str, col_debut: int, col_fin: int) -> None:
    """Bandeau de section (cellules fusionnées sur une seule ligne, hors tableaux de saisie)."""
    ws.merge_cells(start_row=ligne, start_column=col_debut, end_row=ligne, end_column=col_fin)
    c = ws.cell(ligne, col_debut, texte)
    c.font = police(11, True, couleur="FFFFFF")
    c.fill = fill(VERT)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in range(col_debut, col_fin + 1):
        ws.cell(ligne, col).fill = fill(VERT)
    ws.row_dimensions[ligne].height = 22


def entete_tableau(ws, ligne: int, libelles: list, col_debut: int = 1, hauteur: float = 32) -> None:
    for i, texte in enumerate(libelles):
        c = ws.cell(ligne, col_debut + i, texte)
        c.font = police(10, True, couleur="FFFFFF")
        c.fill = fill(ARDOISE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDURE
    ws.row_dimensions[ligne].height = hauteur


def cellule(ws, ligne: int, col: int, valeur=None, fmt=None, gras=False, couleur_fond=None, align=None, taille=10, wrap=False, couleur_texte="000000"):
    c = ws.cell(ligne, col)
    if valeur is not None:
        c.value = valeur
    c.font = police(taille, gras, False, couleur_texte)
    c.border = BORDURE
    if fmt:
        c.number_format = fmt
    if couleur_fond:
        c.fill = fill(couleur_fond)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    return c


def cellule_saisie(ws, ligne: int, col: int, valeur=None, fmt=None, validation=None):
    """Cellule de sélection d'un rapport : jaune, déverrouillée."""
    c = cellule(ws, ligne, col, valeur, fmt, gras=True, couleur_fond=JAUNE, align="center")
    c.protection = Protection(locked=False)
    if validation is not None:
        ws.add_data_validation(validation)
        validation.add(c.coordinate)
    return c


def etiquette(ws, ligne: int, col: int, texte: str):
    c = ws.cell(ligne, col, texte)
    c.font = police(10, True, couleur=VERT_FONCE)
    c.alignment = Alignment(horizontal="left", vertical="center")
    return c


def regle_pourcentage(ws, plage: str) -> None:
    """Vert ≥ 100 %, orange de 70 à 99 %, rouge < 70 % (les cellules vides ou « – » ne sont pas colorées)."""
    premiere = plage.split(":")[0]
    ws.conditional_formatting.add(
        plage, FormulaRule(formula=[f"AND(ISNUMBER({premiere}),{premiere}>=1)"], fill=fill("C6EFCE"), font=Font(color="006100", bold=True))
    )
    ws.conditional_formatting.add(
        plage, FormulaRule(formula=[f"AND(ISNUMBER({premiere}),{premiere}>=0.7,{premiere}<1)"], fill=fill("FFEB9C"), font=Font(color="7F6000", bold=True))
    )
    ws.conditional_formatting.add(
        plage, FormulaRule(formula=[f"AND(ISNUMBER({premiere}),{premiere}<0.7)"], fill=fill(ROUGE_CLAIR), font=Font(color=ROUGE_TEXTE, bold=True))
    )


def proteger(ws) -> None:
    """Protège une feuille de rapport (sans mot de passe) : seules les cellules jaunes se saisissent."""
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False


def mise_en_page(ws, orientation: str = "portrait", zone: Optional[str] = None) -> None:
    ws.page_setup.orientation = orientation
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    if zone:
        ws.print_area = zone
    ws.sheet_view.zoomScale = 90
    ws.sheet_view.showGridLines = False


def habiller_graphique(graphique) -> None:
    """Rend un graphique lisible : axes visibles, titre et légende hors du tracé, quadrillage clair, courbes droites."""
    from openpyxl.chart import LineChart
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties

    if graphique.title is not None:
        graphique.title.overlay = False
    if graphique.legend is not None:
        graphique.legend.overlay = False
    for nom in ("x_axis", "y_axis"):
        axe = getattr(graphique, nom, None)
        if axe is not None:
            axe.delete = False
    if getattr(graphique, "y_axis", None) is not None and graphique.y_axis.majorGridlines is not None:
        graphique.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D9D9D9", w=6350))
    if getattr(graphique, "x_axis", None) is not None:
        graphique.x_axis.tickLblPos = "low"
    for sous in getattr(graphique, "_charts", [graphique]):
        if isinstance(sous, LineChart):
            for serie in sous.series:
                serie.smooth = False
                serie.marker.symbol = "circle"
                serie.marker.size = 5
