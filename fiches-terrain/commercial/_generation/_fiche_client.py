"""Feuille « Fiche_Client » : recherche d'un client (suggestions) et historique complet.

On tape une partie du nom, du téléphone, du quartier ou du type : les clients qui correspondent s'affichent ;
on en choisit un dans la liste déroulante (ou on prend le premier) et la fiche, ses dernières visites, commandes
et encaissements apparaissent. Aucune saisie dans les journaux : tout est lu.
"""

from __future__ import annotations

from copy import copy

from openpyxl.worksheet.datavalidation import DataValidation

from _outils import (
    ARDOISE,
    COULEURS_ONGLETS,
    FMT_DATE,
    FMT_HEURE,
    FMT_XAF,
    GRIS_CALC,
    JAUNE,
    cellule,
    cellule_saisie,
    definir_nom,
    entete_tableau,
    etiquette,
    mise_en_page,
    proteger,
    section,
)
from _rapports import GRAND, _en_tete, _masquer, validation_commercial

NB_RESULTATS = 15
PREMIERE_RESULTAT = 9


def fiche_client(wb) -> None:
    ws = wb.create_sheet("Fiche_Client")
    ws.sheet_properties.tabColor = COULEURS_ONGLETS["rapport"]
    for col, largeur in zip("ABCDEFGH", (38, 16, 18, 30, 16, 16, 16, 16)):
        ws.column_dimensions[col].width = largeur
    _en_tete(
        ws,
        "FICHE CLIENT — recherche et historique",
        "Tapez quelques lettres du nom, du téléphone, du quartier ou du type (cellule jaune) : les clients trouvés s'affichent. Choisissez-en un, sa fiche et son historique apparaissent.",
    )
    etiquette(ws, 4, 1, "Rechercher")
    cellule_saisie(ws, 4, 2, None)
    ws.merge_cells("B4:D4")
    for col in (3, 4):
        ws.cell(4, col).fill = copy(ws.cell(4, 2).fill)
    etiquette(ws, 4, 5, "Commercial")
    cellule_saisie(ws, 4, 6, "=Commercial_Id", None, validation_commercial())
    etiquette(ws, 5, 1, "Client à afficher (liste des résultats)")
    cellule_saisie(ws, 5, 2, None)
    ws.merge_cells("B5:D5")
    for col in (3, 4):
        ws.cell(5, col).fill = copy(ws.cell(5, 2).fill)
    cellule(ws, 5, 5, "Vide = le premier résultat", taille=9, couleur_texte=ARDOISE)
    ws.merge_cells("E5:H5")

    Q, C = "$B$4", "$F$4"
    # Correspondance d'une ligne de la feuille Clients avec la recherche (texte libre, sans casse)
    correspond = f'(Cl_commercial={C})*(Cl_nom<>"")*ISNUMBER(SEARCH({Q},Cl_nom&" "&Cl_tel&" "&Cl_quartier&" "&Cl_type))'
    ws["K4"] = f"=SUMPRODUCT({correspond})"
    ws["K5"] = (
        f'=IF($B$5="",IF(COUNT($J${PREMIERE_RESULTAT}:$J${PREMIERE_RESULTAT + NB_RESULTATS - 1})=0,0,$J${PREMIERE_RESULTAT}-5),'
        f'IFERROR(MATCH(1,INDEX((Cl_label=$B$5)*(Cl_commercial={C}),0),0),0))'
    )
    S = "$K$5"  # ligne (dans la plage des clients) du client affiché ; 0 = aucun

    # --- Résultats ---------------------------------------------------------------------------------------------------
    section(ws, 7, "RÉSULTATS DE LA RECHERCHE (15 premiers)", 1, 8)
    ws["K7"] = "Nb de clients trouvés"
    entete_tableau(ws, 8, ["Client", "Quartier", "Téléphone", "Type", "Statut", "CA cumulé (XAF)", "Reste dû (XAF)", "Dernier contact"])
    for k in range(NB_RESULTATS):
        r = PREMIERE_RESULTAT + k
        ws[f"J{r}"] = f'=IF({k + 1}>$K$4,"",SMALL(INDEX(ROW(Cl_nom)+(1-{correspond})*{GRAND},0),{k + 1}))'
        idx = f"$J{r}-5"
        cellule(ws, r, 1, f'=IF($J{r}="","",INDEX(Cl_label,{idx}))')
        cellule(ws, r, 2, f'=IF($J{r}="","",INDEX(Cl_quartier,{idx})&"")')
        cellule(ws, r, 3, f'=IF($J{r}="","",INDEX(Cl_tel,{idx})&"")')
        cellule(ws, r, 4, f'=IF($J{r}="","",INDEX(Cl_type,{idx})&"")')
        cellule(ws, r, 5, f'=IF($J{r}="","",INDEX(Cl_statut,{idx})&"")', align="center")
        cellule(ws, r, 6, f'=IF($J{r}="","",INDEX(Cl_ca,{idx}))', FMT_XAF, align="right")
        cellule(ws, r, 7, f'=IF($J{r}="","",INDEX(Cl_reste,{idx}))', FMT_XAF, align="right")
        cellule(ws, r, 8, f'=IF($J{r}="","",INDEX(Cl_dernier,{idx}))', FMT_DATE, align="center")
    derniere_res = PREMIERE_RESULTAT + NB_RESULTATS - 1
    definir_nom(wb, "L_ResultatsClient", f"OFFSET('Fiche_Client'!$A${PREMIERE_RESULTAT},0,0,MAX(1,COUNTIF('Fiche_Client'!$A${PREMIERE_RESULTAT}:$A${derniere_res},\"?*\")),1)")
    dv = DataValidation(
        type="list", formula1="=L_ResultatsClient", allow_blank=True, showErrorMessage=True, errorStyle="stop",
        errorTitle="Client hors résultats", error="Choisissez un client parmi les résultats de la recherche (liste déroulante).",
    )
    ws.add_data_validation(dv)
    dv.add("B5")

    # --- Fiche -------------------------------------------------------------------------------------------------------
    section(ws, 25, "FICHE DU CLIENT", 1, 8)
    gauche = [
        ("Enseigne / nom", "Cl_nom", None),
        ("Code client", "Cl_code", None),
        ("Statut (automatique)", "Cl_statut", None),
        ("Type de client", "Cl_type", None),
        ("Interlocuteur", "Cl_interlocuteur", None),
        ("Téléphone principal", "Cl_tel", None),
        ("Téléphone 2 / WhatsApp", "Cl_tel2", None),
        ("Quartier / zone", "Cl_quartier", None),
        ("Point de repère", "Cl_repere", None),
        ("Produit principal / potentiel par jour", None, None),
    ]
    droite = [
        ("CA cumulé (XAF)", "Cl_ca", FMT_XAF),
        ("Encaissé (XAF)", "Cl_encaisse", FMT_XAF),
        ("Reste à encaisser (XAF)", "Cl_reste", FMT_XAF),
        ("Nombre de visites", "Cl_nbvisites", "0"),
        ("Dernier contact", "Cl_dernier", FMT_DATE),
        ("Dernier achat", "Cl_dernierevente", FMT_DATE),
        ("Prochaine action : date", "Cl_actiondate", FMT_DATE),
        ("Prochaine action : à faire", "Cl_actiontexte", None),
        ("Alerte", "Cl_alerte", None),
        ("Observations", "Cl_obs", None),
    ]
    for i, (libelle, nom, fmt) in enumerate(gauche):
        r = 26 + i
        cellule(ws, r, 1, libelle, gras=True)
        if nom:
            valeur = f'=IF({S}<1,"",INDEX({nom},{S})&"")'
        else:
            valeur = f'=IF({S}<1,"",INDEX(Cl_produit,{S})&IF(INDEX(Cl_potentiel,{S})="",""," · "&INDEX(Cl_potentiel,{S})&" "&INDEX(Cl_unite,{S})&" / jour"))'
        cellule(ws, r, 2, valeur, couleur_fond=GRIS_CALC)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.cell(r, 3).border = copy(ws.cell(r, 2).border)
    for i, (libelle, nom, fmt) in enumerate(droite):
        r = 26 + i
        cellule(ws, r, 4, libelle, gras=True)
        if fmt:
            valeur = f'=IF({S}<1,"",INDEX({nom},{S}))'
        else:
            valeur = f'=IF({S}<1,"",INDEX({nom},{S})&"")'
        cellule(ws, r, 5, valeur, fmt, couleur_fond=GRIS_CALC, align="right" if fmt and fmt != FMT_DATE else None)
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
        for col in (6, 7, 8):
            ws.cell(r, col).border = copy(ws.cell(r, 5).border)

    # --- Historique : visites ------------------------------------------------------------------------------------------
    cle = f'INDEX(Cl_cle,{S})'
    section(ws, 37, "DERNIÈRES VISITES ET CONTACTS (8)", 1, 8)
    entete_tableau(ws, 38, ["Objet → résultat", "Date", "Heure", "Type de contact", "Prochaine action", "", "Notes", ""], hauteur=22)
    ws.merge_cells("E38:F38")
    ws.merge_cells("G38:H38")
    for k in range(8):
        r = 39 + k
        ws[f"J{r}"] = (
            f'=IF({S}<1,"",IF(LARGE(INDEX(ROW(VC_cle)*(VC_cle={cle}),0),{k + 1})=0,"",LARGE(INDEX(ROW(VC_cle)*(VC_cle={cle}),0),{k + 1})))'
        )
        idx = f"$J{r}-5"
        cellule(ws, r, 1, f'=IF($J{r}="","",INDEX(VC_objet,{idx})&IF(INDEX(VC_resultat,{idx})="",""," → "&INDEX(VC_resultat,{idx})))')
        cellule(ws, r, 2, f'=IF($J{r}="","",INDEX(VC_date,{idx}))', FMT_DATE, align="center")
        cellule(ws, r, 3, f'=IF($J{r}="","",IF(INDEX(VC_heure,{idx})="","",INDEX(VC_heure,{idx})))', FMT_HEURE, align="center")
        cellule(ws, r, 4, f'=IF($J{r}="","",INDEX(VC_type,{idx})&"")')
        cellule(ws, r, 5, f'=IF($J{r}="","",IF(INDEX(VC_actiondate,{idx})="","",TEXT(INDEX(VC_actiondate,{idx}),"dd/mm/yyyy")&" : ")&INDEX(VC_actiontexte,{idx}))')
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
        ws.cell(r, 6).border = copy(ws.cell(r, 5).border)
        cellule(ws, r, 7, f'=IF($J{r}="","",INDEX(VC_notes,{idx})&"")')
        ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=8)
        ws.cell(r, 8).border = copy(ws.cell(r, 7).border)

    # --- Historique : commandes ------------------------------------------------------------------------------------------
    section(ws, 48, "DERNIÈRES LIGNES DE COMMANDE (6)", 1, 8)
    entete_tableau(ws, 49, ["Produit", "Date", "Quantité", "N° de commande · statut", "Montant (XAF)", "", "", ""], hauteur=22)
    ws.merge_cells("E49:H49")
    for k in range(6):
        r = 50 + k
        ws[f"J{r}"] = (
            f'=IF({S}<1,"",IF(LARGE(INDEX(ROW(Cmd_cle)*(Cmd_cle={cle}),0),{k + 1})=0,"",LARGE(INDEX(ROW(Cmd_cle)*(Cmd_cle={cle}),0),{k + 1})))'
        )
        idx = f"$J{r}-5"
        cellule(ws, r, 1, f'=IF($J{r}="","",INDEX(Cmd_produit,{idx})&"")')
        cellule(ws, r, 2, f'=IF($J{r}="","",INDEX(Cmd_date,{idx}))', FMT_DATE, align="center")
        cellule(ws, r, 3, f'=IF($J{r}="","",INDEX(Cmd_qte,{idx})&" "&INDEX(Cmd_unite,{idx}))', align="center")
        cellule(ws, r, 4, f'=IF($J{r}="","",INDEX(Cmd_bon,{idx})&" · "&INDEX(Cmd_statut,{idx}))')
        cellule(ws, r, 5, f'=IF($J{r}="","",INDEX(Cmd_montant,{idx}))', FMT_XAF, align="right")
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
        for col in (6, 7, 8):
            ws.cell(r, col).border = copy(ws.cell(r, 5).border)

    # --- Historique : encaissements ------------------------------------------------------------------------------------------
    section(ws, 57, "DERNIERS ENCAISSEMENTS (5)", 1, 8)
    entete_tableau(ws, 58, ["Mode de paiement", "Date", "Montant (XAF)", "Référence", "Versé à l'entreprise ?", "", "", ""], hauteur=22)
    ws.merge_cells("E58:H58")
    for k in range(5):
        r = 59 + k
        ws[f"J{r}"] = (
            f'=IF({S}<1,"",IF(LARGE(INDEX(ROW(Enc_cle)*(Enc_cle={cle}),0),{k + 1})=0,"",LARGE(INDEX(ROW(Enc_cle)*(Enc_cle={cle}),0),{k + 1})))'
        )
        idx = f"$J{r}-5"
        cellule(ws, r, 1, f'=IF($J{r}="","",INDEX(Enc_mode,{idx})&"")')
        cellule(ws, r, 2, f'=IF($J{r}="","",INDEX(Enc_date,{idx}))', FMT_DATE, align="center")
        cellule(ws, r, 3, f'=IF($J{r}="","",INDEX(Enc_montant,{idx}))', FMT_XAF, align="right")
        cellule(ws, r, 4, f'=IF($J{r}="","",INDEX(Enc_reference,{idx})&"")')
        cellule(ws, r, 5, f'=IF($J{r}="","",INDEX(Enc_verse,{idx})&"")', align="center")
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
        for col in (6, 7, 8):
            ws.cell(r, col).border = copy(ws.cell(r, 5).border)

    for col in "JK":
        _masquer(ws, col)
    mise_en_page(ws, "portrait", "A1:H63")
    proteger(ws)
