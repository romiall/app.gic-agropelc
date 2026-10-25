"""Feuille « Tableau_de_bord » : les indicateurs clés et leur évolution dans la durée (dernière page).

Tout est calculé à partir des lignes saisies. On choisit la semaine de référence (un lundi) et le commercial ;
le tableau compare la semaine à la précédente, suit le mois, le portefeuille et 12 semaines d'évolution, puis
répartit le chiffre d'affaires par produit et par quartier et les résultats des visites.
"""

from __future__ import annotations

from copy import copy

from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font
from openpyxl.worksheet.pagebreak import Break

from _donnees import (
    ETAPES,
    M_CA,
    M_NVX_CLIENTS,
    M_PROSPECTS,
    M_VISITES,
    ST_CLIENT,
    ST_PERDU,
    ST_PROSPECT,
)
from _kpi import (
    iso_semaine,
    k_ca,
    k_casses,
    k_commandes_xaf,
    k_contacts,
    k_encaisse,
    k_especes_a_verser,
    k_frais,
    k_jours_travailles,
    k_nouveaux_clients,
    k_objectif,
    k_prospects,
    k_ventes_nb,
    k_visites,
    k_visites_commande,
    nom_commercial,
)
from _outils import (
    ARDOISE,
    BORDURE,
    COULEURS_ONGLETS,
    FMT_DATE,
    FMT_ENTIER,
    FMT_POURCENT,
    FMT_XAF,
    GRIS_CALC,
    ROUGE_CLAIR,
    ROUGE_TEXTE,
    cellule,
    cellule_saisie,
    entete_tableau,
    etiquette,
    fill,
    habiller_graphique,
    mise_en_page,
    proteger,
    regle_pourcentage,
    section,
)
from _rapports import GRAND, _en_tete, _masquer, selecteur_lundi, validation_commercial

NB_SEMAINES = 12
VERT_BON = "C6EFCE"
VERT_TEXTE = "006100"


def _evolution(courant: str, precedent: str, en_points: bool = False) -> str:
    """Texte « ▲ +12 % » / « ▼ 8 % » ; « ▲ nouveau » si la semaine précédente est à zéro."""
    if en_points:
        return (
            f'=IF(AND({courant}="",{precedent}=""),"–",IF({courant}="","▼ "&TEXT(N({precedent})*100,"0")&" pts",'
            f'IF({precedent}="","▲ +"&TEXT(N({courant})*100,"0")&" pts",'
            f'IF(N({courant})>=N({precedent}),"▲ +","▼ ")&TEXT(ABS(N({courant})-N({precedent}))*100,"0")&" pts")))'
        )
    return (
        f'=IF(AND(N({courant})=0,N({precedent})=0),"–",IF(N({precedent})=0,"▲ nouveau",'
        f'IF(N({courant})>=N({precedent}),"▲ +","▼ ")&TEXT(ABS(N({courant})/N({precedent})-1),"0%")))'
    )


def _regle_tendance(ws, plage: str, hausse_bonne: bool) -> None:
    """Vert quand la tendance est favorable, rouge quand elle ne l'est pas."""
    premiere = plage.split(":")[0]
    bon = f'LEFT({premiere},1)="{"▲" if hausse_bonne else "▼"}"'
    mauvais = f'LEFT({premiere},1)="{"▼" if hausse_bonne else "▲"}"'
    ws.conditional_formatting.add(plage, FormulaRule(formula=[bon], fill=fill(VERT_BON), font=Font(color=VERT_TEXTE, bold=True)))
    ws.conditional_formatting.add(plage, FormulaRule(formula=[mauvais], fill=fill(ROUGE_CLAIR), font=Font(color=ROUGE_TEXTE, bold=True)))


def _axes_visibles(graphique) -> None:
    habiller_graphique(graphique)


def tableau_de_bord(wb) -> None:
    ws = wb.create_sheet("Tableau_de_bord")
    ws.sheet_properties.tabColor = COULEURS_ONGLETS["rapport"]
    largeurs = [34, 13, 13, 13, 13, 17, 15, 15, 14, 14, 12, 16]
    for i, largeur in enumerate(largeurs, start=1):
        ws.column_dimensions[chr(64 + i)].width = largeur
    ws.column_dimensions["M"].width = 3
    _en_tete(
        ws,
        "TABLEAU DE BORD — la performance dans la durée",
        "Choisissez la semaine de référence (un lundi) et le commercial : les indicateurs, leur évolution sur 12 semaines et les répartitions se calculent seuls.",
    )

    etiquette(ws, 4, 1, "Semaine de référence (lundi)")
    cellule_saisie(ws, 4, 2, "=TODAY()-WEEKDAY(TODAY(),2)+1", FMT_DATE, selecteur_lundi())
    etiquette(ws, 4, 4, "Commercial")
    cellule_saisie(ws, 4, 5, "=Commercial_Id", None, validation_commercial())
    ws.merge_cells("E4:F4")
    D, C = "$B$4", "$E$4"
    FIN = f"({D}+6)"  # dimanche de la semaine
    PREC, PREC_FIN = f"({D}-7)", f"({D}-1)"
    etiquette(ws, 5, 1, "Semaine")
    cellule(ws, 5, 2, f"={iso_semaine(D)}", gras=True, couleur_fond=GRIS_CALC, align="center")
    etiquette(ws, 5, 4, "Nom")
    cellule(ws, 5, 5, f"={nom_commercial(C)}", gras=True, couleur_fond=GRIS_CALC)
    ws.merge_cells("E5:H5")
    # début et fin du mois (celui du jeudi de la semaine), début de la fenêtre de 12 semaines
    ws["Z1"] = f"=DATE(YEAR({D}+3),MONTH({D}+3),1)"
    ws["Z2"] = f"=DATE(YEAR({D}+3),MONTH({D}+3)+1,0)"
    ws["Z3"] = f"={D}-7*{NB_SEMAINES - 1}"
    M1, M2, F1 = "$Z$1", "$Z$2", "$Z$3"

    # --- 1. La semaine et la précédente ----------------------------------------------------------------------------
    section(ws, 7, "1. LA SEMAINE EN CHIFFRES (comparée à la précédente)", 1, 4)
    entete_tableau(ws, 8, ["Indicateur", "Cette semaine", "Semaine précédente", "Évolution"], hauteur=32)
    kpis = [
        ("Chiffre d'affaires réalisé (XAF)", k_ca, FMT_XAF, True, False),
        ("Commandes prises (XAF)", k_commandes_xaf, FMT_XAF, True, False),
        ("Encaissé (XAF)", k_encaisse, FMT_XAF, True, False),
        ("Visites sur place", k_visites, FMT_ENTIER, True, False),
        ("Nouveaux prospects", k_prospects, FMT_ENTIER, True, False),
        ("Nouveaux clients", k_nouveaux_clients, FMT_ENTIER, True, False),
        ("Taux de conversion (visites → commande)", None, FMT_POURCENT, True, True),
        ("Panier moyen (XAF par vente)", None, FMT_XAF, True, False),
        ("Frais engagés (XAF)", k_frais, FMT_XAF, False, False),
        ("Casses et pertes (XAF)", k_casses, FMT_XAF, False, False),
    ]
    for i, (libelle, fonction, fmt, hausse_bonne, en_points) in enumerate(kpis):
        r = 9 + i
        cellule(ws, r, 1, libelle, gras=True)
        for col, (a, b) in ((2, (D, FIN)), (3, (PREC, PREC_FIN))):
            if libelle.startswith("Taux de conversion"):
                formule = f'=IF({k_visites(a, b, C)}=0,"",{k_visites_commande(a, b, C)}/{k_visites(a, b, C)})'
            elif libelle.startswith("Panier moyen"):
                formule = f'=IF({k_ventes_nb(a, b, C)}=0,"",{k_ca(a, b, C)}/{k_ventes_nb(a, b, C)})'
            else:
                formule = f"={fonction(a, b, C)}"
            cellule(ws, r, col, formule, fmt, couleur_fond=GRIS_CALC, align="right")
        cellule(ws, r, 4, _evolution(f"B{r}", f"C{r}", en_points), align="center", gras=True)
        _regle_tendance(ws, f"D{r}:D{r}", hausse_bonne)

    # --- 2. Le mois en cours -----------------------------------------------------------------------------------------
    section(ws, 7, "2. LE MOIS EN COURS (objectif du mois)", 6, 10)
    entete_tableau(ws, 8, ["Indicateur", "", "Réalisé", "Objectif", "Atteinte"], col_debut=6, hauteur=32)
    ws.merge_cells("F8:G8")
    mois = [
        ("Chiffre d'affaires réalisé (XAF)", f"={k_ca(M1, M2, C)}", FMT_XAF, M_CA),
        ("Visites sur place", f"={k_visites(M1, M2, C)}", FMT_ENTIER, M_VISITES),
        ("Nouveaux clients", f"={k_nouveaux_clients(M1, M2, C)}", FMT_ENTIER, M_NVX_CLIENTS),
        ("Prospects créés", f"={k_prospects(M1, M2, C)}", FMT_ENTIER, M_PROSPECTS),
    ]
    for i, (libelle, formule, fmt, metrique) in enumerate(mois):
        r = 9 + i
        _libelle_fusionne(ws, r, libelle)
        cellule(ws, r, 8, formule, fmt, couleur_fond=GRIS_CALC, align="right")
        cellule(ws, r, 9, f"={k_objectif(metrique, M1, M2, C, False)}", fmt, couleur_fond=GRIS_CALC, align="right")
        cellule(ws, r, 10, f'=IF(OR(I{r}="",N(I{r})=0),"",H{r}/I{r})', FMT_POURCENT, align="center", gras=True)
    regle_pourcentage(ws, "J9:J12")
    simples = [
        ("Encaissé dans le mois (XAF)", f"={k_encaisse(M1, M2, C)}", FMT_XAF),
        ("Taux de recouvrement (encaissé ÷ CA)", '=IF(N(H9)=0,"",H13/H9)', FMT_POURCENT),
        ("Jours travaillés (Journal)", f"={k_jours_travailles(M1, M2, C)}", FMT_ENTIER),
        ("CA moyen par jour travaillé (XAF)", '=IF(N(H15)=0,"",H9/H15)', FMT_XAF),
        ("CA par visite sur place (XAF)", '=IF(N(H10)=0,"",H9/H10)', FMT_XAF),
    ]
    for i, (libelle, formule, fmt) in enumerate(simples):
        r = 13 + i
        _libelle_fusionne(ws, r, libelle)
        cellule(ws, r, 8, formule, fmt, couleur_fond=GRIS_CALC, align="right")
        cellule(ws, r, 9, None, couleur_fond=GRIS_CALC)
        cellule(ws, r, 10, None)

    # --- 3. Portefeuille -----------------------------------------------------------------------------------------------
    section(ws, 20, "3. MON PORTEFEUILLE (état à ce jour)", 1, 10)
    gauche = [
        ("Comptes au total", f'=COUNTIFS(Cl_commercial,{C},Cl_nom,"?*")', FMT_ENTIER),
        ("Prospects", f'=COUNTIFS(Cl_commercial,{C},Cl_statut,"{ST_PROSPECT}")', FMT_ENTIER),
        ("Clients", f'=COUNTIFS(Cl_commercial,{C},Cl_statut,"{ST_CLIENT}")', FMT_ENTIER),
        ("Comptes perdus", f'=COUNTIFS(Cl_commercial,{C},Cl_statut,"{ST_PERDU}")', FMT_ENTIER),
        ("Taux de transformation (clients ÷ comptes)", '=IF(B21=0,"",B23/B21)', FMT_POURCENT),
    ]
    droite = [
        ("Clients actifs (achat récent)", f'=COUNTIFS(Cl_commercial,{C},Cl_statut,"{ST_CLIENT}",Cl_dernierevente,">="&TODAY()-P_JoursInactif)', FMT_ENTIER),
        ("Clients inactifs", "=B23-H21", FMT_ENTIER),
        ("Créances : reste à encaisser (XAF)", f'=SUMIFS(Cl_reste,Cl_commercial,{C},Cl_reste,">0")', FMT_XAF),
        ("Espèces détenues, non versées (XAF)", f"={k_especes_a_verser('TODAY()', C)}", FMT_XAF),
        ("Comptes à surveiller (alerte)", f'=COUNTIFS(Cl_commercial,{C},Cl_alerte,"?*")', FMT_ENTIER),
    ]
    for i, (libelle, formule, fmt) in enumerate(gauche):
        r = 21 + i
        cellule(ws, r, 1, libelle, gras=True)
        cellule(ws, r, 2, formule, fmt, couleur_fond=GRIS_CALC, align="right", gras=True)
    for i, (libelle, formule, fmt) in enumerate(droite):
        r = 21 + i
        _libelle_fusionne(ws, r, libelle)
        cellule(ws, r, 8, formule, fmt, couleur_fond=GRIS_CALC, align="right", gras=True)
    # entonnoir : étapes des prospects
    cellule(ws, 21, 3, "Entonnoir des prospects", gras=True, couleur_texte=ARDOISE)
    ws.merge_cells("C21:D21")
    for i, (etape, _) in enumerate(ETAPES):
        r = 22 + i
        cellule(ws, r, 3, etape)
        cellule(ws, r, 4, f'=COUNTIFS(Cl_commercial,{C},Cl_statut,"{ST_PROSPECT}",Cl_etape,"{etape}")', FMT_ENTIER, couleur_fond=GRIS_CALC, align="center")
    cellule(ws, 26, 3, "→ Clients", gras=True)
    cellule(ws, 26, 4, "=B23", FMT_ENTIER, couleur_fond=GRIS_CALC, align="center", gras=True)

    # --- 4. Évolution sur 12 semaines ---------------------------------------------------------------------------------------
    section(ws, 28, f"4. ÉVOLUTION SUR {NB_SEMAINES} SEMAINES (jusqu'à la semaine de référence)", 1, 12)
    entete_tableau(
        ws,
        29,
        ["Semaine du (lundi)", "Visites sur place", "Contacts à distance", "Nouveaux prospects", "Nouveaux clients", "Commandes prises (XAF)", "CA réalisé (XAF)", "Encaissé (XAF)", "Frais (XAF)", "Casses (XAF)", "Conversion visite → commande", "CA cumulé (XAF)"],
        hauteur=46,
    )
    debut_serie = 30
    for i in range(NB_SEMAINES):
        r = debut_serie + i
        cellule(ws, r, 1, f"={D}-7*{NB_SEMAINES - 1 - i}", FMT_DATE, couleur_fond=GRIS_CALC, align="center", gras=True)
        a, b = f"$A{r}", f"($A{r}+6)"
        colonnes = [
            (2, k_visites(a, b, C), FMT_ENTIER),
            (3, k_contacts(a, b, C), FMT_ENTIER),
            (4, k_prospects(a, b, C), FMT_ENTIER),
            (5, k_nouveaux_clients(a, b, C), FMT_ENTIER),
            (6, k_commandes_xaf(a, b, C), FMT_XAF),
            (7, k_ca(a, b, C), FMT_XAF),
            (8, k_encaisse(a, b, C), FMT_XAF),
            (9, k_frais(a, b, C), FMT_XAF),
            (10, k_casses(a, b, C), FMT_XAF),
        ]
        for col, formule, fmt in colonnes:
            cellule(ws, r, col, f"={formule}", fmt, couleur_fond=GRIS_CALC, align="right")
        cellule(ws, r, 11, f'=IF(B{r}=0,"",{k_visites_commande(a, b, C)}/B{r})', FMT_POURCENT, couleur_fond=GRIS_CALC, align="center")
        cellule(ws, r, 12, f"=SUM(G${debut_serie}:G{r})", FMT_XAF, couleur_fond=GRIS_CALC, align="right")
    fin_serie = debut_serie + NB_SEMAINES - 1
    cellule(ws, fin_serie + 1, 1, f"TOTAL SUR {NB_SEMAINES} SEMAINES", gras=True, couleur_fond="D9E8DF")
    for col in range(2, 11):
        lettre = chr(64 + col)
        cellule(ws, fin_serie + 1, col, f"=SUM({lettre}{debut_serie}:{lettre}{fin_serie})", FMT_XAF if col >= 6 else FMT_ENTIER, gras=True, couleur_fond="D9E8DF", align="right")
    cellule(ws, fin_serie + 1, 11, f'=IF(B{fin_serie + 1}=0,"",{k_visites_commande(F1, FIN, C)}/B{fin_serie + 1})', FMT_POURCENT, gras=True, couleur_fond="D9E8DF", align="center")
    cellule(ws, fin_serie + 1, 12, None, couleur_fond="D9E8DF")

    # --- 5. Graphiques d'évolution ------------------------------------------------------------------------------------------------
    section(ws, 44, "5. GRAPHIQUES D'ÉVOLUTION", 1, 12)
    categories = Reference(ws, min_col=1, min_row=debut_serie, max_row=fin_serie)

    ca = BarChart()
    ca.type = "col"
    ca.title = "CA réalisé et encaissé par semaine (XAF)"
    ca.add_data(Reference(ws, min_col=7, max_col=8, min_row=29, max_row=fin_serie), titles_from_data=True)
    ca.set_categories(categories)
    ca.x_axis.number_format = "dd/mm"
    ca.height, ca.width = 8, 13.2
    ca.legend.position = "b"
    _axes_visibles(ca)
    ws.add_chart(ca, "A45")

    activite = LineChart()
    activite.title = "Activité terrain par semaine"
    activite.add_data(Reference(ws, min_col=2, min_row=29, max_row=fin_serie), titles_from_data=True)
    activite.add_data(Reference(ws, min_col=4, min_row=29, max_row=fin_serie), titles_from_data=True)
    activite.set_categories(categories)
    activite.x_axis.number_format = "dd/mm"
    activite.height, activite.width = 8, 13.2
    activite.legend.position = "b"
    _axes_visibles(activite)
    ws.add_chart(activite, "E45")

    clients = BarChart()
    clients.type = "col"
    clients.title = "Nouveaux clients et CA cumulé"
    clients.add_data(Reference(ws, min_col=5, min_row=29, max_row=fin_serie), titles_from_data=True)
    clients.set_categories(categories)
    clients.x_axis.number_format = "dd/mm"
    clients.height, clients.width = 8, 13.2
    clients.legend.position = "b"
    _axes_visibles(clients)
    cumul = LineChart()
    cumul.add_data(Reference(ws, min_col=12, min_row=29, max_row=fin_serie), titles_from_data=True)
    cumul.y_axis.axId = 200
    cumul.y_axis.crosses = "max"
    cumul.y_axis.delete = False
    cumul.y_axis.majorGridlines = None
    clients += cumul
    habiller_graphique(clients)
    ws.add_chart(clients, "J45")

    # --- 6. Répartitions ----------------------------------------------------------------------------------------------------------
    section(ws, 62, f"6. RÉPARTITIONS SUR {NB_SEMAINES} SEMAINES", 1, 12)
    entete_tableau(ws, 63, ["CA par produit", "CA (XAF)", "Part"], hauteur=22)
    for k in range(7):
        r = 64 + k
        cellule(ws, r, 1, f'=IFERROR(INDEX(L_Produit,{k + 1})&"","")')
        cellule(ws, r, 2, f'=IF($A{r}="","",SUMIFS(Cmd_montant,Cmd_produit,$A{r},{_entre12("Cmd_dateca", F1, FIN)},Cmd_commercial,{C}))', FMT_XAF, couleur_fond=GRIS_CALC, align="right")
        cellule(ws, r, 3, f'=IF(OR($A{r}="",N($B$71)=0),"",B{r}/$B$71)', FMT_POURCENT, couleur_fond=GRIS_CALC, align="center")
    cellule(ws, 71, 1, "Total", gras=True, couleur_fond="D9E8DF")
    cellule(ws, 71, 2, "=SUM(B64:B70)", FMT_XAF, gras=True, couleur_fond="D9E8DF", align="right")
    cellule(ws, 71, 3, None, couleur_fond="D9E8DF")

    entete_tableau(ws, 63, ["Top quartiers (CA)", "CA (XAF)", "Part"], col_debut=5, hauteur=22)
    # table d'aide (masquée) : CA de chaque quartier de la liste, avec une clé de classement
    nb_zones = 19
    for k in range(nb_zones):
        r = 30 + k
        ws[f"Q{r}"] = f'=IFERROR(INDEX(L_Zone,{k + 1})&"","")'
        ws[f"R{r}"] = f'=IF(Q{r}="",0,SUMIFS(Cmd_montant,Cmd_quartier,Q{r},{_entre12("Cmd_dateca", F1, FIN)},Cmd_commercial,{C}))'
        ws[f"S{r}"] = f"=IF(R{r}>0,ROUND(R{r},0)+ROW()/1000000,0)"
    for k in range(8):
        r = 64 + k
        cle = f"LARGE($S$30:$S${29 + nb_zones},{k + 1})"
        ws[f"P{r}"] = f'=IF({cle}=0,"",ROUND(MOD({cle},1)*1000000,0))'
        cellule(ws, r, 5, f'=IF($P{r}="","",INDEX($Q$1:$Q$200,$P{r}))')
        cellule(ws, r, 6, f'=IF($P{r}="","",INDEX($R$1:$R$200,$P{r}))', FMT_XAF, couleur_fond=GRIS_CALC, align="right")
        cellule(ws, r, 7, f'=IF(OR($P{r}="",N($B$71)=0),"",F{r}/$B$71)', FMT_POURCENT, couleur_fond=GRIS_CALC, align="center")

    entete_tableau(ws, 63, ["Résultat des visites", "Nombre", "Part"], col_debut=9, hauteur=32)
    for k in range(9):
        r = 64 + k
        cellule(ws, r, 9, f'=IFERROR(INDEX(L_Resultat,{k + 1})&"","")')
        cellule(ws, r, 10, f'=IF($I{r}="","",COUNTIFS(VC_resultat,$I{r},{_entre12("VC_date", F1, FIN)},VC_commercial,{C}))', FMT_ENTIER, couleur_fond=GRIS_CALC, align="center")
        cellule(ws, r, 11, f'=IF(OR($I{r}="",N($J$73)=0),"",J{r}/$J$73)', FMT_POURCENT, couleur_fond=GRIS_CALC, align="center")
    cellule(ws, 73, 9, "Total", gras=True, couleur_fond="D9E8DF")
    cellule(ws, 73, 10, "=SUM(J64:J72)", FMT_ENTIER, gras=True, couleur_fond="D9E8DF", align="center")
    cellule(ws, 73, 11, None, couleur_fond="D9E8DF")

    # --- 7. Graphiques de répartition -------------------------------------------------------------------------------------------------------
    section(ws, 75, "7. GRAPHIQUES DE RÉPARTITION", 1, 12)
    produit = BarChart()
    produit.type = "bar"
    produit.title = "CA par produit (XAF)"
    produit.add_data(Reference(ws, min_col=2, min_row=63, max_row=70), titles_from_data=True)
    produit.set_categories(Reference(ws, min_col=1, min_row=64, max_row=70))
    produit.height, produit.width = 8, 13.2
    produit.legend = None
    produit.y_axis.number_format = '0,"k"'
    _axes_visibles(produit)
    ws.add_chart(produit, "A76")

    quartier = BarChart()
    quartier.type = "bar"
    quartier.title = "Top quartiers par CA (XAF)"
    quartier.add_data(Reference(ws, min_col=6, min_row=63, max_row=71), titles_from_data=True)
    quartier.set_categories(Reference(ws, min_col=5, min_row=64, max_row=71))
    quartier.height, quartier.width = 8, 13.2
    quartier.legend = None
    quartier.y_axis.number_format = '0,"k"'
    _axes_visibles(quartier)
    ws.add_chart(quartier, "E76")

    resultats = DoughnutChart()
    resultats.title = "Résultat des visites et contacts"
    resultats.add_data(Reference(ws, min_col=10, min_row=63, max_row=72), titles_from_data=True)
    resultats.set_categories(Reference(ws, min_col=9, min_row=64, max_row=72))
    resultats.height, resultats.width = 8, 13.2
    resultats.legend.position = "r"
    habiller_graphique(resultats)
    ws.add_chart(resultats, "J76")

    for col in "PQRSZ":
        _masquer(ws, col)
    ws.freeze_panes = "A6"
    mise_en_page(ws, "landscape", "A1:L92")
    ws.row_breaks.append(Break(id=27))
    ws.row_breaks.append(Break(id=61))
    proteger(ws)


def _entre12(nom: str, debut: str, fin: str) -> str:
    return f'{nom},">="&{debut},{nom},"<="&{fin}'


def _libelle_fusionne(ws, ligne: int, texte: str) -> None:
    """Libellé sur deux colonnes (F:G), pour les blocs de droite de la page 1."""
    cellule(ws, ligne, 6, texte, gras=True)
    ws.merge_cells(start_row=ligne, start_column=6, end_row=ligne, end_column=7)
    ws.cell(ligne, 7).border = copy(ws.cell(ligne, 6).border)
