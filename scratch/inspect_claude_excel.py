import openpyxl
import re
from collections import Counter

claude_path = r"c:\xampp\htdocs\app.gic-agropelc\BASE_CLIENTS_ET_PARTENAIRES_AFRIQUE_CENTRALE_ENRICHI_CLAUDE.xlsx"
gemini_path = r"c:\xampp\htdocs\app.gic-agropelc\BASE_CLIENTS_ET_PARTENAIRES_AFRIQUE_CENTRALE.xlsx"

print("=== INSPECTING CLAUDE EXCEL FILE (READ-ONLY) ===")
wb_claude = openpyxl.load_workbook(claude_path, read_only=True, data_only=True)
print("Claude sheet names:", wb_claude.sheetnames)

for name in wb_claude.sheetnames:
    ws = wb_claude[name]
    rows = list(ws.iter_rows(values_only=True))
    print(f"\nSheet '{name}': {len(rows)} total rows (including header)")
    if rows:
        print("  Headers:", list(rows[0][:25]))
        if len(rows) > 1:
            print("  Sample row 2:", list(rows[1][:16]))
            print("  Sample row 3:", list(rows[2][:16]))

ws_c = wb_claude["Base Contacts"]
c_rows = list(ws_c.iter_rows(values_only=True))
headers_c = list(c_rows[0])
data_c = c_rows[1:]

print(f"\nTotal data rows in Claude 'Base Contacts': {len(data_c)}")

# Analyze countries
countries = Counter(r[1] for r in data_c if len(r) > 1)
print("\nCountries in Claude file:")
for k, v in countries.most_common():
    print(f"  {k}: {v}")

# Analyze phone numbers (Primary vs Secondary vs Other)
def classify_phone(phone_str, country_str):
    if not phone_str or str(phone_str).strip() in ['', 'None', 'null']:
        return 'EMPTY'
    s = str(phone_str).strip()
    digits = re.sub(r'\D', '', s)
    if not digits:
        return 'EMPTY'
    c = str(country_str or '').lower()
    if 'cameroun' in c or digits.startswith('237'):
        rest = digits[3:] if digits.startswith('237') else digits
        if rest.startswith('2'):
            return 'FIXED_CAMTEL'
        elif rest.startswith('6'):
            return 'MOBILE_CMR'
        return 'OTHER_CMR'
    elif 'gabon' in c or digits.startswith('241'):
        rest = digits[3:] if digits.startswith('241') else digits
        if rest.startswith(('01', '11', '011')):
            return 'FIXED_GAB'
        elif rest.startswith(('06', '07', '6', '7', '04', '4', '02', '2', '05', '5')):
            return 'MOBILE_GAB'
        return 'OTHER_GAB'
    elif 'congo' in c or digits.startswith('242'):
        rest = digits[3:] if digits.startswith('242') else digits
        if rest.startswith('22'):
            return 'FIXED_COG'
        elif rest.startswith(('04', '05', '06')):
            return 'MOBILE_COG'
        return 'OTHER_COG'
    elif 'guin' in c or digits.startswith('240'):
        rest = digits[3:] if digits.startswith('240') else digits
        if rest.startswith(('333', '350', '33')):
            return 'FIXED_GNQ'
        elif rest.startswith(('222', '555', '666')):
            return 'MOBILE_GNQ'
        return 'OTHER_GNQ'
    elif 'tchad' in c or digits.startswith('235'):
        rest = digits[3:] if digits.startswith('235') else digits
        if rest.startswith('22'):
            return 'FIXED_TCH'
        elif rest.startswith(('6', '9', '7')):
            return 'MOBILE_TCH'
        return 'OTHER_TCH'
    elif 'centrafric' in c or digits.startswith('236'):
        rest = digits[3:] if digits.startswith('236') else digits
        if rest.startswith('21'):
            return 'FIXED_RCA'
        elif rest.startswith('7'):
            return 'MOBILE_RCA'
        return 'OTHER_RCA'
    return 'UNKNOWN'

pri_types = Counter()
fixed_with_mobile_sec = 0
fixed_without_mobile = []

for idx, r in enumerate(data_c, 2):
    country = r[1] if len(r) > 1 else ''
    name = r[4] if len(r) > 4 else ''
    p1 = r[9] if len(r) > 9 else ''
    p2 = r[10] if len(r) > 10 else ''
    p3 = r[11] if len(r) > 11 else ''
    
    t1 = classify_phone(p1, country)
    t2 = classify_phone(p2, country)
    t3 = classify_phone(p3, country)
    pri_types[t1] += 1
    
    if t1.startswith('FIXED'):
        if t2.startswith('MOBILE') or t3.startswith('MOBILE'):
            fixed_with_mobile_sec += 1
        else:
            fixed_without_mobile.append((idx, country, r[2], name, p1, p2, p3, r[12] if len(r) > 12 else '', r[15] if len(r) > 15 else ''))

print("\nPrimary Phone Classification in Claude file:")
for k, v in pri_types.most_common():
    print(f"  {k}: {v}")

print(f"\nRows with FIXED primary phone that ALREADY have a MOBILE in secondary/other: {fixed_with_mobile_sec}")
print(f"Rows with FIXED primary phone and NO mobile number yet: {len(fixed_without_mobile)}")

fixed_by_country = Counter(x[1] for x in fixed_without_mobile)
print("Fixed without mobile breakdown by country:")
for k, v in fixed_by_country.most_common():
    print(f"  {k}: {v}")

print("\nSample fixed without mobile (first 25):")
for item in fixed_without_mobile[:25]:
    print(f"  Row {item[0]} [{item[1]} - {item[2]}] {item[3]} | P1={item[4]} | P2={item[5]} | Email={item[7]} | Web={item[8]}")

# Check emails in Claude file
emails_c = sum(1 for r in data_c if len(r) > 12 and r[12] and '@' in str(r[12]))
emails2_c = sum(1 for r in data_c if len(r) > 13 and r[13] and '@' in str(r[13]))
print(f"\nEmails in Claude file: Primary={emails_c}, Secondary={emails2_c}")

wb_claude.close()

print("\n=== INSPECTING CURRENT GEMINI EXCEL FILE ===")
wb_gem = openpyxl.load_workbook(gemini_path, read_only=True, data_only=True)
ws_g = wb_gem["Base Contacts"]
g_rows = list(ws_g.iter_rows(values_only=True))
print(f"Total rows in Gemini 'Base Contacts': {len(g_rows) - 1}")
wb_gem.close()
