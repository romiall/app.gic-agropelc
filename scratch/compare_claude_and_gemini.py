import openpyxl
import re
from collections import Counter

claude_path = r"c:\xampp\htdocs\app.gic-agropelc\BASE_CLIENTS_ET_PARTENAIRES_AFRIQUE_CENTRALE_ENRICHI_CLAUDE.xlsx"
gemini_path = r"c:\xampp\htdocs\app.gic-agropelc\BASE_CLIENTS_ET_PARTENAIRES_AFRIQUE_CENTRALE.xlsx"

wb_c = openpyxl.load_workbook(claude_path, read_only=True, data_only=True)
print("Exact sheet names:", wb_c.sheetnames)
ws_c = wb_c["Base Contacts"]
c_rows = list(ws_c.iter_rows(values_only=True))
c_headers = list(c_rows[0])
c_data = c_rows[1:]

# Check Origine de la ligne (col 31, index 30)
origins = Counter(r[30] for r in c_data if len(r) > 30)
print("\nOrigine de la ligne in Claude:")
for k, v in origins.most_common():
    print(f"  {k}: {v}")

# Check Autres numéros confirmés (col 25, index 24)
other_confirmed_count = sum(1 for r in c_data if len(r) > 24 and r[24] and str(r[24]).strip() not in ['', 'None'])
print(f"\nRows with 'Autres numéros confirmés' (Col 25): {other_confirmed_count}")

def is_fixed(phone_str, country_str):
    if not phone_str:
        return False
    digits = re.sub(r'\D', '', str(phone_str))
    c = str(country_str or '').lower()
    if 'cameroun' in c or digits.startswith('237'):
        rest = digits[3:] if digits.startswith('237') else digits
        return rest.startswith('2')
    elif 'gabon' in c or digits.startswith('241'):
        rest = digits[3:] if digits.startswith('241') else digits
        return rest.startswith(('01', '11', '011'))
    elif 'congo' in c or digits.startswith('242'):
        rest = digits[3:] if digits.startswith('242') else digits
        return rest.startswith('22')
    elif 'guin' in c or digits.startswith('240'):
        rest = digits[3:] if digits.startswith('240') else digits
        return rest.startswith(('333', '350', '33'))
    elif 'tchad' in c or digits.startswith('235'):
        rest = digits[3:] if digits.startswith('235') else digits
        return rest.startswith('22')
    elif 'centrafric' in c or digits.startswith('236'):
        rest = digits[3:] if digits.startswith('236') else digits
        return rest.startswith('21')
    return False

def is_mobile(phone_str, country_str):
    if not phone_str or str(phone_str).strip() in ['', 'None']:
        return False
    digits = re.sub(r'\D', '', str(phone_str))
    if len(digits) < 8:
        return False
    return not is_fixed(phone_str, country_str)

url_domains = Counter()
fixed_only_rows = []
fixed_with_mobile_in_col25 = 0

for idx, r in enumerate(c_data, 2):
    country = r[1]
    p1, p2, p3 = r[9], r[10], r[11]
    other_phones = r[24] if len(r) > 24 else None
    if is_fixed(p1, country):
        if is_mobile(p2, country) or is_mobile(p3, country):
            pass
        elif other_phones and any(is_mobile(p.strip(), country) for p in re.split(r'[;|,/]', str(other_phones))):
            fixed_with_mobile_in_col25 += 1
        else:
            url = str(r[15] or r[27] or '')
            m = re.search(r'https?://([^/]+)', url)
            dom = m.group(1) if m else 'NO_URL'
            url_domains[dom] += 1
            fixed_only_rows.append((idx, country, r[2], r[4], p1, p2, p3, other_phones, url))

print(f"Fixed primary with mobile in Col 25: {fixed_with_mobile_in_col25}")
print(f"\nFixed-only rows URL domains ({len(fixed_only_rows)} total):")
for d, cnt in url_domains.most_common(20):
    print(f"  {d}: {cnt}")

# Compare with Gemini
wb_g = openpyxl.load_workbook(gemini_path, read_only=True, data_only=True)
ws_g = wb_g["Base Contacts"]
g_rows = list(ws_g.iter_rows(values_only=True))[1:]

print(f"\nGemini rows count: {len(g_rows)}")
gnq_g = [r for r in g_rows if 'guin' in str(r[1] or '').lower()]
gnq_c = [r for r in c_data if 'guin' in str(r[1] or '').lower()]
print(f"Equatorial Guinea in Gemini: {len(gnq_g)} vs in Claude: {len(gnq_c)}")

# How many rows in Gemini have mobile numbers or emails not in Claude?
c_phones = set()
for r in c_data:
    for col_i in [9, 10, 11]:
        if r[col_i]:
            d = re.sub(r'\D', '', str(r[col_i]))
            if d: c_phones.add(d)

g_unique_to_gemini = []
for r in g_rows:
    d = re.sub(r'\D', '', str(r[9] or ''))
    if d and d not in c_phones:
        g_unique_to_gemini.append(r)

print(f"Rows in Gemini whose primary phone is NOT in Claude: {len(g_unique_to_gemini)}")
g_unique_by_country = Counter(r[1] for r in g_unique_to_gemini)
for k, v in g_unique_by_country.most_common():
    print(f"  {k}: {v}")

wb_c.close()
wb_g.close()
