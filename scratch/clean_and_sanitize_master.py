import os
import re
import json
import shutil
import unicodedata
from collections import Counter, defaultdict
import openpyxl

BASE_DIR = r"c:\xampp\htdocs\app.gic-agropelc"
MASTER_EXCEL = os.path.join(BASE_DIR, "BASE_CLIENTS_ET_PARTENAIRES_AFRIQUE_CENTRALE.xlsx")
ENRICHI_EXCEL = os.path.join(BASE_DIR, "BASE_CLIENTS_ET_PARTENAIRES_AFRIQUE_CENTRALE_ENRICHI.xlsx")
CSV_OUTPUT = os.path.join(BASE_DIR, "scratch", "whatsapp_verifier", "numeros_a_tester.csv")
CHECKPOINT_PATH = os.path.join(BASE_DIR, "scratch", "whatsapp_verifier", "whatsapp_checkpoint.json")

# Blacklisted directory/portal numbers that leaked into multiple distinct companies
BLACK_PHONES = {
    '240376347860',  # Malabo directory web developer / footer
    '240555911787',  # Paginas Amarillas GE directory admin/support
    '240222633687',  # Directory support contact
    '240555841073',  # Directory medical contact
    '240555802997',  # Directory sports contact
    '240555891628',  # Directory beauty contact
    '237677128870',  # Petit Futé contact on unrelated hotels
    '241011448000',  # Radisson Blu fixed line incorrectly set on external restaurants
    '24107048989',   # Directory contact on 3 unrelated restaurants
    '24105928205',   # Directory contact on 3 unrelated restaurants
    '237696826609',  # Maligah / GoAfrica directory footer
    '237698545454',  # Maligah directory footer
    '237655223396',  # Directory footer
    '237675335252',  # Dummy portal
    '237655658012',  # Dummy portal
    '237699000000',  # Dummy sequence
    '237677000000',  # Dummy sequence
    '237222222222',  # Dummy sequence
    '237600000000',  # Dummy sequence
}

# Blacklisted dummy / directory emails
BLACK_EMAILS = {
    'info@salireneg.com',
    'guinealia@gmail.com',
    'contact@lepratiquedugabon.com',
    'adamaoua@armp.cm',
    'abuse@lws.fr',
}

BLACK_EMAIL_SUBSTRINGS = [
    'wght@', 'segmenter', 'core-js', '11.7.10', '3.2.1', 'npm', 'github.com',
    'example.com', 'exemple.com', 'domain.com', 'yoursite.com', 'sentry.io',
    'wixpress.com', 'wordpress.org', 'w3.org', 'schema.org'
]

def is_fixed_digits(d):
    if not d or len(d) < 8:
        return False
    if d.startswith('237'):
        return d[3:].startswith('2')
    if d.startswith('241'):
        return d[3:].startswith(('01', '11', '011'))
    if d.startswith('242'):
        return d[3:].startswith('22')
    if d.startswith('240'):
        return d[3:].startswith(('333', '350', '33'))
    if d.startswith('235'):
        return d[3:].startswith('22')
    if d.startswith('236'):
        return d[3:].startswith('21')
    return False

def is_valid_mobile_digits(d):
    if not d or len(d) < 8 or d in BLACK_PHONES:
        return False
    if d.startswith('237'):
        sub = d[3:]
        return len(sub) == 9 and sub[:2] in ('62', '65', '66', '67', '68', '69') and len(set(sub)) > 2
    if d.startswith('241'):
        sub = d[3:]
        return not sub.startswith(('01', '11', '011')) and len(sub) in (7, 8, 9)
    if d.startswith('242'):
        sub = d[3:]
        return len(sub) == 9 and sub.startswith(('04', '05', '06'))
    if d.startswith('240'):
        sub = d[3:]
        return len(sub) == 9 and sub.startswith(('222', '555', '666'))
    if d.startswith('235'):
        sub = d[3:]
        return len(sub) == 8 and sub[0] in ('6', '9', '7')
    if d.startswith('236'):
        sub = d[3:]
        return len(sub) == 8 and sub.startswith('7')
    return not is_fixed_digits(d)

def normalize_single_phone(raw, country=""):
    d = re.sub(r'\D', '', str(raw or ''))
    if not d:
        return ""
    c_low = str(country or '').lower()
    if d.startswith('00'):
        d = d[2:]
    if 'cameroun' in c_low and len(d) == 9 and d[0] in ('6', '2'):
        d = '237' + d
    elif 'gabon' in c_low and len(d) in (7, 8, 9) and not d.startswith('241'):
        d = '241' + d
    elif 'congo' in c_low and len(d) == 9 and d.startswith('0') and not d.startswith('242'):
        d = '242' + d
    elif 'tchad' in c_low and len(d) == 8 and not d.startswith('235'):
        d = '235' + d
    elif 'centraf' in c_low and len(d) == 8 and not d.startswith('236'):
        d = '236' + d
    elif 'guin' in c_low and len(d) == 9 and not d.startswith('240'):
        d = '240' + d
    return d

def split_and_extract_phones(raw_str, country=""):
    if not raw_str:
        return []
    s = str(raw_str)
    # Check for glued 24-digit or 23-digit strings (e.g. 237676...237695...)
    matches_237 = re.findall(r'237[26]\d{8}', s)
    matches_241 = re.findall(r'241(?:0[1-7]|11|[67])\d{6,7}', s)
    matches_240 = re.findall(r'240(?:222|555|666|333|350)\d{6}', s)
    matches_242 = re.findall(r'242(?:0[456]|22)\d{7}', s)
    matches_235 = re.findall(r'235(?:[697]|22)\d{7}', s)
    matches_236 = re.findall(r'236(?:7|21)\d{7}', s)

    combined_regex_matches = matches_237 + matches_241 + matches_240 + matches_242 + matches_235 + matches_236
    if combined_regex_matches:
        results = []
        for m in combined_regex_matches:
            if m not in BLACK_PHONES and m not in results:
                results.append(m)
        if results:
            return results

    # Standard split by delimiters
    parts = re.split(r'[/|;,]+', s)
    results = []
    for p in parts:
        if 'WA :' in p:
            continue
        d = normalize_single_phone(p, country)
        if len(d) >= 8 and d not in BLACK_PHONES and d not in results:
            results.append(d)
    return results

def is_clean_email(em):
    if not em or '@' not in str(em):
        return False
    em_low = str(em).lower().strip().strip('.')
    if ' ' in em_low or ';' in em_low or ',' in em_low:
        em_low = re.split(r'[\s;,|]+', em_low)[0]
    if em_low in BLACK_EMAILS:
        return False
    if any(sub in em_low for sub in BLACK_EMAIL_SUBSTRINGS):
        return False
    domain = em_low.split('@')[-1]
    if '.' not in domain or len(domain) < 4:
        return False
    # Check if domain has invalid characters or is a package version (e.g. 11.7.10)
    if re.search(r'^\d', domain) or '..' in domain:
        return False
    if any(em_low.endswith(ext) for ext in ('.png', '.jpg', '.jpeg', '.gif', '.svg', '.webp', '.js', '.css', '.pdf')):
        return False
    return True

def detect_operator(digits, country):
    d = str(digits)
    c_low = country.lower()
    if d.startswith('237'):
        sub = d[3:]
        if sub.startswith(('67', '650', '651', '652', '653', '654', '680', '681', '682', '683')):
            return 'MTN Cameroun'
        if sub.startswith(('69', '655', '656', '657', '658', '659', '685', '686', '687', '688', '689')):
            return 'Orange Cameroun'
        if sub.startswith('66'):
            return 'Nexttel'
        if sub.startswith('62'):
            return 'Camtel Blue (Mobile)'
        if sub.startswith('2'):
            return 'Camtel (Ligne Fixe)'
    elif d.startswith('241'):
        sub = d[3:]
        if sub.startswith(('06', '62', '66', '65')):
            return 'Moov Africa Gabon'
        if sub.startswith(('07', '74', '77')):
            return 'Airtel Gabon'
        if sub.startswith(('01', '11', '011')):
            return 'Gabon Telecom (Fixe)'
    elif d.startswith('242'):
        sub = d[3:]
        if sub.startswith(('06', '05')):
            return 'MTN Congo'
        if sub.startswith('04'):
            return 'Airtel Congo'
        if sub.startswith('22'):
            return 'Congo Telecom (Fixe)'
    elif d.startswith('240'):
        sub = d[3:]
        if sub.startswith(('222', '555')):
            return 'GETESA (GE)'
        if sub.startswith('666'):
            return 'Muni (GE)'
        if sub.startswith(('333', '350', '33')):
            return 'GETESA Fixe (GE)'
    elif d.startswith('235'):
        sub = d[3:]
        if sub.startswith(('6', '9')):
            return 'Airtel Tchad'
        if sub.startswith('7'):
            return 'Moov Tchad'
        if sub.startswith('22'):
            return 'Sotel Tchad (Fixe)'
    elif d.startswith('236'):
        sub = d[3:]
        if sub.startswith('70'):
            return 'Telecel Centrafrique'
        if sub.startswith('72'):
            return 'Orange Centrafrique'
        if sub.startswith(('75', '77')):
            return 'Moov Centrafrique'
        if sub.startswith('21'):
            return 'Socatel (Fixe RCA)'
    return 'Autre / Inconnu'

def main():
    print("=================================================================")
    print("ASSAINISSEMENT RIGOUREUX DU CLASSEUR MAITRE & GENERATION CSV WA")
    print("=================================================================")

    wb = openpyxl.load_workbook(MASTER_EXCEL)
    ws = wb["Base Contacts"]

    stats = Counter()
    all_unique_phones = {} # digits -> {pays, type, operateur, priorite, entreprise, nb}

    max_r = ws.max_row
    max_c = ws.max_column
    has_col25 = max_c >= 25
    print(f"Dimensions: {max_r} lignes, {max_c} colonnes. Démarrage de l'analyse...", flush=True)

    # First pass: clean each row in Base Contacts
    for row_idx in range(2, max_r + 1):
        if row_idx % 1000 == 0:
            print(f"   [Progression] Traitement ligne {row_idx}/{max_r}...", flush=True)

        rec_id = str(ws.cell(row=row_idx, column=1).value or "").strip()
        country = str(ws.cell(row=row_idx, column=2).value or "").strip()
        city = str(ws.cell(row=row_idx, column=3).value or "").strip()
        name = str(ws.cell(row=row_idx, column=5).value or "").strip()

        p1_raw = ws.cell(row=row_idx, column=10).value
        p2_raw = ws.cell(row=row_idx, column=11).value
        wa_raw = ws.cell(row=row_idx, column=12).value
        em_raw = ws.cell(row=row_idx, column=13).value
        c25_raw = ws.cell(row=row_idx, column=25).value if has_col25 else ""

        # Extract all numbers from all phone columns
        extracted = []
        for src in (p1_raw, p2_raw, c25_raw):
            for d in split_and_extract_phones(src, country):
                if d and d not in extracted:
                    extracted.append(d)

        # Separate into valid mobiles and valid fixed
        mobiles = [d for d in extracted if is_valid_mobile_digits(d)]
        fixed = [d for d in extracted if is_fixed_digits(d)]

        # Sort mobiles by operator priority for Cameroon (MTN & Orange first)
        if 'cameroun' in country.lower() and len(mobiles) > 1:
            def cm_prio(d):
                sub = d[3:] if d.startswith('237') else d
                if sub.startswith(('67', '69', '65', '68')):
                    return 0
                return 1
            mobiles.sort(key=cm_prio)

        # Determine Primary, Secondary, and Others
        if mobiles:
            primary_val = f"+{mobiles[0]}"
            # Secondary: 2nd mobile if available, else 1st fixed
            if len(mobiles) > 1:
                secondary_val = f"+{mobiles[1]}"
                remaining_mobs = mobiles[2:]
                remaining_fixed = fixed
            elif fixed:
                secondary_val = f"+{fixed[0]}"
                remaining_mobs = []
                remaining_fixed = fixed[1:]
            else:
                secondary_val = ""
                remaining_mobs = []
                remaining_fixed = []

            # Column 25 (Others): remaining mobiles + remaining fixed
            others = [f"+{m}" for m in remaining_mobs] + [f"+{f}" for f in remaining_fixed]
            col25_val = "; ".join(others)
            stats["has_mobile_primary"] += 1
        elif fixed:
            primary_val = f"+{fixed[0]}"
            secondary_val = f"+{fixed[1]}" if len(fixed) > 1 else ""
            others = [f"+{f}" for f in fixed[2:]]
            col25_val = "; ".join(others)
            stats["fixed_only"] += 1
        else:
            primary_val = ""
            secondary_val = ""
            col25_val = ""

        # Update cells in Excel
        ws.cell(row=row_idx, column=10).value = primary_val
        ws.cell(row=row_idx, column=11).value = secondary_val
        if has_col25:
            ws.cell(row=row_idx, column=25).value = col25_val

        # Clean WhatsApp Column (Col 12)
        cur_wa = str(wa_raw or "").strip()
        # Clean any raw phone strings or slashes from WA column
        if cur_wa:
            wa_parts = [p.strip() for p in re.split(r'[/|]+', cur_wa) if p.strip()]
            status_parts = [p for p in wa_parts if p.startswith('WA :')]
            if status_parts:
                ws.cell(row=row_idx, column=12).value = " | ".join(status_parts)
            else:
                # If only raw phone was in WA column, leave it empty or clearly marked
                ws.cell(row=row_idx, column=12).value = ""
        else:
            ws.cell(row=row_idx, column=12).value = ""

        # Clean Email Column (Col 13)
        cur_em = str(em_raw or "").strip()
        if cur_em:
            clean_em = ""
            for part in re.split(r'[\s;,/]+', cur_em):
                p = part.strip().strip('.').lower()
                if is_clean_email(p):
                    clean_em = p
                    break
            ws.cell(row=row_idx, column=13).value = clean_em
            if clean_em:
                stats["has_email"] += 1
            if clean_em != cur_em:
                stats["email_purged"] += 1

        # Register all valid numbers into all_unique_phones for CSV generation
        for d in mobiles + fixed:
            is_mob = is_valid_mobile_digits(d)
            op = detect_operator(d, country)
            if is_mob:
                if 'MTN' in op or 'Orange' in op or 'Airtel' in op or 'Moov' in op or 'GETESA' in op:
                    prio = 1
                else:
                    prio = 2
                t_label = 'mobile'
            else:
                prio = 3
                t_label = 'fixe'

            if d not in all_unique_phones:
                all_unique_phones[d] = {
                    "numero": d,
                    "pays": country,
                    "type": t_label,
                    "operateur": op,
                    "priorite": prio,
                    "entreprise": name,
                    "nb_entreprises": 1
                }
            else:
                all_unique_phones[d]["nb_entreprises"] += 1

    wb.save(MASTER_EXCEL)
    wb.close()
    shutil.copy2(MASTER_EXCEL, ENRICHI_EXCEL)

    print(f"Classeur Excel mis à jour : {MASTER_EXCEL}")
    print(f"- Total entreprises traitées : {max_r - 1}")
    print(f"- Entreprises avec MOBILE en Principal : {stats['has_mobile_primary']} ({stats['has_mobile_primary']*100/(max_r - 1):.1f}%)")
    print(f"- Entreprises restant uniquement fixes : {stats['fixed_only']}")
    print(f"- Entreprises avec e-mail professionnel vérifié : {stats['has_email']}")
    print(f"- E-mails pollués nettoyés / supprimés : {stats['email_purged']}")

    # Write numeros_a_tester.csv for verifier_whatsapp_complet.js
    sorted_numbers = sorted(all_unique_phones.values(), key=lambda x: (x["priorite"], -x["nb_entreprises"]))

    os.makedirs(os.path.dirname(CSV_OUTPUT), exist_ok=True)
    with open(CSV_OUTPUT, "w", encoding="utf-8-sig") as f:
        f.write("numero;pays;type;operateur;priorite;entreprise;nb_entreprises\n")
        for n in sorted_numbers:
            f.write(f"{n['numero']};{n['pays']};{n['type']};{n['operateur']};{n['priorite']};{n['entreprise']};{n['nb_entreprises']}\n")

    mob_count = sum(1 for n in sorted_numbers if n["type"] == "mobile")
    fixe_count = sum(1 for n in sorted_numbers if n["type"] == "fixe")
    prio1_count = sum(1 for n in sorted_numbers if n["priorite"] == 1)

    print("\n=================================================================")
    print(f"FICHIER CSV POUR VERIFICATION WHATSAPP GENERE : {CSV_OUTPUT}")
    print(f"- Total numéros uniques documentés : {len(sorted_numbers)}")
    print(f"- Numéros mobiles à tester : {mob_count} (dont {prio1_count} en priorité 1)")
    print(f"- Numéros fixes (écartables via --sans-fixes) : {fixe_count}")
    print("=================================================================")

if __name__ == "__main__":
    main()
