import json
import os
import re
import time
import shutil
import openpyxl

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
VERIFIER_DIR = os.path.dirname(__file__)
CHECKPOINT_PATH = os.path.join(VERIFIER_DIR, "whatsapp_checkpoint.json")
EXCEL_MAIN = os.path.join(BASE_DIR, "BASE_CLIENTS_ET_PARTENAIRES_AFRIQUE_CENTRALE.xlsx")
EXCEL_ENRICHI = os.path.join(BASE_DIR, "BASE_CLIENTS_ET_PARTENAIRES_AFRIQUE_CENTRALE_ENRICHI.xlsx")
LOCK_FILE = os.path.join(VERIFIER_DIR, ".excel_sync.lock")

def acquire_lock(timeout_sec=30):
    start = time.time()
    while time.time() - start < timeout_sec:
        try:
            # Atomic file creation in exclusive mode
            fd = os.open(LOCK_FILE, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.write(fd, str(os.getpid()).encode('utf-8'))
            os.close(fd)
            return True
        except FileExistsError:
            time.sleep(1.5)
        except Exception:
            time.sleep(1.5)
    return False

def release_lock():
    try:
        if os.path.exists(LOCK_FILE):
            os.remove(LOCK_FILE)
    except Exception:
        pass

def main():
    if not os.path.exists(EXCEL_MAIN):
        print(f"[ERREUR] Fichier Excel principal introuvable : {EXCEL_MAIN}")
        return

    # 1. Discover and load ALL resultats_whatsapp*.jsonl files
    jsonl_files = []
    if os.path.exists(VERIFIER_DIR):
        for f in os.listdir(VERIFIER_DIR):
            if f.startswith("resultats_whatsapp") and f.endswith(".jsonl"):
                jsonl_files.append(os.path.join(VERIFIER_DIR, f))

    tested_numbers = {} # digits -> 'oui' | 'non'
    for jf in jsonl_files:
        try:
            with open(jf, "r", encoding="utf-8") as f:
                for line in f:
                    line = line.strip()
                    if not line:
                        continue
                    try:
                        obj = json.loads(line)
                        if "numero" in obj and "statut" in obj:
                            num = re.sub(r'\D', '', str(obj["numero"]))
                            if num:
                                tested_numbers[num] = obj["statut"]
                    except Exception:
                        pass
        except Exception as e:
            print(f"[WARN] Erreur lecture {os.path.basename(jf)}: {e}")

    # 2. Load results from whatsapp_checkpoint.json
    checkpoint = {}
    if os.path.exists(CHECKPOINT_PATH):
        try:
            with open(CHECKPOINT_PATH, "r", encoding="utf-8") as f:
                checkpoint = json.load(f)
        except Exception:
            checkpoint = {}

    if not tested_numbers and not checkpoint:
        print("[INFO] Aucun résultat WhatsApp à synchroniser pour le moment.")
        return

    print(f"[SYNC] Consolidation de {len(tested_numbers)} numéros vérifiés depuis {len(jsonl_files)} journal(aux) JSONL et {len(checkpoint)} entrées checkpoint...")

    # Acquire lock before touching the Excel files to prevent concurrent write collisions
    if not acquire_lock():
        print("[ATTENTION] Une autre synchronisation Excel est en cours. Nouvel essai au prochain lot.")
        return

    try:
        wb = openpyxl.load_workbook(EXCEL_MAIN)
        if "Base Contacts" not in wb.sheetnames:
            wb.close()
            return
        ws = wb["Base Contacts"]
        max_r = ws.max_row
        max_c = ws.max_column
        has_col25 = max_c >= 25

        updated_count = 0
        promoted_count = 0

        for row_idx in range(2, max_r + 1):
            rec_id = str(ws.cell(row=row_idx, column=1).value or "").strip()
            p1 = str(ws.cell(row=row_idx, column=10).value or "").strip()
            p2 = str(ws.cell(row=row_idx, column=11).value or "").strip()
            cur_wa = str(ws.cell(row=row_idx, column=12).value or "").strip()
            c25 = str(ws.cell(row=row_idx, column=25).value or "").strip() if has_col25 else ""

            p1_d = re.sub(r'\D', '', p1)
            p2_d = re.sub(r'\D', '', p2)

            matched_status = None
            promoted_p1 = None
            demoted_p2 = None

            # Priority 1: Checkpoint data
            key = rec_id if rec_id in checkpoint else p1
            if key in checkpoint:
                info = checkpoint[key]
                matched_status = info.get("status_str", "")
                promoted_p1 = info.get("promoted_primary")
                demoted_p2 = info.get("demoted_secondary")

            # Priority 2: Consolidated JSONL data from all parallel workers
            if not matched_status:
                if p1_d in tested_numbers:
                    st = tested_numbers[p1_d]
                    if st == 'oui':
                        matched_status = f"WA : OUI (+{p1_d})"
                    else:
                        if p2_d in tested_numbers and tested_numbers[p2_d] == 'oui':
                            matched_status = f"WA : OUI (+{p2_d})"
                            promoted_p1 = f"+{p2_d}"
                            demoted_p2 = f"+{p1_d}"
                        elif p2_d and p2_d in tested_numbers and tested_numbers[p2_d] == 'non':
                            matched_status = "WA : NON"
                        elif not p2_d:
                            matched_status = "WA : NON"
                elif p2_d in tested_numbers and tested_numbers[p2_d] == 'oui':
                    matched_status = f"WA : OUI (+{p2_d})"
                    promoted_p1 = f"+{p2_d}"
                    demoted_p2 = f"+{p1_d}"

            if matched_status and matched_status not in cur_wa:
                parts = [p.strip() for p in cur_wa.split("|") if p.strip() and not p.strip().startswith("WA :")]
                parts.append(matched_status)
                ws.cell(row=row_idx, column=12).value = " | ".join(parts)
                updated_count += 1

            if promoted_p1 and demoted_p2 and ws.cell(row=row_idx, column=10).value != promoted_p1:
                ws.cell(row=row_idx, column=10).value = promoted_p1
                ws.cell(row=row_idx, column=11).value = demoted_p2
                promoted_count += 1

        # Attempt to save with retry if user currently has Excel open
        saved = False
        for attempt in range(1, 4):
            try:
                wb.save(EXCEL_MAIN)
                wb.close()
                saved = True
                break
            except PermissionError:
                print(f"[REESSAI {attempt}/3] Le classeur Excel semble ouvert dans une autre application. En attente de fermeture (3s)...")
                time.sleep(3)
            except Exception as e:
                print(f"[ERREUR SAUVEGARDE] {e}")
                break

        if saved:
            try:
                shutil.copy2(EXCEL_MAIN, ENRICHI_EXCEL)
            except Exception:
                pass
            print(f"[SYNC EXCEL] Succès : {updated_count} statuts WhatsApp synchronisés dans le classeur maître.")
            if promoted_count > 0:
                print(f"[PROMOTION] {promoted_count} numéros secondaires actifs sur WhatsApp promus en Téléphone Principal !")
        else:
            print("[ATTENTION] Sauvegarde Excel reportée : veuillez fermer Excel s'il est ouvert sur votre écran.")

    finally:
        release_lock()

if __name__ == "__main__":
    main()
