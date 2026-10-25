import os
import re
import json
import shutil
import unicodedata
from copy import copy
from collections import Counter, defaultdict
import openpyxl

BASE_DIR = r"c:\xampp\htdocs\app.gic-agropelc"
CLAUDE_PATH = os.path.join(BASE_DIR, "BASE_CLIENTS_ET_PARTENAIRES_AFRIQUE_CENTRALE_ENRICHI_CLAUDE.xlsx")
GEMINI_OLD_PATH = os.path.join(BASE_DIR, "BASE_CLIENTS_ET_PARTENAIRES_AFRIQUE_CENTRALE.xlsx")
CHATGPT_PATH = os.path.join(BASE_DIR, "BASE_CLIENTS_ET_PARTENAIRES_AFRIQUE_CENTRALE_VERSION_CHATGPT.xlsx")
TARGET_MAIN = os.path.join(BASE_DIR, "BASE_CLIENTS_ET_PARTENAIRES_AFRIQUE_CENTRALE.xlsx")
TARGET_ENRICHI = os.path.join(BASE_DIR, "BASE_CLIENTS_ET_PARTENAIRES_AFRIQUE_CENTRALE_ENRICHI.xlsx")
CHECKPOINT_PATH = os.path.join(BASE_DIR, "scratch", "whatsapp_verifier", "whatsapp_checkpoint.json")

PORTAL_PHONES = {
    '237675335252', '237655658012', '237699000000', '237677000000',
    '237222222222', '237600000000'
}

PORTAL_DOMAINS = (
    'maligah.com', 'goafricaonline.com', 'chadyp.com', 'armp.cm',
    'example.com', 'exemple.cm', 'exemple.com', 'domain.com', 'yoursite.com', 'email.com',
    'sentry.io', 'wixpress.com', 'wordpress.org', 'w3.org', 'schema.org',
    'lws.fr', 'ovh.net', 'ovh.com', 'domotics.studio', 'iea-tech.net',
    'hostinger.com', 'namecheap.com', 'ionos.fr', 'ionos.com', 'gandi.net',
    'petitfute.com', 'tripadvisor.com', 'booking.com'
)

def norm_name(s):
    if not s:
        return ""
    s = unicodedata.normalize('NFKD', str(s)).encode('ASCII', 'ignore').decode('utf-8').lower()
    s = re.sub(r'\b(sarl|sa|sas|ets|etablissement|societe|group|groupe|ltd|inc|gic|cooperative|hotel|restaurant|pharmacie|complexe|centre|agence|direction|generale|siege|social|cameroun|gabon|congo|tchad)\b', ' ', s)
    return re.sub(r'[^a-z0-9]', '', s)

def is_fixed_digits(d):
    if not d or len(d) < 8:
        return False
    if d.startswith('237'):
        return d[3:].startswith('2')
    if d.startswith('241'):
        return d[3:].startswith(('01', '11', '011'))
    if d.startswith('242'):
        return d[3:].startswith('22')
    if d.startswith('240'):
        return d[3:].startswith(('333', '350', '33'))
    if d.startswith('235'):
        return d[3:].startswith('22')
    if d.startswith('236'):
        return d[3:].startswith('21')
    return False

def is_valid_mobile_digits(d):
    if not d or len(d) < 8 or d in PORTAL_PHONES:
        return False
    if d.startswith('237'):
        sub = d[3:]
        return len(sub) == 9 and sub[:2] in ('62', '65', '66', '67', '68', '69') and len(set(sub)) > 2
    if d.startswith('241'):
        sub = d[3:]
        return not sub.startswith(('01', '11', '011')) and len(sub) in (7, 8, 9)
    if d.startswith('242'):
        sub = d[3:]
        return len(sub) == 9 and sub.startswith(('04', '05', '06'))
    if d.startswith('240'):
        sub = d[3:]
        return len(sub) == 9 and sub.startswith(('222', '555', '666'))
    if d.startswith('235'):
        sub = d[3:]
        return len(sub) == 8 and sub[0] in ('6', '9', '7')
    if d.startswith('236'):
        sub = d[3:]
        return len(sub) == 8 and sub.startswith('7')
    return not is_fixed_digits(d)

def normalize_phone_digits(raw, country=""):
    d = re.sub(r'\D', '', str(raw or ''))
    if not d:
        return ""
    c_low = str(country or '').lower()
    if d.startswith('00'):
        d = d[2:]
    if 'cameroun' in c_low and len(d) == 9 and d[0] in ('6', '2'):
        d = '237' + d
    elif 'gabon' in c_low and len(d) in (7, 8, 9) and not d.startswith('241'):
        d = '241' + d
    elif 'congo' in c_low and len(d) == 9 and d.startswith('0') and not d.startswith('242'):
        d = '242' + d
    elif 'tchad' in c_low and len(d) == 8 and not d.startswith('235'):
        d = '235' + d
    elif 'centraf' in c_low and len(d) == 8 and not d.startswith('236'):
        d = '236' + d
    elif 'guin' in c_low and len(d) == 9 and not d.startswith('240'):
        d = '240' + d
    return d

def is_clean_email(em):
    if not em or '@' not in str(em):
        return False
    em_low = str(em).lower().strip().strip('.')
    if ' ' in em_low or ';' in em_low or ',' in em_low:
        em_low = re.split(r'[\s;,|]+', em_low)[0]
    domain = em_low.split('@')[-1]
    if '.' not in domain or len(domain) < 4:
        return False
    if any(d in domain for d in PORTAL_DOMAINS):
        return False
    if any(em_low.endswith(ext) for ext in ('.png', '.jpg', '.jpeg', '.gif', '.svg', '.webp', '.js', '.css', '.pdf')):
        return False
    return True

def clean_email_str(em):
    if not em:
        return ""
    for part in re.split(r'[\s;,|/]+', str(em)):
        p = part.strip().strip('.').lower()
        if is_clean_email(p):
            return p
    return ""

def main():
    print("1. Lecture des données existantes (Gemini, ChatGPT, Claude) avant fusion...")
    # Collect from current Gemini file:
    # - 40 WhatsApp test statuses
    # - Verified emails
    # - Net-new Equatorial Guinea rows
    wa_statuses_by_phone = {}
    gemini_emails_by_phone = {}
    gemini_emails_by_name = {}
    gemini_mobiles_by_name = {}
    gemini_gq_rows = []

    if os.path.exists(GEMINI_OLD_PATH):
        wb_g = openpyxl.load_workbook(GEMINI_OLD_PATH, read_only=True, data_only=True)
        if "Base Contacts" in wb_g.sheetnames:
            for r in list(wb_g["Base Contacts"].iter_rows(values_only=True))[1:]:
                if not r or not r[0]:
                    continue
                country = str(r[1] or '').strip()
                city = str(r[2] or '').strip().lower()
                name = str(r[4] or '').strip()
                n_key = (norm_name(name), city)
                p1_d = normalize_phone_digits(r[9], country)
                p2_d = normalize_phone_digits(r[10], country)
                wa_val = str(r[11] or '').strip()
                em_val = clean_email_str(r[12])

                if 'WA : OUI' in wa_val or 'WA : NON' in wa_val:
                    if p1_d:
                        wa_statuses_by_phone[p1_d] = wa_val

                if em_val:
                    if p1_d:
                        gemini_emails_by_phone[p1_d] = em_val
                    if n_key[0]:
                        gemini_emails_by_name[n_key] = em_val
                        gemini_emails_by_name[(n_key[0], '')] = em_val

                for d in (p1_d, p2_d):
                    if is_valid_mobile_digits(d) and n_key[0]:
                        gemini_mobiles_by_name.setdefault(n_key, []).append(d)
                        gemini_mobiles_by_name.setdefault((n_key[0], ''), []).append(d)

                if 'guin' in country.lower():
                    gemini_gq_rows.append(r)
        wb_g.close()

    # Collect from ChatGPT file
    if os.path.exists(CHATGPT_PATH):
        wb_cg = openpyxl.load_workbook(CHATGPT_PATH, read_only=True, data_only=True)
        for sname in wb_cg.sheetnames:
            if "Contact" in sname or "Base" in sname:
                for r in list(wb_cg[sname].iter_rows(values_only=True))[1:]:
                    if not r or len(r) < 13:
                        continue
                    country = str(r[1] or '').strip()
                    city = str(r[2] or '').strip().lower()
                    n_key = (norm_name(r[4]), city)
                    p1_d = normalize_phone_digits(r[9], country)
                    p2_d = normalize_phone_digits(r[10], country)
                    em_val = clean_email_str(r[12])
                    if em_val:
                        if p1_d and p1_d not in gemini_emails_by_phone:
                            gemini_emails_by_phone[p1_d] = em_val
                        if n_key[0] and n_key not in gemini_emails_by_name:
                            gemini_emails_by_name[n_key] = em_val
                    for d in (p1_d, p2_d):
                        if is_valid_mobile_digits(d) and n_key[0]:
                            gemini_mobiles_by_name.setdefault(n_key, []).append(d)
                            gemini_mobiles_by_name.setdefault((n_key[0], ''), []).append(d)
        wb_cg.close()

    # Load Claude's auxiliary sheets for cross-matching
    wb_c_ro = openpyxl.load_workbook(CLAUDE_PATH, read_only=True, data_only=True)
    for aux_sheet, col_country, col_city, col_name, col_p1, col_p2, col_em in [
        ("Cibles secondaires", 1, 2, 4, 9, 10, 12),
        ("Numéros partagés (à arbitrer)", 2, 3, 5, 10, 11, 13),
        ("Pistes Facebook (sans tél.)", 1, 2, 4, None, None, 8),
    ]:
        if aux_sheet in wb_c_ro.sheetnames:
            for r in list(wb_c_ro[aux_sheet].iter_rows(values_only=True))[1:]:
                if not r:
                    continue
                country = str(r[col_country] or '').strip() if col_country is not None and len(r) > col_country else ''
                city = str(r[col_city] or '').strip().lower() if col_city is not None and len(r) > col_city else ''
                name = str(r[col_name] or '').strip() if col_name is not None and len(r) > col_name else ''
                n_key = (norm_name(name), city)
                if not n_key[0]:
                    continue
                if col_p1 is not None and len(r) > col_p1:
                    d1 = normalize_phone_digits(r[col_p1], country)
                    if is_valid_mobile_digits(d1):
                        gemini_mobiles_by_name.setdefault(n_key, []).append(d1)
                        gemini_mobiles_by_name.setdefault((n_key[0], ''), []).append(d1)
                if col_p2 is not None and len(r) > col_p2:
                    d2 = normalize_phone_digits(r[col_p2], country)
                    if is_valid_mobile_digits(d2):
                        gemini_mobiles_by_name.setdefault(n_key, []).append(d2)
                        gemini_mobiles_by_name.setdefault((n_key[0], ''), []).append(d2)
                if col_em is not None and len(r) > col_em:
                    em_val = clean_email_str(r[col_em])
                    if em_val:
                        gemini_emails_by_name.setdefault(n_key, em_val)
                        gemini_emails_by_name.setdefault((n_key[0], ''), em_val)
    wb_c_ro.close()

    # Load JSON enrichments (corporate websites, directory pages, subagents)
    row_enrichments = {}   # row_idx -> {"mobiles": [], "emails": []}
    id_enrichments = {}    # record_id -> {"mobiles": [], "emails": []}
    name_enrichments = {}  # norm_name -> {"mobiles": [], "emails": []}

    json_files = [
        os.path.join(BASE_DIR, "scratch", "corporate_websites_enriched.json"),
        os.path.join(BASE_DIR, "scratch", "directory_fixed_enriched.json"),
        os.path.join(BASE_DIR, "scratch", "subagent_enrichment_1.json"),
        os.path.join(BASE_DIR, "scratch", "subagent_enrichment_2.json"),
        os.path.join(BASE_DIR, "scratch", "subagent_enrichment_3.json"),
        os.path.join(BASE_DIR, "scratch", "subagent_enrichment_4.json"),
    ]
    for jf in json_files:
        if os.path.exists(jf):
            try:
                with open(jf, "r", encoding="utf-8") as f:
                    items = json.load(f)
                print(f"   Chargé {len(items)} enrichissements depuis {os.path.basename(jf)}")
                # Pre-calculate frequency by distinct company name for directory scrapes
                is_dir_scrape = "directory_fixed_enriched" in jf
                mob_name_freq = defaultdict(set)
                em_name_freq = defaultdict(set)
                if is_dir_scrape:
                    for it in items:
                        nm = norm_name(it.get("name"))
                        country = it.get("country", "")
                        for m in it.get("mobiles", []):
                            d = normalize_phone_digits(m, country)
                            if d: mob_name_freq[d].add(nm)
                        for e in it.get("emails", []):
                            ce = clean_email_str(e)
                            if ce: em_name_freq[ce].add(nm)

                for it in items:
                    r_idx = it.get("row")
                    rec_id = str(it.get("id") or "").strip()
                    nm = norm_name(it.get("name"))
                    country = it.get("country", "")

                    mobs = []
                    if "mobiles" in it and isinstance(it["mobiles"], list):
                        for m in it["mobiles"]:
                            d = normalize_phone_digits(m, country)
                            if is_valid_mobile_digits(d):
                                if not is_dir_scrape or len(mob_name_freq[d]) <= 1:
                                    mobs.append(d)
                    for k_m in ("mobile_phone", "secondary_mobile"):
                        if it.get(k_m):
                            for part in re.split(r'[|;,/]+', str(it[k_m])):
                                d = normalize_phone_digits(part, country)
                                if is_valid_mobile_digits(d):
                                    mobs.append(d)

                    ems = []
                    if "emails" in it and isinstance(it["emails"], list):
                        for e in it["emails"]:
                            ce = clean_email_str(e)
                            if ce:
                                if not is_dir_scrape or len(em_name_freq[ce]) <= 1:
                                    ems.append(ce)
                    if it.get("email"):
                        ce = clean_email_str(it["email"])
                        if ce:
                            ems.append(ce)

                    if r_idx:
                        entry = row_enrichments.setdefault(int(r_idx), {"mobiles": [], "emails": []})
                        for m in mobs:
                            if m not in entry["mobiles"]: entry["mobiles"].append(m)
                        for e in ems:
                            if e not in entry["emails"]: entry["emails"].append(e)
                    if rec_id:
                        entry = id_enrichments.setdefault(rec_id, {"mobiles": [], "emails": []})
                        for m in mobs:
                            if m not in entry["mobiles"]: entry["mobiles"].append(m)
                        for e in ems:
                            if e not in entry["emails"]: entry["emails"].append(e)
                    if nm:
                        entry = name_enrichments.setdefault(nm, {"mobiles": [], "emails": []})
                        for m in mobs:
                            if m not in entry["mobiles"]: entry["mobiles"].append(m)
                        for e in ems:
                            if e not in entry["emails"]: entry["emails"].append(e)
            except Exception as e:
                print(f"   [WARN] Erreur lecture {jf}: {e}")

    print("2. Copie intégrale du classeur Claude vers le classeur maître Gemini...")
    shutil.copy2(CLAUDE_PATH, TARGET_MAIN)

    wb = openpyxl.load_workbook(TARGET_MAIN)
    ws = wb["Base Contacts"]

    existing_phones = set()
    existing_names = set()

    stats = Counter()
    checkpoint_data = {}
    if os.path.exists(CHECKPOINT_PATH):
        try:
            with open(CHECKPOINT_PATH, "r", encoding="utf-8") as f:
                checkpoint_data = json.load(f)
        except Exception:
            checkpoint_data = {}

    print(f"3. Enrichissement et inversion Fixe->Secondaire / Mobile->Principal sur les {ws.max_row - 1} lignes de Claude...")
    for row_idx in range(2, ws.max_row + 1):
        rec_id = str(ws.cell(row=row_idx, column=1).value or "").strip()
        country = str(ws.cell(row=row_idx, column=2).value or "").strip()
        city = str(ws.cell(row=row_idx, column=3).value or "").strip().lower()
        name = str(ws.cell(row=row_idx, column=5).value or "").strip()
        n_key = (norm_name(name), city)

        p1_raw = str(ws.cell(row=row_idx, column=10).value or "").strip()
        p2_raw = str(ws.cell(row=row_idx, column=11).value or "").strip()
        wa_raw = str(ws.cell(row=row_idx, column=12).value or "").strip()
        em_raw = str(ws.cell(row=row_idx, column=13).value or "").strip()
        other_nums_raw = str(ws.cell(row=row_idx, column=25).value or "").strip()

        p1_d = normalize_phone_digits(p1_raw, country)
        if p1_d:
            existing_phones.add(p1_d)
        if n_key[0]:
            existing_names.add((country.lower(), n_key[0]))

        # Gather all known numbers for this row
        all_mobiles = []
        all_fixed = []

        for src_str in (p1_raw, p2_raw, wa_raw, other_nums_raw):
            for part in re.split(r'[|;,/]+', src_str):
                if 'WA :' in part:
                    continue
                d = normalize_phone_digits(part, country)
                if not d or len(d) < 8:
                    continue
                if is_valid_mobile_digits(d):
                    if d not in all_mobiles:
                        all_mobiles.append(d)
                elif is_fixed_digits(d):
                    if d not in all_fixed:
                        all_fixed.append(d)

        # Check enrichments for this row
        for enr_dict, key in [
            (row_enrichments, row_idx),
            (id_enrichments, rec_id),
            (name_enrichments, n_key[0])
        ]:
            if key and key in enr_dict:
                for m in enr_dict[key]["mobiles"]:
                    if m not in all_mobiles:
                        all_mobiles.append(m)

        # If still no mobile and primary is fixed, check local cross-sheet matches by name
        if not all_mobiles and is_fixed_digits(p1_d):
            for cand in gemini_mobiles_by_name.get(n_key, []) + gemini_mobiles_by_name.get((n_key[0], ''), []):
                if is_valid_mobile_digits(cand) and cand not in all_mobiles:
                    all_mobiles.append(cand)

        # Prioritize MTN (67, 650-654, 680-683) and Orange (69, 655-659, 686-689) over Camtel Blue (62) or Nexttel (66) in Cameroon
        if len(all_mobiles) > 1 and 'cameroun' in country.lower():
            def cm_priority(d):
                sub = d[3:] if d.startswith('237') else d
                if sub.startswith(('67', '69', '65', '68')):
                    return 0
                return 1
            all_mobiles.sort(key=cm_priority)

        # Apply the swap/enrichment rule!
        if is_fixed_digits(p1_d):
            if all_mobiles:
                best_mobile = f"+{all_mobiles[0]}"
                ws.cell(row=row_idx, column=10).value = best_mobile
                existing_phones.add(all_mobiles[0])

                # Build secondary phone list: remaining mobiles + the old fixed primary
                sec_parts = []
                for m in all_mobiles[1:]:
                    sec_parts.append(f"+{m}")
                for fx in all_fixed:
                    sec_parts.append(f"+{fx}")
                ws.cell(row=row_idx, column=11).value = " / ".join(sec_parts[:3])
                stats["fixed_replaced_by_mobile"] += 1
            else:
                # Only fixed number exists: keep in Principal AND copy to Secondaire if requested, or keep clearly marked
                if not p2_raw:
                    ws.cell(row=row_idx, column=11).value = p1_raw
                stats["still_fixed_only"] += 1
        else:
            # Primary is already mobile! Ensure secondary has any extra discovered mobile or fixed number
            if not p2_raw and (len(all_mobiles) > 1 or all_fixed):
                sec_parts = [f"+{m}" for m in all_mobiles[1:]] + [f"+{fx}" for fx in all_fixed]
                if sec_parts:
                    ws.cell(row=row_idx, column=11).value = " / ".join(sec_parts[:2])
                    stats["secondary_enriched"] += 1

        # Email enrichment & cleanup
        cur_email = clean_email_str(em_raw)
        if not cur_email:
            # Look in row_enrichments, id_enrichments, name_enrichments, gemini_emails
            for enr_dict, key in [
                (row_enrichments, row_idx),
                (id_enrichments, rec_id),
                (name_enrichments, n_key[0])
            ]:
                if key and key in enr_dict and enr_dict[key]["emails"]:
                    cur_email = enr_dict[key]["emails"][0]
                    break
            if not cur_email and p1_d in gemini_emails_by_phone:
                cur_email = gemini_emails_by_phone[p1_d]
            if not cur_email and n_key in gemini_emails_by_name:
                cur_email = gemini_emails_by_name[n_key]
            if not cur_email and (n_key[0], '') in gemini_emails_by_name:
                cur_email = gemini_emails_by_name[(n_key[0], '')]

            if cur_email:
                ws.cell(row=row_idx, column=13).value = cur_email
                stats["emails_enriched"] += 1
        elif cur_email != em_raw:
            ws.cell(row=row_idx, column=13).value = cur_email

        if cur_email:
            stats["total_with_email"] += 1

        # Restore WhatsApp status if already tested
        final_p1_d = normalize_phone_digits(ws.cell(row=row_idx, column=10).value, country)
        wa_status = wa_statuses_by_phone.get(final_p1_d) or wa_statuses_by_phone.get(p1_d)
        if wa_status and 'WA :' not in str(ws.cell(row=row_idx, column=12).value or ''):
            cur_w = str(ws.cell(row=row_idx, column=12).value or '').strip()
            ws.cell(row=row_idx, column=12).value = f"{cur_w} | {wa_status}" if cur_w else wa_status
            checkpoint_data[rec_id or f"+{final_p1_d}"] = {
                "status_str": wa_status,
                "checked_at": "2026-10-15T23:00:00Z"
            }
            stats["wa_restored"] += 1

    # 4. Append net-new verified Equatorial Guinea rows from Gemini
    print("4. Intégration des nouvelles entreprises vérifiées de Guinée équatoriale (absentes de la version Claude)...")
    template_row_idx = ws.max_row
    template_styles = [copy(ws.cell(row=template_row_idx, column=c)._style) for c in range(1, 33)]

    gq_added = 0
    for r in gemini_gq_rows:
        country = str(r[1] or 'Guinée équatoriale').strip()
        name = str(r[4] or '').strip()
        n_clean = norm_name(name)
        p1_d = normalize_phone_digits(r[9], country)
        p2_d = normalize_phone_digits(r[10], country)

        if not p1_d or p1_d in existing_phones or (country.lower(), n_clean) in existing_names:
            continue

        # Check if p1_d is fixed and p2_d is mobile
        p1_str = str(r[9] or '').strip()
        p2_str = str(r[10] or '').strip()
        if is_fixed_digits(p1_d) and is_valid_mobile_digits(p2_d):
            p1_str, p2_str = f"+{p2_d}", f"+{p1_d}"
            p1_d, p2_d = p2_d, p1_d
        elif is_fixed_digits(p1_d) and n_clean in name_enrichments and name_enrichments[n_clean]["mobiles"]:
            mob_d = name_enrichments[n_clean]["mobiles"][0]
            p2_str = p1_str
            p1_str = f"+{mob_d}"
            p1_d = mob_d

        new_row_idx = ws.max_row + 1
        for c in range(1, 24):
            val = r[c - 1] if c - 1 < len(r) else ""
            cell = ws.cell(row=new_row_idx, column=c, value=val)
            cell._style = copy(template_styles[c - 1])

        ws.cell(row=new_row_idx, column=10).value = p1_str
        ws.cell(row=new_row_idx, column=11).value = p2_str
        em_clean = clean_email_str(r[12] if len(r) > 12 else "")
        if not em_clean and n_clean in name_enrichments and name_enrichments[n_clean]["emails"]:
            em_clean = name_enrichments[n_clean]["emails"][0]
        ws.cell(row=new_row_idx, column=13).value = em_clean
        if em_clean:
            stats["total_with_email"] += 1

        # Fill audit columns 24-32
        extra_vals = {
            24: r[15] if len(r) > 15 and 'facebook.com' in str(r[15] or '').lower() else "",
            25: p2_str,
            26: "Confirmé (annuaire national GE vérifié)",
            27: 1,
            28: r[15] if len(r) > 15 else "",
            29: "2026-10-16",
            30: "Ajouté depuis l'extension Guinée équatoriale vérifiée (EmpresasGuinea / Guinealia / SalirenEG)",
            31: "Extension GE Gemini vérifiée",
            32: r[0]
        }
        for c, val in extra_vals.items():
            cell = ws.cell(row=new_row_idx, column=c, value=val)
            cell._style = copy(template_styles[c - 1])

        existing_phones.add(p1_d)
        existing_names.add((country.lower(), n_clean))
        gq_added += 1

    print(f"   +{gq_added} nouvelles entreprises vérifiées de Guinée équatoriale ajoutées !")

    # Count final mobile vs fixed in Base Contacts
    final_total = ws.max_row - 1
    final_mobile = 0
    final_fixed = 0
    for row_idx in range(2, ws.max_row + 1):
        c_try = str(ws.cell(row=row_idx, column=2).value or "")
        d = normalize_phone_digits(ws.cell(row=row_idx, column=10).value, c_try)
        if is_fixed_digits(d):
            final_fixed += 1
        else:
            final_mobile += 1

    wb.save(TARGET_MAIN)
    wb.close()
    shutil.copy2(TARGET_MAIN, TARGET_ENRICHI)

    os.makedirs(os.path.dirname(CHECKPOINT_PATH), exist_ok=True)
    with open(CHECKPOINT_PATH, "w", encoding="utf-8") as f:
        json.dump(checkpoint_data, f, ensure_ascii=False, indent=2)

    print("\n=================================================================")
    print("BILAN DE LA FUSION ET DE L'ENRICHISSEMENT :")
    print(f"- Total entreprises dans Base Contacts : {final_total}")
    print(f"- Entreprises avec Numéro MOBILE en Principal : {final_mobile} ({final_mobile*100/final_total:.1f}%)")
    print(f"- Lignes fixes remplacées par un Mobile en Principal : {stats['fixed_replaced_by_mobile']}")
    print(f"- Lignes restant uniquement fixes (sans mobile public) : {final_fixed}")
    print(f"- Nouvelles adresses e-mail enrichies : +{stats['emails_enriched']}")
    print(f"- Total entreprises avec adresse e-mail vérifiée : {stats['total_with_email']}")
    print(f"- Statuts WhatsApp restaurés : {stats['wa_restored']}")
    print("=================================================================")

if __name__ == "__main__":
    main()
