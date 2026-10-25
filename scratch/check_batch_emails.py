import openpyxl
import glob
import os
import re

wb_paths = [
    r"c:\xampp\htdocs\app.gic-agropelc\BASE_CLIENTS_ET_PARTENAIRES_AFRIQUE_CENTRALE.xlsx",
    r"c:\xampp\htdocs\app.gic-agropelc\BASE_CLIENTS_ET_PARTENAIRES_AFRIQUE_CENTRALE_ENRICHI.xlsx"
]

wb = openpyxl.load_workbook(wb_paths[0], data_only=True)
ws = wb["Base Contacts"]

# Map name to row
name_to_row = {}
phone_to_row = {}
for r in range(2, ws.max_row + 1):
    name = str(ws.cell(r, 5).value or '').strip().lower()
    phone = re.sub(r'\D', '', str(ws.cell(r, 10).value or ''))
    name_to_row[name] = r
    if phone:
        phone_to_row[phone] = r

# Find all email matches in batch files
batch_files = glob.glob(r"C:\Users\SMART\.gemini\antigravity\brain\bf4e43c6-b734-4bea-9d78-496e7525f4e8\scratch\batch_*.py")
new_found = []
for bf in batch_files:
    with open(bf, "r", encoding="utf-8") as f:
        content = f.read()
        # Find dictionaries with regex
        blocks = re.findall(r'\{[^{}]+?\}', content)
        for b in blocks:
            name_m = re.search(r'"display_name":\s*"([^"]+)"', b)
            phone_m = re.search(r'"phone_primary":\s*"([^"]+)"', b)
            email_m = re.search(r'"email":\s*"([^"]+)"', b)
            phone2_m = re.search(r'"phone_secondary":\s*"([^"]+)"', b)
            if email_m:
                em = email_m.group(1).strip()
                if em and '@' in em and em != 'None':
                    dname = name_m.group(1).strip() if name_m else ''
                    ph = phone_m.group(1).strip() if phone_m else ''
                    ph2 = phone2_m.group(1).strip() if phone2_m else ''
                    new_found.append((dname, ph, em, ph2))

print(f"Extracted {len(new_found)} contacts with emails from batch files.")

# Enrich workbooks
for p in wb_paths:
    wb = openpyxl.load_workbook(p)
    ws = wb["Base Contacts"]
    enriched = 0
    for dname, ph, em, ph2 in new_found:
        digits = re.sub(r'\D', '', ph)
        target_r = name_to_row.get(dname.lower()) or phone_to_row.get(digits)
        if target_r:
            curr_em = ws.cell(target_r, 13).value
            curr_ph2 = ws.cell(target_r, 11).value
            if not curr_em or '@' not in str(curr_em):
                ws.cell(target_r, 13, value=em)
                enriched += 1
            if ph2 and ph2 != 'None' and not curr_ph2:
                ws.cell(target_r, 11, value=ph2)
                
    tot_em = sum(1 for row in range(2, ws.max_row + 1) if ws.cell(row, 13).value and '@' in str(ws.cell(row, 13).value))
    tot_rows = ws.max_row - 1
    ws_dash = wb["Synthèse & Tableau de bord"]
    ws_dash.cell(row=17, column=2, value=f"{tot_em} ({(tot_em/tot_rows)*100:.1f}%)")
    wb.save(p)
    print(f"{p}: enriched {enriched} additional email fields. Grand total emails: {tot_em}")
