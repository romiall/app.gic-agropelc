import openpyxl
import re
import json
import urllib.request
import ssl
from concurrent.futures import ThreadPoolExecutor, as_completed
from collections import defaultdict, Counter

claude_path = r"c:\xampp\htdocs\app.gic-agropelc\BASE_CLIENTS_ET_PARTENAIRES_AFRIQUE_CENTRALE_ENRICHI_CLAUDE.xlsx"
gemini_path = r"c:\xampp\htdocs\app.gic-agropelc\BASE_CLIENTS_ET_PARTENAIRES_AFRIQUE_CENTRALE.xlsx"
gpt_path = r"c:\xampp\htdocs\app.gic-agropelc\BASE_CLIENTS_ET_PARTENAIRES_AFRIQUE_CENTRALE_VERSION_CHATGPT.xlsx"

def is_valid_mobile(digits, country_str=''):
    c = str(country_str or '').lower()
    if digits.startswith('237') or 'cameroun' in c:
        rest = digits[3:] if digits.startswith('237') else digits
        if len(rest) == 9 and rest[:2] in ('62', '65', '66', '67', '68', '69'):
            return '237' + rest
        return None
    if digits.startswith('241') or 'gabon' in c:
        rest = digits[3:] if digits.startswith('241') else digits
        if len(rest) in (7, 8) and not rest.startswith(('01', '11', '011')):
            if rest[0] in ('6', '7', '2', '4', '5') or rest[:2] in ('06', '07', '02', '04', '05', '60', '62', '65', '66', '74', '76', '77'):
                return '241' + rest
        return None
    if digits.startswith('242') or 'congo' in c:
        rest = digits[3:] if digits.startswith('242') else digits
        if len(rest) == 9 and rest[:2] in ('04', '05', '06'):
            return '242' + rest
        return None
    if digits.startswith('240') or 'guin' in c:
        rest = digits[3:] if digits.startswith('240') else digits
        if len(rest) == 9 and rest[:3] in ('222', '555', '666'):
            return '240' + rest
        return None
    if digits.startswith('235') or 'tchad' in c:
        rest = digits[3:] if digits.startswith('235') else digits
        if len(rest) == 8 and rest[0] in ('6', '9'):
            return '235' + rest
        return None
    if digits.startswith('236') or 'centrafric' in c:
        rest = digits[3:] if digits.startswith('236') else digits
        if len(rest) == 8 and rest[0] == '7':
            return '236' + rest
        return None
    return None

def is_fixed_phone(digits, country_str=''):
    if not digits:
        return False
    return is_valid_mobile(digits, country_str) is None

def is_clean_email(em):
    if not em or '@' not in str(em):
        return False
    em = str(em).strip().lower().strip('.')
    if not re.match(r'^[a-z0-9_.+-]+@[a-z0-9-]+\.[a-z]{2,10}$', em):
        return False
    bad_patterns = [
        '.png', '.jpg', '.jpeg', '.gif', '.svg', '.webp', '.js', '.css', '.woff', '.ttf',
        'example.', 'exemple.', 'sentry.', 'wixpress.', 'wordpress.', 'domain.', 'email.',
        'yoursite.', 'maligah.com', 'goafricaonline.com', 'armp.cm', 'chadyp.com',
        'wght@', 'react@', 'lodash@', 'bootstrap@', 'jquery@', 'vue@'
    ]
    return not any(b in em for b in bad_patterns)

def norm_name(s):
    s = str(s or '').lower()
    s = re.sub(r'\([^)]*\)', ' ', s)
    for w in ['sarl', 'sa', 'snc', 'ltd', 'ets', 'group', 'groupe', 'de', 'du', 'des', 'la', 'le', 'les', 'cameroun', 'cameroon', 'gabon', 'congo', 'tchad', 'siege', 'siège', 'social', 'direction', 'generale', 'générale', 'agence']:
        s = re.sub(r'\b' + w + r'\b', ' ', s)
    return re.sub(r'[\W_]+', '', s)

wb_c = openpyxl.load_workbook(claude_path, read_only=True, data_only=True)
c_rows = list(wb_c["Base Contacts"].iter_rows(values_only=True))[1:]
sec_rows = list(wb_c["Cibles secondaires"].iter_rows(values_only=True))[1:]
shared_rows = list(wb_c["Numéros partagés (à arbitrer)"].iter_rows(values_only=True))[1:]
fb_rows = list(wb_c["Pistes Facebook (sans tél.)"].iter_rows(values_only=True))[1:]
wb_c.close()

wb_g = openpyxl.load_workbook(gemini_path, read_only=True, data_only=True)
g_rows = list(wb_g["Base Contacts"].iter_rows(values_only=True))[1:]
wb_g.close()

wb_gpt = openpyxl.load_workbook(gpt_path, read_only=True, data_only=True)
gpt_rows = list(wb_gpt["Base Contacts"].iter_rows(values_only=True))[1:]
gpt_audit = list(wb_gpt["Audit Gemini"].iter_rows(values_only=True))[1:]
wb_gpt.close()

# Build lookup of valid mobiles and clean emails by normalized name + city
mobiles_by_key = defaultdict(list)
emails_by_key = defaultdict(list)
mobiles_by_name = defaultdict(list)
emails_by_name = defaultdict(list)

for ds in [c_rows, sec_rows, shared_rows, g_rows, gpt_rows, gpt_audit]:
    for r in ds:
        if not r or len(r) < 13: continue
        country = str(r[1] or '')
        city = re.sub(r'[\W_]+', '', str(r[2] or '').lower())
        n = norm_name(r[4])
        if len(n) < 4: continue
        for idx_p in [9, 10, 11]:
            if idx_p < len(r) and r[idx_p]:
                for part in re.split(r'[;|,/]', str(r[idx_p])):
                    d = re.sub(r'\D', '', part)
                    mob = is_valid_mobile(d, country)
                    if mob:
                        if mob not in mobiles_by_key[(n, city)]:
                            mobiles_by_key[(n, city)].append(mob)
                        if mob not in mobiles_by_name[n]:
                            mobiles_by_name[n].append(mob)
        for idx_e in [12, 13]:
            if idx_e < len(r) and r[idx_e]:
                for part in re.split(r'[;|,/\s]', str(r[idx_e])):
                    if is_clean_email(part):
                        em = part.strip().lower()
                        if em not in emails_by_key[(n, city)]:
                            emails_by_key[(n, city)].append(em)
                        if em not in emails_by_name[n]:
                            emails_by_name[n].append(em)

# Check how many fixed-only rows remain after internal swap + cross-match
still_fixed = []
swapped_internal = 0
matched_cross = 0

for idx, r in enumerate(c_rows, 2):
    country = str(r[1] or '')
    city = str(r[2] or '')
    name = str(r[4] or '')
    cat = str(r[6] or '')
    d1 = re.sub(r'\D', '', str(r[9] or ''))
    if is_valid_mobile(d1, country):
        continue
    # Fixed primary! Check p2, p3, col 25
    found_mob = None
    for col_i in [10, 11, 24]:
        if col_i < len(r) and r[col_i]:
            for part in re.split(r'[;|,/]', str(r[col_i])):
                m = is_valid_mobile(re.sub(r'\D', '', part), country)
                if m:
                    found_mob = m
                    break
        if found_mob:
            break
    if found_mob:
        swapped_internal += 1
        continue
    
    # Check cross-match by (name, city)
    n = norm_name(name)
    c_norm = re.sub(r'[\W_]+', '', city.lower())
    if mobiles_by_key.get((n, c_norm)):
        matched_cross += 1
        continue
    
    still_fixed.append({
        'row': idx,
        'id': r[0],
        'country': country,
        'city': city,
        'quarter': r[3],
        'name': name,
        'category': cat,
        'activity': r[7],
        'fixed_phone': r[9],
        'email': r[12],
        'address': r[14],
        'url': r[15]
    })

print(f"Swapped from secondary/other column: {swapped_internal}")
print(f"Matched by (name, city) across local sheets: {matched_cross}")
print(f"Remaining fixed-only rows to enrich: {len(still_fixed)}")

by_cat = Counter(x['category'] for x in still_fixed)
print("Remaining fixed-only by category:", by_cat)
by_country = Counter(x['country'] for x in still_fixed)
print("Remaining fixed-only by country:", by_country)

with open(r"c:\xampp\htdocs\app.gic-agropelc\scratch\still_fixed_to_enrich.json", "w", encoding="utf-8") as f:
    json.dump(still_fixed, f, ensure_ascii=False, indent=2)
print("Saved scratch/still_fixed_to_enrich.json")
