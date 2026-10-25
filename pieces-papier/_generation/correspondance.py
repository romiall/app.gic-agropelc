"""Génère CORRESPONDANCE-APPLICATION.xlsx : pour chaque pièce papier, les champs et référentiels de l'application.
Sources lues dans le dépôt (lecture seule) : db/seeds/*.ts, apps/server/src/modules/*/application/commands/*.ts,
docs/03-data/dictionnaire/*.md, docs/01-functional/01-glossaire.md, docs/A-VALIDER.md.
Lancer : python -I pieces-papier/_generation/correspondance.py"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

SORTIE = Path(__file__).resolve().parents[1] / "CORRESPONDANCE-APPLICATION.xlsx"
VERT = "3A6351"

PIECES = [
    ("BV", "Bon de vente", "Enregistrer une vente (commande `sales.sale.record`)", "sales.sales, sales.sale_lines, sales.customer_payments", "VTE", "vérifié (code)"),
    ("RE", "Reçu d'encaissement", "Enregistrer un encaissement affecté à une vente ou une commande", "sales.customer_payments, sales.payment_allocations", "ENC", "vérifié (code)"),
    ("BD", "Bon de décaissement", "Dépense (finance.expenses) ; paiement fournisseur, avance, remboursement client = autres commandes de la finance", "finance.expenses ; finance.supplier_payments", "DEP / PF", "vérifié (dictionnaire, code de la dépense)"),
    ("RF", "Remise de fonds", "Remise de fonds entre deux comptes de trésorerie (envoi puis réception)", "finance.cash_transfers", "RMF", "vérifié (dictionnaire)"),
    ("JE", "Journal des encaissements", "Aucune écriture : état reconstitué depuis les encaissements enregistrés", "lecture des encaissements", "—", "déduit"),
    ("FA", "Facture (modèle provisoire)", "AUCUNE : l'application n'émet pas de facture client (AV-041, décision du 30/09/2026 : prix TTC, reçu numérique simple)", "—", "—", "vérifié (A-VALIDER AV-041)"),
    ("DA", "Demande d'achat", "Soumettre une demande d'achat", "procurement.purchase_requests", "DA", "vérifié (code)"),
    ("BC", "Bon de commande fournisseur", "Bon de commande fournisseur (approbation au-delà d'un seuil)", "procurement.purchase_orders", "BC", "vérifié (glossaire, seuil)"),
    ("BR", "Bon de réception", "Enregistrer une réception (livré, rejeté, accepté)", "procurement.goods_receipts", "REC", "vérifié (code)"),
    ("BL", "Bon de livraison", "Livraison d'une commande client (partielle possible)", "sales.delivery_notes", "LIV", "vérifié (code)"),
    ("BT", "Bon de transfert de stock", "Transfert en deux temps : expédition puis réception", "inventory.stock_transfers", "TRF", "vérifié (dictionnaire)"),
    ("FP", "Fiche de perte / consommation", "Déclaration de perte ; consommation d'intrant", "inventory.loss_declarations ; inventory.consumptions", "PRT", "vérifié (dictionnaire)"),
    ("FI", "Fiche d'inventaire", "Comptage physique d'un emplacement", "inventory.inventory_counts", "INV", "vérifié (dictionnaire, code)"),
    ("FL", "Suivi de lot de production", "Saisie du jour d'un lot : mortalité, aliment, pesée, collecte d'œufs", "inventory.loss_declarations (MORTALITE), inventory.consumptions, production.lot_weighings, production.egg_collections", "LOT / —", "vérifié (dictionnaire, code)"),
    ("BH", "Bilan hebdomadaire du commercial", "AUCUNE aujourd'hui : nouvelle entité proposée (PROPOSITION-INTEGRATION-APPLICATION.md, P-07)", "(à créer)", "—", "proposition"),
    ("CC", "Bon de commande client", "Commande client (avec acompte éventuel)", "sales.sales_orders", "CMD", "vérifié (code)"),
    ("VC", "Fiche de visite commerciale", "Visite d'un compte client", "crm.visits", "—", "vérifié (dictionnaire, référentiel)"),
    ("PT", "Fiche de présence et de pointage", "Session de travail terrain (prise et fin de service)", "fieldwork.work_sessions", "—", "vérifié (dictionnaire)"),
    ("RC", "Registre de suivi des carnets", "AUCUNE : registre interne papier", "—", "—", "—"),
    ("EL / BN / CV", "Papier à en-tête, blocs-notes, couvertures", "AUCUNE : papeterie", "—", "—", "—"),
]
PIECES_COLS = ["Code", "Pièce", "Opération de l'application", "Table(s)", "Préfixe du numéro officiel", "Vérification"]

MOY = "ESPECES (Espèces) · MOBILE_MONEY_ORANGE (Orange Money) · MOBILE_MONEY_MTN (MTN Mobile Money) · VIREMENT (Virement bancaire) · CHEQUE (Chèque)"
UNITES = "TETE, OEUF, PIECE, KG, G, L, SAC, PLATEAU, CARTON, DOSE"
CHAMPS = [
    # (pièce, rubrique papier, saisie, obligatoire, champ cible, référentiel, source)
    ("BV", "Numéro imprimé de la pièce", "recopié", "non", "localRef (20 car. max)", "—", "sale-payload.ts"),
    ("BV", "Point de vente", "texte", "oui", "fromLocationId (emplacement de vente)", "emplacements de l'organisation", "sale-payload.ts"),
    ("BV", "Client / « client anonyme »", "texte / case", "non", "customerId (absent = vente anonyme)", "comptes clients", "sale-payload.ts ; glossaire « vente anonyme »"),
    ("BV", "Désignation", "texte", "oui", "lines[].productId", "catalogue produits", "sale-payload.ts"),
    ("BV", "Unité", "texte", "oui", "lines[].unitCode", UNITES, "sale-payload.ts ; 03-catalog-pricing.md"),
    ("BV", "Quantité", "nombre", "oui", "lines[].quantity (+ quantityBase)", "—", "sale-payload.ts"),
    ("BV", "Poids (kg)", "nombre", "si vente au poids", "lines[].weightKg", "mode de tarification PER_WEIGHT", "sale-payload.ts"),
    ("BV", "Prix unitaire (XAF, TTC)", "montant", "oui", "lines[].unitPriceXaf", "prix en TTC (AV-041)", "sale-payload.ts ; AV-041"),
    ("BV", "Remise (XAF)", "montant", "non", "lines[].discountXaf (+ overrideReasonCodeId si le prix diffère du prix résolu)", "motifs de dérogation de prix", "sale-payload.ts"),
    ("BV", "Mode de paiement", "case", "non", "payments[].methodCode", MOY, "finance-references.ts"),
    ("BV", "Montant encaissé", "montant", "non", "payments[].amountXaf (> 0)", "—", "sale-payload.ts"),
    ("BV", "Réf. transaction / chèque", "texte", "oui hors espèces", "payments[].reference (80 car. max)", "requiresReference = vrai sauf ESPECES", "finance-references.ts"),
    ("BV", "Reste dû (crédit)", "montant", "calculé", "créance = ventes − affectations", "—", "glossaire « créance »"),
    ("BV", "Date et heure", "date-heure", "oui", "occurred_at (heure métier, Africa/Douala)", "—", "ADR-016"),
    ("RE", "Reçu de", "texte", "non", "customerId", "comptes clients", "payment-commands.ts"),
    ("RE", "Point de vente / site", "texte", "non", "siteId", "sites", "payment-commands.ts"),
    ("RE", "Montant reçu", "montant", "oui", "amountXaf (> 0)", "—", "payment-commands.ts"),
    ("RE", "Mode de paiement", "case", "oui", "methodCode", MOY, "finance-references.ts"),
    ("RE", "Réf. transaction / n° du chèque", "texte", "oui hors espèces", "reference (80 car. max)", "—", "payment-commands.ts"),
    ("RE", "Numéro imprimé", "recopié", "non", "localRef (20 car. max)", "—", "payment-commands.ts"),
    ("RE", "Pour : règlement de la vente n° / acompte sur la commande n°", "cases + numéro", "non", "allocations[] : saleId ou orderId (exactement un) + amountXaf", "—", "payment-commands.ts"),
    ("RE", "Reste dû après ce paiement", "montant", "calculé", "—", "—", "—"),
    ("RE", "Compte crédité", "(non écrit sur la pièce)", "non", "cashAccountId (défaut selon le moyen : espèces vers la caisse de l'encaisseur)", "types : CAISSE_PDV, CAISSE_UTILISATEUR, CAISSE_CENTRALE, MOBILE_MONEY, BANQUE", "finance-references.ts"),
    ("BD", "Payé à (bénéficiaire)", "texte", "non", "payee_name", "—", "09-finance.md (finance.expenses)"),
    ("BD", "Catégorie", "case", "oui", "category_id", "ALIMENT_HORS_STOCK, VETERINAIRE, TRANSPORT, CARBURANT, ENERGIE_EAU, MAIN_OEUVRE_OCCASIONNELLE, ENTRETIEN, FOURNITURES, DIVERS (+ NON_JUSTIFIEE, catégorie système)", "finance-references.ts"),
    ("BD", "Objet détaillé", "texte", "oui", "description", "—", "09-finance.md"),
    ("BD", "Imputé à (lot, site, PDV)", "texte", "non", "cost_object_type, cost_object_id (+ species_group)", "lot de production, site, PDV", "09-finance.md"),
    ("BD", "Facture / bon de commande n°", "texte", "non", "supplier_id ; rattachement à une facture fournisseur", "fournisseurs", "09-finance.md"),
    ("BD", "Montant payé", "montant", "oui", "amount_xaf", "—", "09-finance.md"),
    ("BD", "Justificatifs joints", "cases", "selon seuil", "pièces jointes, propriétaire EXPENSE (natures RECEIPT, INVOICE)", "seuil : finance.expense_photo_threshold_xaf", "ADR-046, ADR-047"),
    ("BD", "Autorisé par (visa)", "signature", "selon seuil", "demande de validation de type EXPENSE", "seuil : finance.expense_approval_threshold_xaf ; plafond PDV : finance.pos_expense_max_xaf", "system-settings.ts"),
    ("BD", "Nature : paiement fournisseur / avance fournisseur / remboursement client", "case", "—", "autres commandes de la finance (paiement fournisseur, avance, remboursement)", "—", "ADR-043"),
    ("FJC", "Caisse : du PDV / du commercial / centrale", "case", "oui", "cash_account_id (types CAISSE_PDV, CAISSE_UTILISATEUR, CAISSE_CENTRALE)", "BR-FIN-010", "schéma ; D09"),
    ("FJC", "Jour", "date", "oui", "business_date du mouvement ou de la session", "—", "finance.cash_movements"),
    ("FJC", "Session n°", "numéro", "PDV seulement", "cash_session_id (la caisse d'un commercial n'a pas de session)", "—", "glossaire « session de caisse »"),
    ("FJC", "Fonds d'ouverture compté", "montant", "PDV", "opening_counted_xaf", "—", "finance.cash_sessions"),
    ("FJC", "Ligne : heure", "heure", "oui", "occurred_at du mouvement", "—", "finance.cash_movements"),
    ("FJC", "Ligne : n° de la pièce", "recopié", "non", "doc_number ou local_ref du document source (source_doc_type, source_doc_id)", "—", "finance.cash_movements"),
    ("FJC", "Ligne : client, bénéficiaire ou objet", "texte", "non", "libellé tiré du document source (encaissement, dépense, remise de fonds)", "—", "—"),
    ("FJC", "Ligne : mode (E O M V C)", "code", "oui", "payment_method_code du document source", MOY, "finance-references.ts"),
    ("FJC", "Ligne : ENTRÉE / SORTIE", "montant", "oui", "direction + amount_xaf du mouvement", "—", "finance.cash_movements"),
    ("FJC", "Espèces comptées", "montant", "PDV", "closing_counted_xaf", "—", "finance.cash_sessions"),
    ("FJC", "Solde théorique des espèces", "montant", "calculé", "closing_expected_xaf", "—", "finance.cash_sessions"),
    ("FJC", "ÉCART (compté − théorique)", "montant", "calculé", "variance_xaf", "tolérance : finance.cash_variance_tolerance_xaf (0)", "system-settings.ts"),
    ("FJC", "Motif de l'écart", "texte", "si écart", "variance_reason ; validation Finance (finance.cash_session.validate)", "—", "finance-references.ts"),
    ("FJC", "Comptage par coupures", "(supprimé)", "—", "retiré de la feuille sur décision de la Direction", "—", "—"),
    ("RF", "Montant envoyé / reçu / écart", "montants", "oui", "amount_sent_xaf / amount_received_xaf / discrepancy_xaf", "—", "09-finance.md"),
    ("RF", "De / Vers", "cases", "oui", "comptes source et destinataire", "paires admises : finance.cash_transfer_allowed_pairs", "system-settings.ts"),
    ("RF", "Motif de l'écart", "texte", "si écart", "décision d'écart (politique CASH_TRANSFER_DISCREPANCY)", "alerte si non reçue : 48 h", "finance-references.ts"),
    ("DA", "Site demandeur", "texte", "oui", "siteId", "sites", "request-commands.ts"),
    ("DA", "Besoin attendu pour le", "date", "non", "neededByDate", "—", "request-commands.ts"),
    ("DA", "Justification du besoin", "texte", "OUI (1 à 4000 car.)", "justification", "—", "request-commands.ts"),
    ("DA", "Produit, unité, quantité, prix estimé, observations", "lignes", "oui", "lines[] : productId, unitCode, quantity, estimatedUnitPriceXaf, notes", UNITES, "request-commands.ts"),
    ("DA", "Numéro imprimé", "recopié", "non", "localRef", "—", "request-commands.ts"),
    ("BC", "Fournisseur ; livraison attendue ; lieu de livraison ; lignes ; prix", "—", "oui", "bon de commande fournisseur (purchase_orders et lignes)", "approbation au-delà de procurement.po_approval_threshold_xaf", "system-settings.ts ; glossaire"),
    ("BR", "Bon de commande n°", "texte", "non", "purchaseOrderId (absent = réception sans BC)", "—", "receipt-commands.ts"),
    ("BR", "Fournisseur", "texte", "oui", "supplierId", "fournisseurs", "receipt-commands.ts"),
    ("BR", "Site / emplacement de réception", "texte", "oui", "locationId", "emplacements physiques", "receipt-commands.ts"),
    ("BR", "Bon de livraison du fournisseur n°", "texte", "non", "supplierDeliveryNoteRef (60 car. max)", "—", "receipt-commands.ts"),
    ("BR", "Livré / Rejeté / Accepté", "nombres", "oui", "lines[].qtyDeliveredBase / qtyRejectedBase ; accepté = livré − rejeté", "—", "receipt-commands.ts ; glossaire"),
    ("BR", "Motif du rejet", "texte", "si rejet", "lines[].rejectionReasonCodeId", "motifs de rejet", "receipt-commands.ts"),
    ("BR", "N° de lot fournisseur ; péremption", "texte ; date", "non", "lines[].supplierLotRef (60) ; lines[].expiryDate", "—", "receipt-commands.ts"),
    ("BR", "Observations", "texte", "non", "observations (4000 car. max)", "—", "receipt-commands.ts"),
    ("BR", "Photo du bon de livraison", "case", "selon seuil", "pièce jointe GOODS_RECEIPT, nature DELIVERY_NOTE (attachment_ids)", "exigée sans BC ou à partir de procurement.receipt_photo_threshold_xaf (100 000 XAF)", "ADR-046 ; system-settings.ts"),
    ("BR", "Numéro imprimé", "recopié", "non", "localRef", "—", "receipt-commands.ts"),
    ("BL", "Commande client n°", "texte", "oui", "orderId", "commandes clients", "delivery-commands.ts"),
    ("BL", "Livré par", "texte", "non", "deliveredByUserId", "utilisateurs", "delivery-commands.ts"),
    ("BL", "Reçu par (nom)", "texte", "non", "recipientName (200 car. max)", "—", "delivery-commands.ts"),
    ("BL", "Livrée ce jour", "nombre", "oui", "lines[].quantityBase (par ligne de commande)", "—", "delivery-commands.ts"),
    ("BL", "Signature photographiée", "photo", "non", "proofAttachmentId", "—", "delivery-commands.ts"),
    ("BL", "Observations", "texte", "non", "notes (2000 car. max)", "—", "delivery-commands.ts"),
    ("BT", "Expédié de / Vers", "texte", "oui", "from_location_id / to_location_id", "emplacements", "06-inventory.md"),
    ("BT", "Nature", "case", "oui", "transfer_kind", "—", "06-inventory.md"),
    ("BT", "Transporteur / véhicule", "texte", "non", "carrier_user_id ou carrier_name", "—", "06-inventory.md"),
    ("BT", "Quantité reçue ; écart ; motif", "nombres", "à la réception", "lines[].receivedQtyBase ; discrepancyReasonCodeId", "écart = perte en transit à justifier ; alerte si non rapproché : 48 h", "stock-transfer-commands.ts ; system-settings.ts"),
    ("FP", "Catégorie", "case", "oui", "category", "MORTALITE, CASSE, DETERIORATION, IMPROPRE, DESTRUCTION, INEXPLIQUEE, VOL_SUSPECTE, ECART_TRANSFERT", "06-inventory.md"),
    ("FP", "Produit, quantité, lot", "lignes", "oui", "product_id, quantity_base (> 0), lot_id ; production_lot_id OBLIGATOIRE si MORTALITE", "—", "06-inventory.md"),
    ("FP", "Cause, circonstances", "texte", "si INEXPLIQUEE ou VOL_SUSPECTE", "comment (2000 car. max) ; reason_code_id", "—", "06-inventory.md"),
    ("FP", "Personne responsable", "texte", "non", "responsibility_user_id", "—", "06-inventory.md"),
    ("FP", "Photo", "case", "selon politique", "pièce jointe STOCK_LOSS, nature PHOTO", "mortalité : politique MORTALITY", "ADR-046"),
    ("FP", "Consommation : imputé à", "texte", "si consommation", "cost_object_type, cost_object_id, cost_type", "lot, site, PDV", "06-inventory.md"),
    ("FI", "Emplacement ; type", "texte ; case", "oui", "location_id ; count_type (OPENING pour l'inventaire d'ouverture)", "—", "06-inventory.md"),
    ("FI", "Produit, lot, stock physique compté, motif d'écart", "lignes", "oui", "productId, lotId, countedQtyBase, varianceReasonCodeId, comment, countedAt", "validation au-delà de inventory.count_approval_threshold_xaf (25 000 XAF)", "count-commands.ts ; system-settings.ts"),
    ("FL", "Lot de production n°", "texte", "oui", "productionLotId", "lots de production (LOT-{site}-{année}-{séquence})", "daily-commands.ts"),
    ("FL", "Mortalité : nombre ; cause", "nombre ; code", "non", "quantity (entier > 0) ; reasonCodeId ; comment", "—", "daily-commands.ts"),
    ("FL", "Pesée : échantillon ; poids moyen (g)", "nombres", "non", "sampleSize (entier > 0) ; avgWeightG ; totalWeightKg ; source = MANUAL", "—", "daily-commands.ts"),
    ("FL", "Œufs : collectés / cassés / non conformes / commercialisables / à couver", "nombres", "pondeuses", "collected_qty, broken_qty, nonconforming_qty, marketable_qty, hatching_qty (collectés = somme des quatre autres)", "—", "07-production.md ; glossaire"),
    ("FL", "Aliment : produit, quantité", "texte ; nombre", "non", "consommation d'intrant imputée au lot", "—", "06-inventory.md"),
    ("CC", "Client", "texte", "oui", "customerId", "comptes clients", "order-commands.ts"),
    ("CC", "Adresse de livraison ; livraison souhaitée le", "texte ; date", "non", "deliveryAddress (1000 car. max) ; requestedDeliveryDate", "—", "order-commands.ts"),
    ("CC", "Lignes (désignation, unité, quantité, prix convenu)", "lignes", "oui", "lines[] (mêmes champs qu'une ligne de vente)", UNITES, "order-commands.ts"),
    ("CC", "Acompte versé ; mode", "montant ; case", "non", "advancePayments[] : methodCode, amountXaf, reference", MOY, "order-commands.ts"),
    ("CC", "Numéro imprimé", "recopié", "non", "localRef", "—", "order-commands.ts"),
    ("VC", "Client existant / nouveau prospect", "case", "oui", "customerId (existant) ; création d'un compte client au stade PROSPECT (nouveau)", "—", "glossaire"),
    ("VC", "Résultat de la visite", "case", "oui", "outcome_reason_code_id", "COMMANDE_PRISE (Commande prise), VENTE_REALISEE (Vente réalisée sur place), A_RELANCER (À relancer), PAS_INTERESSE (Pas intéressé), ABSENT (Client absent), AUTRE (Autre : commentaire obligatoire)", "crm-references.ts"),
    ("VC", "Commentaire", "texte", "si AUTRE", "notes", "—", "crm-references.ts"),
    ("VC", "Prochaine action ; le", "texte ; date", "non", "next_action_note ; next_action_at", "—", "04-crm-fieldwork.md"),
    ("VC", "Position GPS", "(non écrite sur le papier)", "non", "position (lat, lng) ; distance au client", "visite à moins de crm.visit.max_distance_m (500 m)", "system-settings.ts"),
    ("VC", "Produit d'intérêt", "texte", "non", "PROPOSITION P-09 : colonne du compte client (aucune aujourd'hui)", "catalogue produits", "Correspondances Excel : « hors modèle client »"),
    ("VC", "Potentiel par jour (quantité, unité)", "nombre", "non", "PROPOSITION P-09 : colonnes du compte client", UNITES, "Correspondances Excel"),
    ("VC", "Paiement habituel", "case", "non", "PROPOSITION P-09 : colonne du compte client", MOY, "Correspondances Excel"),
    ("VC", "Différé demandé (oui / non)", "case", "non", "PROPOSITION P-09 : demande distincte de credit_allowed (décision du responsable, existante)", "—", "schéma crm_customers"),
    ("VC", "Fournisseur ou concurrent actuel ; produit ; prix constaté ; unité", "texte ; nombre", "non", "PROPOSITION P-08 : table fille de la visite (observation de marché)", "—", "Correspondances Excel : « veille concurrentielle »"),
    ("BH", "Commercial ; semaine du … au …", "texte ; dates", "oui", "PROPOSITION P-07 : utilisateur, lundi de la semaine", "—", "Correspondances Excel (feuille Bilans_Hebdo)"),
    ("BH", "Visites sur place ; nouveaux prospects et clients ; CA ; encaissé ; casses et pertes", "nombres", "calculés", "calculés par l'application depuis les visites, ventes, encaissements, pertes", "—", "principe : aucun indicateur saisi"),
    ("BH", "Réussites ; difficultés ; besoins et demandes à la direction ; priorités ; informations du marché", "texte", "non", "PROPOSITION P-07 : champs de texte du bilan", "—", "Correspondances Excel"),
    ("BH", "Vu par le responsable ; commentaire", "case ; texte", "non", "PROPOSITION P-07 : validation du responsable", "—", "Correspondances Excel"),
    ("PT", "Zone / mission", "texte", "oui", "declared_zone_id", "zones", "04-crm-fieldwork.md"),
    ("PT", "Arrivée / départ", "heures", "oui", "started_at / ended_at", "géorepère : fieldwork.geofence_radius_m (500 m)", "04-crm-fieldwork.md ; system-settings.ts"),
    ("PT", "Codes d'absence (A, M, C, R, MI)", "code", "non", "(aucun équivalent : usage interne du papier)", "—", "—"),
]
CHAMPS_COLS = ["Pièce", "Rubrique de la pièce papier", "Type de saisie", "Obligatoire dans l'application", "Champ, colonne ou propriété de commande", "Référentiel / valeurs / seuil", "Source lue"]

REFERENTIELS = [
    ("Moyens de paiement", "ESPECES", "Espèces", "sans référence"),
    ("Moyens de paiement", "MOBILE_MONEY_ORANGE", "Orange Money", "référence exigée"),
    ("Moyens de paiement", "MOBILE_MONEY_MTN", "MTN Mobile Money", "référence exigée"),
    ("Moyens de paiement", "VIREMENT", "Virement bancaire", "référence exigée"),
    ("Moyens de paiement", "CHEQUE", "Chèque", "référence exigée (n° du chèque)"),
    ("Catégories de dépense", "ALIMENT_HORS_STOCK", "Aliment hors stock", ""),
    ("Catégories de dépense", "VETERINAIRE", "Vétérinaire", ""),
    ("Catégories de dépense", "TRANSPORT", "Transport", ""),
    ("Catégories de dépense", "CARBURANT", "Carburant", ""),
    ("Catégories de dépense", "ENERGIE_EAU", "Énergie et eau", ""),
    ("Catégories de dépense", "MAIN_OEUVRE_OCCASIONNELLE", "Main-d'œuvre occasionnelle", ""),
    ("Catégories de dépense", "ENTRETIEN", "Entretien", ""),
    ("Catégories de dépense", "FOURNITURES", "Fournitures", ""),
    ("Catégories de dépense", "DIVERS", "Divers", ""),
    ("Catégories de perte", "MORTALITE", "mortalité", "lot de production obligatoire"),
    ("Catégories de perte", "CASSE", "casse", ""),
    ("Catégories de perte", "DETERIORATION", "détérioration", ""),
    ("Catégories de perte", "IMPROPRE", "impropre à la vente", ""),
    ("Catégories de perte", "DESTRUCTION", "destruction", ""),
    ("Catégories de perte", "INEXPLIQUEE", "inexpliquée", "commentaire obligatoire"),
    ("Catégories de perte", "VOL_SUSPECTE", "vol suspecté", "commentaire obligatoire"),
    ("Catégories de perte", "ECART_TRANSFERT", "écart de transfert", ""),
    ("Résultats de visite", "COMMANDE_PRISE", "Commande prise", "code figé"),
    ("Résultats de visite", "VENTE_REALISEE", "Vente réalisée sur place", ""),
    ("Résultats de visite", "A_RELANCER", "À relancer", ""),
    ("Résultats de visite", "PAS_INTERESSE", "Pas intéressé", ""),
    ("Résultats de visite", "ABSENT", "Client absent", ""),
    ("Résultats de visite", "AUTRE", "Autre", "commentaire obligatoire"),
    ("Motifs d'annulation", "ERREUR_SAISIE", "Erreur de saisie", ""),
    ("Motifs d'annulation", "DOUBLON", "Doublon", ""),
    ("Motifs d'annulation", "DEMANDE_CLIENT", "Demande du client", ""),
    ("Motifs d'annulation", "AUTRE", "Autre", "commentaire obligatoire"),
    ("Unités", "TETE · OEUF · PIECE · KG · G · L · SAC · PLATEAU · CARTON · DOSE", "tête, œuf, pièce, kg, g, litre, sac (50 kg), plateau (30 œufs), carton, dose", "facteur de conversion par produit"),
    ("Types de compte de trésorerie", "CAISSE_PDV · CAISSE_UTILISATEUR · CAISSE_CENTRALE · MOBILE_MONEY · BANQUE", "caisse du PDV, caisse d'un utilisateur, caisse centrale, mobile money, banque", ""),
]
REF_COLS = ["Liste", "Code de l'application", "Libellé à employer sur le papier", "Remarque"]

SEUILS = [
    ("finance.expense_photo_threshold_xaf", "10 000 XAF", "Dépense : justificatif exigé (comparaison stricte)", "BD : case « justificatifs joints »"),
    ("finance.expense_approval_threshold_xaf", "50 000 XAF", "Dépense : validation exigée", "BD : visa « autorisé par »"),
    ("finance.pos_expense_max_xaf", "50 000 XAF", "Plafond d'une dépense en caisse de PDV", "BD"),
    ("finance.supplier_payment_approval_xaf", "500 000 XAF", "Paiement fournisseur : décision de la Direction", "BD (paiement fournisseur)"),
    ("procurement.po_approval_threshold_xaf", "500 000 XAF", "Bon de commande : approbation", "BC : signature « approuvé par »"),
    ("procurement.receipt_photo_threshold_xaf", "100 000 XAF", "Réception : photo du bon de livraison (toujours exigée sans bon de commande)", "BR"),
    ("procurement.receipt_over_tolerance_pct", "0 %", "Réception : tolérance au-delà du commandé", "BR : « Accepté » ne dépasse pas le reliquat"),
    ("inventory.count_approval_threshold_xaf", "25 000 XAF", "Inventaire : validation de l'écart", "FI"),
    ("inventory.transfer_unmatched_hours", "48 h", "Transfert non rapproché : signalement", "BT : renvoyer le feuillet de réception sous 48 h"),
    ("finance.cash_transfer_unreceived_hours", "48 h", "Remise de fonds non reçue : alerte", "RF"),
    ("finance.cash_holding_max_xaf / _days", "200 000 XAF / 3 jours", "Encaisse maximale avant remise", "FJC, RF"),
    ("finance.cash_session_force_close_hours", "24 h", "Session de caisse ouverte : clôture d'office possible", "FJC"),
    ("finance.cash_variance_critical_xaf", "10 000 XAF", "Écart de caisse critique (alerte)", "FJC"),
    ("finance.cash_variance_tolerance_xaf", "0 XAF", "Tout écart de caisse est soumis à validation de la Finance", "FJC"),
    ("sales.direct_cancel_minutes", "15 min", "Délai d'annulation directe d'une vente (AV-030)", "BV : voir la règle d'annulation de l'application"),
    ("sync.justification_after_hours", "24 h", "Heure métier antérieure : justification demandée (AV-078, OUVERT)", "Toutes pièces : délai de report"),
    ("sync.backdate_max_hours", "72 h", "Heure métier antérieure : plafond, sauf validation (AV-078, OUVERT)", "Toutes pièces : délai de report"),
    ("fieldwork.geofence_radius_m", "500 m", "Rayon du pointage", "PT, VC"),
    ("offline.max_autonomy_hours / warning_hours", "168 h / 48 h", "Autonomie hors ligne d'un appareil", "Procédure de secours"),
]
SEUILS_COLS = ["Paramètre", "Valeur initiale (seed)", "Effet", "Pièce concernée"]

ECARTS = [
    ("E1", "L'application n'émet pas de facture client (AV-041 : prix TTC, reçu numérique simple, taxes à 0).", "Facture hors application pour l'instant (D3).", "P-03"),
    ("E2", "Numéro imprimé de la pièce : `local_ref` (20 car.) EXISTE sur 13 tables (vente, commande, encaissement, livraison, annulation de vente, demande d'achat, bon de commande, réception, transfert, perte, inventaire, dépense, remise de fonds) ; acceptation par la commande vérifiée pour vente, commande, encaissement, livraison, réception, demande d'achat, dépense. Il manque sur visite, session de travail, session de caisse, paiement et facture fournisseur, pesée, collecte d'œufs. Attention : `local_ref` est aujourd'hui une référence d'appareil générée.", "Un champ « pièce papier » commun à toutes les opérations (D2). Colonne dédiée ou généralisation de `local_ref` : à arbitrer par la session de développement.", "P-02"),
    ("E3", "Comptage par coupures : la session de caisse ne garde que les totaux.", "Section supprimée de la feuille de caisse (décision de la Direction).", "—"),
    ("E4", "Fiche de présence pour tout le personnel : le pointage de l'application est géolocalisé et réservé aux commerciaux de terrain.", "Table « présence » simple alimentée chaque semaine (D6).", "P-06"),
    ("E5", "Journal des encaissements et journal de caisse : états reconstructibles depuis les mouvements de trésorerie.", "Un journal par compte de trésorerie, imprimé par l'application (D4, D5).", "P-04"),
    ("E6", "Seuils paramétrables : non imprimés sur les pièces.", "Inchangé : voir l'onglet Seuils.", "—"),
    ("E7", "Fenêtre d'antidatage (AV-078) non encore appliquée à toutes les commandes.", "24 h puis 72 h au maximum (D12).", "P-11"),
    ("E8", "Bilan hebdomadaire absent de l'application et du papier (l'Excel le mappe sur `internal_notes`, qui est la note de direction).", "Fiche BH + écran (D8).", "P-07"),
    ("E9", "Veille marché (fournisseur ou concurrent, prix constaté) : aucune cible dans le schéma.", "Rubriques structurées (D9).", "P-08"),
    ("E10", "Profil commercial du client : produit d'intérêt, potentiel, mode de paiement habituel absents ; `credit_allowed` est une décision, non une demande.", "Rubriques ajoutées (D10).", "P-09"),
    ("E11", "Reprise des données : modèles CSV prévus (AV-072, ouvert).", "Écran de report + import par lots (D11).", "P-10"),
    ("E12", "Registre des carnets : pas de suivi des numéros dans l'application.", "Registre papier + contrôle de continuité par les numéros saisis (D7).", "P-05"),
    ("E13", "Le reçu d'une vente est déjà imprimable depuis la PWA (ECR-VEN-05, `window.print()`).", "Impression de toutes les opérations dans la mise en page des pièces, depuis la PWA (D1).", "P-01"),
]
ECARTS_COLS = ["N°", "Constat", "Orientation retenue par le porteur", "Proposition (PROPOSITION-INTEGRATION-APPLICATION.md)"]

REPORT = [
    ("1", "Prendre les pièces dans l'ordre des numéros.", "Vérifier que le numéro suivant est bien celui du carnet en cours (registre RC)."),
    ("2", "Saisir l'opération avec la date et l'heure écrites sur la pièce.", "Elles deviennent occurred_at (heure métier) ; la date de saisie reste distincte (client_created_at, received_at)."),
    ("3", "Recopier le numéro imprimé dans « référence locale ».", "20 caractères au plus, ex. RE-000123."),
    ("4", "Joindre la photo de la pièce.", "Pièce justificative numérique (ADR-012, ADR-046)."),
    ("5", "Renseigner le cartouche de la souche.", "Saisi le, par, n° opération, photo jointe."),
    ("6", "Doublon.", "Ne pas ressaisir un fait déjà saisi ; écrire « déjà saisi — n° … » sur la souche."),
    ("7", "Annulation.", "Annuler dans l'application (contre-écriture) et annuler la pièce papier."),
]
REPORT_COLS = ["Étape", "Action", "Précision"]


def feuille(wb, titre, entetes, lignes, largeurs):
    ws = wb.create_sheet(titre)
    ws.append(entetes)
    for l in lignes:
        ws.append(list(l))
    fin = Side(style="thin", color="B8C7BF")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF", name="Segoe UI", size=10)
        c.fill = PatternFill("solid", fgColor=VERT)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.font = Font(name="Segoe UI", size=9)
            c.border = Border(bottom=fin)
    for i, w in enumerate(largeurs, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


wb = Workbook()
wb.remove(wb.active)
lisez = wb.create_sheet("Lisez-moi")
for ligne in [
    "CORRESPONDANCE PIÈCES PAPIER ↔ APPLICATION — GIC AGROPELC (20/10/2026)",
    "",
    "Pour chaque rubrique d'une pièce papier : le champ de l'application qui la reçoit, les valeurs du référentiel et les seuils.",
    "Sources lues dans le dépôt, en lecture seule : db/seeds, apps/server/src/modules/*/application/commands, docs/03-data/dictionnaire, glossaire, A-VALIDER.",
    "Les valeurs de seuils sont les valeurs INITIALES des paramètres système : elles sont modifiables dans l'application.",
    "Un champ marqué « non trouvé » ou « à confirmer » n'a pas été vérifié dans le code : ne pas le considérer comme acquis.",
    "Onglets : Pièces · Champs · Référentiels · Seuils · Écarts et orientations · Report. Les lignes « PROPOSITION » ne décrivent pas un champ existant : elles renvoient à la proposition d'intégration.",
]:
    lisez.append([ligne])
lisez["A1"].font = Font(bold=True, size=13, color=VERT, name="Segoe UI")
lisez.column_dimensions["A"].width = 140
feuille(wb, "Pièces", PIECES_COLS, PIECES, [12, 30, 60, 52, 18, 30])
feuille(wb, "Champs", CHAMPS_COLS, CHAMPS, [8, 44, 16, 22, 56, 60, 32])
feuille(wb, "Référentiels", REF_COLS, REFERENTIELS, [28, 52, 52, 34])
feuille(wb, "Seuils", SEUILS_COLS, SEUILS, [44, 22, 64, 44])
feuille(wb, "Écarts et orientations", ECARTS_COLS, ECARTS, [8, 90, 70, 22])
feuille(wb, "Report", REPORT_COLS, REPORT, [8, 60, 84])
wb.save(SORTIE)
print("ok", SORTIE, len(CHAMPS), "lignes de champs")
